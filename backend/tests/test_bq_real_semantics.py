"""REAL BigQuery check of the SQL the app sends (exact-duplicate gate + day x SKU x FC aggregation), on INLINE rows:
the source is a literal `UNNEST(ARRAY<STRUCT<...>>[...])` run through the history-SQL path, so no table is read
(0 bytes) and nothing is written anywhere. Opt-in: needs Google credentials of this Windows account (ADC) and
`DF_REAL_BQ=1`; skipped otherwise.

    cd backend; $env:DF_REAL_BQ = "1"; ..\\.venv\\Scripts\\python.exe -m pytest -q -p no:cacheprovider tests\\test_bq_real_semantics.py
"""
import os
from datetime import date

import pytest
from openpyxl import load_workbook

from app import settings_service as ss
from app.bq import connector, dup_report
from app.messages import AppError

pytestmark = pytest.mark.skipif(os.environ.get("DF_REAL_BQ") != "1",
                                reason="real BigQuery (ADC of this Windows account): set DF_REAL_BQ=1")

# 1.2.9: total_open_hours is a mandatory D-1 input; the same value on every line changes no duplicate group
COLS = ("so_code STRING, created_date DATE, sku_code STRING, fc_code STRING, demand_quantity FLOAT64, "
        "outbound_quantity FLOAT64, sales_price FLOAT64, order_status STRING, seller_code STRING, seller_class STRING, "
        "refresh_time STRING, total_open_hours FLOAT64, customer_id INT64")
D1, D2, OUT = "DATE '2026-10-01'", "DATE '2026-10-02'", "DATE '2026-09-30'"
ROWS = [  # so_code, day, sku, qty, outbound, price, status
    ("'O1'", D1, "A", 5, 5, 10.0, "COMPLETED"), ("'O2'", D1, "A", 8, 8, 10.0, "COMPLETED"),
    ("'O3'", D1, "A", 2, 2, 10.0, "COMPLETED"),                                     # 1: three orders, sum 15
    ("'O4'", D1, "B", 3, 3, 5.0, "COMPLETED"), ("'O4'", D1, "B", 4, 4, 5.0, "COMPLETED"),     # 2: same key, qty differs
    ("'O5'", D1, "C", 2, 2, 10.0, "COMPLETED"), ("'O5'", D1, "C", 2, 2, 12.62, "COMPLETED"),  # 3: price differs
    ("'O6'", D1, "D", 6, 6, 7.0, "COMPLETED"), ("'O6'", D1, "D", 6, 5, 7.0, "COMPLETED"),     # 4: outbound differs
    ("'O7'", D1, "E", 20, 20, 3.0, "COMPLETED"), ("'O7'", D1, "E", 20, 20, 3.0, "COMPLETED"),  # 5: exact duplicate
    ("NULL", D1, "F", 1, None, 9.0, "WAIT_TO_CONFIRM"), ("NULL", D1, "F", 2, None, 9.0, "WAIT_TO_CONFIRM"),
    ("NULL", D1, "F", 10, None, 9.0, "WAIT_TO_CONFIRM"),                             # 7: NULL so_code, kept, summed
    # 8 (owner rule 21/10/2026): pending orders without order code identical in every mapped column are NOT exact
    # duplicates when another source column tells them apart (here customer_id, unmapped - fix 7.1 compares their whole
    # source row): kept and summed, counted apart (connector.PENDING_STATUSES)
    ("NULL", D1, "G", 3, None, 4.0, "WAIT_TO_CONFIRM"), ("NULL", D1, "G", 3, None, 4.0, "WAIT_TO_CONFIRM"),
    ("'O8'", D1, "H", 1, None, 2.0, "COMPLETED"), ("'O8'", D1, "H", 1, 0, 2.0, "COMPLETED"),  # 8: NULL is not 0
    ("'O9'", D2, "A", 1, 1, 10.0, "COMPLETED"),                                      # another day
    ("'O0'", OUT, "A", 100, 100, 10.0, "COMPLETED"),                                 # outside the window
    ("NULL", D1, "I", 4, None, 6.0, "CONFIRMED"), ("NULL", D1, "I", 4, None, 6.0, "CONFIRMED"),  # 9: NULL=NULL
]
EXTRA = {10, 20}          # second copy of the two exact duplicates (indexes in ROWS)


def _sql(rows) -> str:
    def v(x):
        return "NULL" if x is None else repr(x)
    # every line has its own customer (as in the source): never part of the comparison of a non-pending row
    body = ",\n  ".join(f"({so}, {d}, '{sku}', 'F1', {v(q)}, {v(ob)}, {v(p)}, '{st}', 'SE1', 'X', 'r', 24.0, {100 + i})"
                        for i, (so, d, sku, q, ob, p, st) in enumerate(rows))
    return f"SELECT * FROM UNNEST(ARRAY<STRUCT<{COLS}>>[\n  {body}\n])"


def _setup(rows):
    ss.update("bigquery", {"dataset_id": "gold_buymed_vn2", "source_table": "demand_planing_lightgbm",
                           "location": "asia-southeast1", "partition_column": "",
                           "maximum_bytes_billed": 10 * 1024 ** 2})
    sql = _sql(rows)
    cols = connector.sql_dry_run(sql)["columns"]                  # real dry run: the output schema of the rows
    sug = connector.suggest_mapping(cols)
    ss.update("bigquery", {"field_mapping": sug["field_mapping"], "reference_mapping": sug["reference_mapping"]})
    return {"mode": "SQL", "kind": "history", "version_id": 0, "sha256": "inline-test", "sql": sql,
            "key_columns": ["so_code", "sku_code", "fc_code", "created_date"], "key_duplicate_policy": "SUM_ALL"}


def test_real_bigquery_exact_duplicates_only_when_every_mapped_column_is_equal():
    snap = _setup(ROWS)
    chk = connector.duplicate_check(date(2026, 10, 1), date(2026, 10, 2), snap, context="real-bq-test")
    assert chk["rows_checked"] == 21 and chk["estimated_bytes"] == 0
    assert (chk["duplicate_groups"], chk["duplicate_rows"], chk["extra_duplicate_rows"]) == (2, 4, 2)
    found = {(g["values"]["order_id"], g["values"]["sku_id"], g["occurrences"]) for g in chk["groups"]}
    assert found == {(None, "I", 2), ("O7", "E", 2)}         # a NULL so_code of another status still compares equal
    assert chk["pending_no_order_code"] == {"applies": True, "statuses": ["WAIT_TO_CONFIRM"], "rows": 5, "qty": 19.0,
                                            "identical_groups": 1, "identical_rows": 2, "duplicate_groups": 0,
                                            "duplicate_rows": 0}
    assert [c["column"] for c in chk["columns"]][:5] == ["so_code", "created_date", "sku_code", "fc_code",
                                                          "demand_quantity"]
    assert {"sales_price", "outbound_quantity", "order_status", "seller_code"} <= {c["column"] for c in chk["columns"]}
    assert set(chk["not_compared"]) == {"seller_class", "refresh_time", "customer_id"}
    wb = load_workbook(dup_report.folder() / chk["report"]["file"])
    assert wb["Duplicate_Rows"].max_row - 1 == 4 and wb["Duplicate_Groups"].max_row - 1 == 2
    with pytest.raises(AppError) as e:                             # the load stops before any aggregation
        connector.fetch_daily(date(2026, 10, 1), date(2026, 10, 2), snap)
    assert e.value.code == "BQ_DATA_DUPLICATES"


def test_real_bigquery_clean_rows_are_summed_per_day_sku_fc():
    snap = _setup([r for i, r in enumerate(ROWS) if i not in EXTRA])
    res = connector.fetch_daily(date(2026, 10, 1), date(2026, 10, 2), snap)
    got = {(str(r["order_date"]), r["sku_id"]): float(r["qty"]) for r in res["rows"]}
    assert got == {("2026-10-01", "A"): 15.0, ("2026-10-01", "B"): 7.0, ("2026-10-01", "C"): 4.0,
                   ("2026-10-01", "D"): 12.0, ("2026-10-01", "E"): 20.0, ("2026-10-01", "F"): 13.0,
                   ("2026-10-01", "G"): 6.0, ("2026-10-01", "H"): 2.0, ("2026-10-01", "I"): 4.0,
                   ("2026-10-02", "A"): 1.0}                  # G: both pending orders kept
    assert res["rows_read"] == 19 and res["rows_valid"] == 19 and res["qty_valid"] == 84.0
    assert res["exact_duplicate_rows"] == 0 and res["key_duplicate_rows"] == 7        # repeated keys: summed
    a = next(r for r in res["rows"] if r["sku_id"] == "A" and str(r["order_date"]) == "2026-10-01")
    assert a["n_orders"] == 3 and a["n_rows"] == 3
