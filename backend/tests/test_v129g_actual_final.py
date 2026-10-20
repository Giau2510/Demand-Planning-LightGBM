"""1.2.9 (seventh batch, owner rule 20/10/2026): the actual of a completed day is FINALIZED (frozen) from the successful
D-1 load that wrote it and never changes afterwards; the history page's KPI use the finalized rows only.

Kind of evidence: integration on isolated temporary databases (conftest); the D-1 path uses the fake BigQuery client
(FakeBQ), the other cases write demand rows exactly as a committed D-1 load does (one loaded_at = the load's synced_at,
its audit D1_SYNC in the same transaction). Nothing here reaches production or BigQuery."""
from __future__ import annotations

import io
import json
import sqlite3
from datetime import date, datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from app import actual_final as af
from app import data_checks, db, paths, timeutil
from app import forecast_history as fh
from app.ml import metrics as mt
from test_v125_forecast_history import SCOPE, _run, _snapshot, _utc

KEYS = (("S1", "HN"), ("S2", "HN"), ("S1", "BD"))
FINAL_TABLES = ("actual_finals", "actual_final_values")


def reverse_0009(c) -> None:
    """The reverse statements written in the header of migration 0009 (for a database COPY only)."""
    for (t,) in c.execute("SELECT name FROM sqlite_master WHERE type='trigger' AND (name LIKE 'trg_afn_%' "
                          "OR name LIKE 'trg_afv_%')").fetchall():
        c.execute(f"DROP TRIGGER {t}")
    c.execute("DROP TABLE actual_final_values")
    c.execute("DROP TABLE actual_finals")
    c.execute("DELETE FROM schema_migrations WHERE version = 9")


@pytest.fixture
def frozen():
    """20/10/2026 10:00 Vietnam time (03:00 UTC); set_now_override can move it."""
    timeutil.set_now_override(lambda: datetime(2026, 10, 20, 3, 0, tzinfo=timezone.utc))
    yield
    timeutil.set_now_override(None)


def _at(vn: str):
    t = datetime.fromisoformat(vn).replace(tzinfo=timeutil.VN_TZ)
    timeutil.set_now_override(lambda: t.astimezone(timezone.utc))


def _load(days: dict, synced_vn: str, window: tuple | None = None, mode: str = "FALLBACK_TIME_VERIFIED",
          backfill: bool = False, audit: bool = True, scope: str = SCOPE) -> dict:
    """One committed D-1 load as sync.run_d1_sync writes it: the days of its window replaced, every row with ONE
    loaded_at = synced_at, and its audit D1_SYNC (same synced_at, a unique load id) in the same transaction."""
    at = _utc(synced_vn)
    w0, w1 = window or (min(days), max(days))
    lid = db.new_id("D1L")
    with db.tx() as c:
        c.execute("DELETE FROM demand_daily WHERE scope_key=? AND order_date BETWEEN ? AND ?", (scope, w0, w1))
        c.executemany("INSERT INTO demand_daily(scope_key, data_source, order_date, sku_id, fc_id, qty, n_orders, loaded_at) "
                      "VALUES (?, 'BIGQUERY', ?,?,?,?,?,?)",
                      [(scope, d, s, f, q, 1, at) for d, vals in days.items() for (s, f), q in vals.items()])
        if audit:
            db.audit(c, "D1_SYNC", "scope", scope, {"synced_at": at, "watermark": w1, "window_start": w0,
                                                    "window_end": w1, "mode": mode, "backfill": backfill,
                                                    "load_id": lid, "job_id": f"JOB-{lid}", "full_reload": False})
    return {"load_id": lid, "synced_at": at}


def _final(day, scope=SCOPE):
    with db.read() as c:
        h = db.one(c, "SELECT * FROM actual_finals WHERE scope_key=? AND business_date=?", (scope, day))
        v = {(r[0], r[1]): r[2] for r in c.execute("SELECT sku_id, fc_id, qty FROM actual_final_values WHERE final_id=?",
                                                    ((h or {}).get("id"),))}
    return h, v


def _count(sql, p=()):
    with db.read() as c:
        return c.execute(sql, p).fetchone()[0]


def _fp():
    """Forecast side: published values and the D+1 history (must never move because of a finalization)."""
    with db.read() as c:
        return [c.execute(f"SELECT COUNT(*), TOTAL(yhat) FROM {t}").fetchone()
                for t in ("forecast_values", "forecast_history_values")] + \
               [c.execute("SELECT COUNT(*) FROM forecast_history_sessions").fetchone()[0],
                tuple(c.execute("SELECT run_date, forecast_run_id FROM forecast_history_days ORDER BY run_date"))]


def _q(**kw):
    with db.read() as c:
        qy = fh.Query(c, kw.pop("scope", SCOPE), **kw)
        daily = qy.daily()
        return qy.kpi(daily), daily, list(qy.rows(limit=None))


# ================================================================== 1, 4, 5, 6, 10 - the real D-1 path
def test_a_successful_d1_load_finalizes_the_completed_day_and_a_later_load_never_changes_it(client):
    from test_v129_backfill_history_patterns import SCOPE as BQS, _production, _row
    from app import sync
    timeutil.set_now_override(None)
    today = timeutil.today_vn()
    y = today - timedelta(days=1)
    rid = _run(f"{y - timedelta(days=1)} 09:00", scope=BQS, origin=str(y - timedelta(days=1)),
               keys=(("A", "F1"), ("B", "F1"), ("C", "F1")))
    fh.capture(rid)
    fake = _production([_row(y, "A", 7), _row(y, "B", 3)])
    f0 = _fp()
    res = sync.run_d1_sync("J1")
    h, v = _final(str(y), BQS)
    assert (h["status"], h["sealed"], h["reason"], h["d1_load_id"], h["d1_synced_at"], h["d1_mode"]) == (
        "FINALIZED", 1, "D1_LOAD", res["load_id"], res["synced_at"], "SIGNAL")
    assert (h["n_rows"], h["total_qty"], h["timezone"], h["finalized_by"], h["job_id"]) == (
        2, 10.0, "Asia/Ho_Chi_Minh", "AFTER_D1", "J1")
    assert v == {("A", "F1"): 7.0, ("B", "F1"): 3.0} and h["values_sha256"] == af._sha(
        [("A", "F1", 7.0, 1), ("B", "F1", 3.0, 1)])
    assert _count("SELECT COUNT(*) FROM actual_finals WHERE business_date >= ?", (str(today),)) == 0   # never today
    assert _fp() == f0                                                     # forecast and D+1 history untouched
    # a later successful load reports another quantity and omits B: the finalized day does not move
    fake.responses["daily"] = [_row(y, "A", 9)]
    sync.run_d1_sync("J2")
    with db.read() as c:
        assert {r[0]: r[1] for r in c.execute("SELECT sku_id, qty FROM demand_daily WHERE scope_key=? AND order_date=?",
                                              (BQS, str(y)))} == {"A": 9.0}           # the live data did change
    h2, v2 = _final(str(y), BQS)
    assert h2 == h and v2 == v                                             # same header, same values, same provenance
    assert _count("SELECT COUNT(*) FROM audit_log WHERE action='ACTUAL_FINALIZED'") == 1
    assert af.finalize(BQS, "BIGQUERY")["finalized"] == []                 # retry: nothing new
    assert _count("SELECT COUNT(*) FROM actual_final_values") == 2         # no duplicate
    with db.read() as c:
        qy = fh.Query(c, BQS)
        rows = {r["sku_id"]: r for r in qy.rows(limit=None)}
        k = qy.kpi(qy.daily())
    assert {s: (r["actual"], r["actual_status"], r["actual_source"]) for s, r in rows.items()} == {
        "A": (7.0, "COMPLETE", "FINAL"), "B": (3.0, "COMPLETE", "FINAL"), "C": (0.0, "COMPLETE", "FINAL")}
    assert k["actual_total"] == 10.0 and k["n_zero_absent"] == 1 and k["status"] == "COMPLETE"
    assert _fp() == f0


# ================================================================== 2 - a failed D-1 load finalizes nothing
def test_a_failed_d1_load_finalizes_nothing(client):
    from test_v129_backfill_history_patterns import SCOPE as BQS, _production, _row
    from app import sync
    from app.messages import AppError
    timeutil.set_now_override(None)
    y = timeutil.today_vn() - timedelta(days=1)
    _production([_row(y, "A", 5), _row(y, "A", 6)])                       # two output rows for one key: invalid
    with pytest.raises(AppError) as e:
        sync.run_d1_sync("J1")
    assert e.value.code == "BQ_AGGREGATE_INVALID"
    assert _count("SELECT COUNT(*) FROM demand_daily WHERE scope_key=?", (BQS,)) == 0
    assert _count("SELECT COUNT(*) FROM actual_finals") == 0 and _count("SELECT COUNT(*) FROM actual_final_values") == 0


# ================================================================== 3, 17 - no verified load, no finalization
def test_days_without_a_verified_load_are_blocked_with_a_reason_and_stay_provisional(frozen):
    _load({"2026-10-14": {("S1", "HN"): 1}}, "2026-10-15 05:00", mode="SOMETHING_ELSE")           # not an eligible mode
    _load({"2026-10-15": {("S1", "HN"): 2}}, "2026-10-16 05:00", window=("2026-10-16", "2026-10-16"))   # window misses it
    _load({"2026-10-16": {("S1", "HN"): 3}}, "2026-10-17 05:00")
    with db.tx() as c:                                                    # a second writer touched part of the day
        c.execute("UPDATE demand_daily SET loaded_at='2026-10-17T00:00:00+00:00' WHERE order_date='2026-10-16'")
        c.execute("INSERT INTO demand_daily(scope_key, data_source, order_date, sku_id, fc_id, qty, n_orders, loaded_at) "
                  "VALUES (?, 'BIGQUERY', '2026-10-16', 'S2', 'HN', 1, 1, '2026-10-16T22:00:00+00:00')", (SCOPE,))
    _load({"2026-10-17": {("S1", "HN"): 4}}, "2026-10-18 05:00", audit=False)                    # legacy: no audit
    _load({"2026-10-18": {("S1", "HN"): 5}}, "2026-10-18 23:30")                                 # read before the day ended
    out = af.finalize(SCOPE, "BIGQUERY")
    assert out["finalized"] == [] and not out["errors"]
    assert {b["business_date"]: b["reason"] for b in out["blocked"]} == {
        "2026-10-14": "D1_MODE_NOT_ELIGIBLE", "2026-10-15": "LOAD_WINDOW_DOES_NOT_COVER_DAY",
        "2026-10-16": "MIXED_LOADS", "2026-10-17": "NO_VERIFIED_D1_LOAD", "2026-10-18": "LOADED_BEFORE_DAY_ENDED"}
    assert _count("SELECT COUNT(*) FROM actual_final_values") == 0
    n_audit = _count("SELECT COUNT(*) FROM audit_log WHERE action='ACTUAL_FINALIZATION_BLOCKED'")
    again = af.finalize(SCOPE, "BIGQUERY")                                 # idempotent: nothing rewritten
    assert again["blocked"] == [] and again["unchanged"] == 5
    assert _count("SELECT COUNT(*) FROM audit_log WHERE action='ACTUAL_FINALIZATION_BLOCKED'") == n_audit
    r = _run("2026-10-16 09:00", origin="2026-10-16", keys=KEYS)          # D+1 = 17/10
    fh.capture(r)
    _k, daily, _rows = _q(run_from="2026-10-16", run_to="2026-10-16")
    assert (daily[0]["actual_status"], daily[0]["actual_reason"]) == ("PROVISIONAL", "D1_NOT_FINALIZED")
    assert daily[0]["actual_finalization"]["reason"] == "NO_VERIFIED_D1_LOAD" and daily[0]["actual_final"] is None
    # a later verified load that re-reads 17/10 makes it final: BLOCKED -> FINALIZED (never the reverse)
    _load({"2026-10-17": {("S1", "HN"): 6}}, "2026-10-19 05:00")
    res = af.finalize(SCOPE, "BIGQUERY")
    assert [x["business_date"] for x in res["finalized"]] == ["2026-10-17"]
    h, v = _final("2026-10-17")
    assert (h["status"], h["attempts"], v) == ("FINALIZED", 2, {("S1", "HN"): 6.0})


# ================================================================== 15 - Vietnam business day around midnight
def test_the_business_day_is_the_vietnam_date_around_midnight(frozen):
    _at("2026-10-20 00:30")                                               # = 2026-10-19 17:30 UTC
    _load({"2026-10-19": {("S1", "HN"): 1}}, "2026-10-19 23:30")          # Vietnam date of the load = the day itself
    assert {b["business_date"]: b["reason"] for b in af.finalize(SCOPE, "BIGQUERY")["blocked"]} == {
        "2026-10-19": "LOADED_BEFORE_DAY_ENDED"}
    _load({"2026-10-19": {("S1", "HN"): 2}}, "2026-10-20 00:05")          # UTC date still 19/10, Vietnam date 20/10
    _load({"2026-10-20": {("S1", "HN"): 9}}, "2026-10-20 00:10")          # today: never a candidate
    out = af.finalize(SCOPE, "BIGQUERY")
    assert [x["business_date"] for x in out["finalized"]] == ["2026-10-19"] and out["blocked"] == []
    assert _final("2026-10-19")[1] == {("S1", "HN"): 2.0} and _final("2026-10-20")[0] is None
    with db.read() as c:
        assert [i["business_date"] for i in af.plan(c, SCOPE, "BIGQUERY")["items"]] == []


# ================================================================== 9 - today stays provisional
def test_the_current_day_stays_provisional(frozen):
    r = _run("2026-10-19 09:00", origin="2026-10-19", keys=KEYS)          # D+1 = 20/10 = today
    fh.capture(r)
    _snapshot("2026-10-20", {("S2", "HN"): 3}, "2026-10-20 08:00")
    assert af.finalize(SCOPE, "BIGQUERY")["finalized"] == []
    k, daily, rows = _q(run_from="2026-10-19", run_to="2026-10-19")
    assert {(x["actual_status"], x["actual_reason"]) for x in rows} == {("PROVISIONAL", "D0_INTRADAY")}
    assert k["actual_total"] is None and k["provisional"]["actual_total"] == 3 and k["status"] == "NO_ACTUAL"


# ================================================================== 7, 8 - atomic per day, retried after a failure
def test_a_failed_finalization_rolls_back_its_day_alone_and_is_retried(frozen, monkeypatch):
    _load({"2026-10-17": {("S1", "HN"): 1}, "2026-10-18": {("S1", "HN"): 2, ("S2", "HN"): 3}}, "2026-10-19 05:00")
    real = af._sha
    calls = {"n": 0}

    def flaky(rows):
        calls["n"] += 1
        if calls["n"] == 2:
            raise OSError("disk hiccup")                                  # the second day fails mid-way
        return real(rows)

    with monkeypatch.context() as m:
        m.setattr(af, "_sha", flaky)
        out = af.finalize(SCOPE, "BIGQUERY")
    assert [x["business_date"] for x in out["finalized"]] == ["2026-10-17"]
    assert [x["business_date"] for x in out["errors"]] == ["2026-10-18"]
    assert _final("2026-10-18") == (None, {})                             # nothing of the failed day was kept
    assert _count("SELECT COUNT(*) FROM alerts WHERE code='ACTUAL_FINALIZE_FAILED' AND level='ERROR'") == 1
    h17 = _final("2026-10-17")
    retry = af.finalize(SCOPE, "BIGQUERY")
    assert [x["business_date"] for x in retry["finalized"]] == ["2026-10-18"] and not retry["errors"]
    assert _final("2026-10-18")[1] == {("S1", "HN"): 2.0, ("S2", "HN"): 3.0} and _final("2026-10-17") == h17


def test_the_database_refuses_any_change_to_a_finalized_actual_and_any_partial_day(frozen):
    _load({"2026-10-18": {("S1", "HN"): 2, ("S2", "HN"): 3}}, "2026-10-19 05:00")
    af.finalize(SCOPE, "BIGQUERY")
    h, _v = _final("2026-10-18")
    for sql in ("UPDATE actual_finals SET total_qty = 99 WHERE id = '{id}'",
                "UPDATE actual_finals SET sealed = 0 WHERE id = '{id}'",
                "UPDATE actual_finals SET status = 'BLOCKED' WHERE id = '{id}'",
                "UPDATE actual_finals SET business_date = '2026-10-17' WHERE id = '{id}'",
                "DELETE FROM actual_finals WHERE id = '{id}'",
                "UPDATE actual_final_values SET qty = 0 WHERE final_id = '{id}'",
                "DELETE FROM actual_final_values WHERE final_id = '{id}'",
                "INSERT INTO actual_final_values(final_id, sku_id, fc_id, qty) VALUES ('{id}', 'NEW', 'HN', 1)",
                "INSERT OR REPLACE INTO actual_final_values(final_id, sku_id, fc_id, qty) VALUES ('{id}', 'S1', 'HN', 7)",
                "INSERT OR REPLACE INTO actual_finals(id, scope_key, data_source, business_date, status, reason, timezone, "
                "first_attempt_at, last_attempt_at) VALUES ('X', '" + SCOPE + "', 'BIGQUERY', '2026-10-18', 'BLOCKED', 'R', "
                "'Asia/Ho_Chi_Minh', 'a', 'b')"):
        with pytest.raises(sqlite3.IntegrityError):
            with db.tx() as c:
                c.execute(sql.format(id=h["id"]))
    assert _final("2026-10-18")[0] == h
    with pytest.raises(sqlite3.IntegrityError):                          # a header is never born sealed
        with db.tx() as c:
            c.execute("INSERT INTO actual_finals(id, scope_key, data_source, business_date, status, reason, sealed, "
                      "timezone, first_attempt_at, last_attempt_at) VALUES ('Y', ?, 'BIGQUERY', '2026-10-10', 'BLOCKED', "
                      "'R', 1, 'Asia/Ho_Chi_Minh', 'a', 'b')", (SCOPE,))
    with pytest.raises(sqlite3.IntegrityError):                          # half a day can never be sealed
        with db.tx() as c:
            c.execute("INSERT INTO actual_finals(id, scope_key, data_source, business_date, status, reason, d1_synced_at, "
                      "n_rows, total_qty, values_sha256, timezone, first_attempt_at, last_attempt_at, finalized_at) "
                      "VALUES ('Z', ?, 'BIGQUERY', '2026-10-11', 'FINALIZED', 'D1_LOAD', 't', 2, 5, 'x', "
                      "'Asia/Ho_Chi_Minh', 'a', 'b', 'c')", (SCOPE,))
            c.execute("INSERT INTO actual_final_values(final_id, sku_id, fc_id, qty) VALUES ('Z', 'S1', 'HN', 2)")
            c.execute("UPDATE actual_finals SET sealed = 1 WHERE id = 'Z'")
    assert _count("SELECT COUNT(*) FROM actual_finals WHERE id IN ('Y', 'Z')") == 0
    dc = {x["check"]: x["status"] for x in data_checks.run(paths.db_path())["checks"]}
    assert dc["actual_final_sealed"] == "OK" and dc["actual_final_protection_triggers"] == "OK"


# ================================================================== 11-14, 16 - KPI / coverage / filters
def test_kpi_use_finalized_rows_only_and_report_the_other_groups_apart(frozen):
    for d in ("2026-10-17", "2026-10-18", "2026-10-19", "2026-10-20"):     # D+1 = 18/10 .. 21/10
        fh.capture(_run(f"{d} 09:00", origin=d, keys=KEYS))
    _load({"2026-10-18": {("S1", "HN"): 5, ("S2", "HN"): 30}}, "2026-10-19 05:00")          # S1|BD absent -> 0
    _load({"2026-10-19": {("S1", "HN"): 4, ("S2", "HN"): 1, ("S1", "BD"): 2}}, "2026-10-20 05:00", audit=False)
    _snapshot("2026-10-20", {("S2", "HN"): 3}, "2026-10-20 08:00")       # today
    af.finalize(SCOPE, "BIGQUERY")
    k, daily, rows = _q(run_from="2026-10-17", run_to="2026-10-20")
    by = {d["target_date"]: d for d in daily}
    assert [by[d]["actual_status"] for d in ("2026-10-18", "2026-10-19", "2026-10-20", "2026-10-21")] == [
        "COMPLETE", "PROVISIONAL", "PROVISIONAL", "NO_ACTUAL"]
    fin = [r for r in rows if r["actual_status"] == "COMPLETE"]
    prov = [r for r in rows if r["actual_status"] == "PROVISIONAL"]
    ref = mt.compute([r["actual"] for r in fin], [r["forecast_qty"] for r in fin])
    assert k["basis"] == "FINALIZED_ONLY" and k["status"] == "COMPLETE"
    assert k["actual_total"] == 35 and k["n_rows_compared"] == 3 and k["days_compared"] == 1
    assert k["wape"] == pytest.approx(ref["wape"]) and k["bias"] == pytest.approx(ref["bias"])
    assert k["deviation"] == sum(r["forecast_qty"] for r in fin) - 35 and k["n_zero_absent"] == 1
    pref = mt.compute([r["actual"] for r in prov], [r["forecast_qty"] for r in prov])
    assert k["provisional"]["actual_total"] == 7 + 3 and k["provisional"]["n_rows"] == 6
    assert k["provisional"]["wape"] == pytest.approx(pref["wape"]) and k["closed"]["n_rows"] == 0
    assert k["coverage"] == {"rows_total": 12, "rows_final": 3, "rows_provisional": 6, "rows_closed": 0,
                             "rows_no_actual": 3, "days_final": 1, "days_provisional": 2, "days_closed": 0,
                             "days_no_actual": 1}
    nope = [r for r in rows if r["target_date"] == "2026-10-21"]
    assert all(r["actual"] is None and r["deviation"] is None for r in nope)       # missing is never 0
    k1, _d, r1 = _q(run_from="2026-10-17", run_to="2026-10-20", actual_status="COMPLETE")
    assert {r["target_date"] for r in r1} == {"2026-10-18"} and k1["actual_total"] == 35
    k2, _d, r2 = _q(run_from="2026-10-17", run_to="2026-10-20", actual_status="PROVISIONAL", fc_id="HN")
    assert {r["actual_status"] for r in r2} == {"PROVISIONAL"} and {r["fc_id"] for r in r2} == {"HN"}
    assert k2["actual_total"] is None and k2["provisional"]["actual_total"] == 4 + 1 + 3
    k3, _d, r3 = _q(run_from="2026-10-17", run_to="2026-10-20", q="s2")
    assert {r["sku_id"] for r in r3} == {"S2"} and k3["actual_total"] == 30
    # a finalized day whose rows sum to 0: no WAPE, a clear note, still finalized
    _load({"2026-10-12": {("OTHER", "HN"): 0.0}}, "2026-10-13 05:00")
    fh.capture(_run("2026-10-11 09:00", origin="2026-10-11", keys=KEYS))
    af.finalize(SCOPE, "BIGQUERY")
    kz, _d, _r = _q(run_from="2026-10-11", run_to="2026-10-11")
    assert (kz["actual_total"], kz["wape"], kz["notes"], kz["n_zero_absent"]) == (0, None, ["SUM_ACTUAL_ZERO"], 3)


# ================================================================== API, Excel, read-only
def test_api_and_excel_show_the_finalized_basis_and_reading_writes_nothing(client, frozen):
    fh.capture(_run("2026-10-17 09:00", origin="2026-10-17", keys=KEYS))
    fh.capture(_run("2026-10-18 09:00", origin="2026-10-18", keys=KEYS))
    _load({"2026-10-18": {("S1", "HN"): 5, ("S2", "HN"): 30, ("S1", "BD"): 2}}, "2026-10-19 05:00")
    _load({"2026-10-19": {("S1", "HN"): 1}}, "2026-10-20 05:00", audit=False)
    af.finalize(SCOPE, "BIGQUERY")
    tables = ("actual_finals", "actual_final_values", "audit_log", "alerts", "job_runs", "demand_daily",
              "forecast_history_sessions", "forecast_history_values", "forecast_history_days")
    before = {t: _count(f"SELECT COUNT(*) FROM {t}") for t in tables}
    p = {"scope_key": SCOPE, "run_from": "2026-10-17", "run_to": "2026-10-18"}
    ov = client.get("/api/history/overview", params={"scope_key": SCOPE}).json()
    assert ov["finalization"] == {"finalized_days": 1, "blocked_days": 1, "last_finalized_date": "2026-10-18",
                                  "applicable": True}
    sm = client.get("/api/history/summary", params=p).json()
    assert sm["kpi"]["basis"] == "FINALIZED_ONLY" and sm["kpi"]["actual_total"] == 37
    assert sm["kpi"]["provisional"]["actual_total"] == 1
    d18 = next(d for d in sm["daily"] if d["target_date"] == "2026-10-18")
    assert d18["actual_final"]["d1_mode"] == "FALLBACK_TIME_VERIFIED" and d18["actual_reason"] == "FINALIZED_D1"
    rows = client.get("/api/history/rows", params={**p, "limit": 1000}).json()["rows"]
    assert {r["actual_source"] for r in rows if r["target_date"] == "2026-10-18"} == {"FINAL"}
    x = client.get("/api/history/export.xlsx", params={**p, "lang": "vi"})
    info = {r[0]: r[1] for r in load_workbook(io.BytesIO(x.content))["Thông tin"].iter_rows(min_row=2, values_only=True)}
    assert info["Phiên bản định dạng file"] == "H4" and info["Tổng thực tế đã hoàn tất (chỉ dòng đã chốt)"] == 37
    assert info["Số dòng có thực tế đã hoàn tất"] == 3 and info["Số dòng thực tế tạm tính"] == 3
    assert info["Tổng thực tế tạm tính (riêng, không phải kết quả mô hình)"] == 1
    xe = client.get("/api/history/export.xlsx", params={**p, "lang": "en"})
    info_en = {r[0]: r[1] for r in load_workbook(io.BytesIO(xe.content))["Info"].iter_rows(min_row=2,
                                                                                                  values_only=True)}
    assert info_en["Complete actual total (finalized rows only)"] == 37
    assert {t: _count(f"SELECT COUNT(*) FROM {t}") for t in tables} == before          # GET never writes


# ================================================================== scope, close, legacy database
def test_demo_and_excel_scopes_are_never_finalized(frozen):
    with db.read() as c:
        assert af.plan(c, "DEMO-X", "DEMO") == {"applicable": False, "reason": "NOT_A_D1_SOURCE", "scope_key": "DEMO-X",
                                                "items": []}
    assert af.finalize("DEMO-X", "MANUAL_IMPORT")["applicable"] is False


def test_legacy_rows_get_no_status_from_the_migration_and_the_dry_run_measures_the_impact(tmp_path, frozen):
    src = tmp_path / "v8.db"
    c = sqlite3.connect(src)
    c.execute("CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY, name TEXT NOT NULL, applied_at TEXT NOT NULL)")
    for v, p in db.available_migrations():
        if v > 8:
            continue
        for stmt in db._split_sql(p.read_text(encoding="utf-8")):
            c.execute(stmt)
        c.execute("INSERT INTO schema_migrations VALUES (?,?,?)", (v, p.name, "2026-10-20T00:00:00+00:00"))
    sync_at = "2026-10-17T01:30:00+00:00"                                 # a verified full reload of 15/10 - 16/10
    c.executemany("INSERT INTO demand_daily(scope_key, data_source, order_date, sku_id, fc_id, qty, n_orders, loaded_at) "
                  "VALUES (?, 'BIGQUERY', ?,?,?,?,1,?)",
                  [(SCOPE, d, s, "HN", q, sync_at) for d in ("2026-10-15", "2026-10-16") for s, q in (("S1", 2), ("S2", 3))]
                  + [(SCOPE, "2026-10-14", "S1", "HN", 1, "2026-10-14T20:00:00+00:00")])     # written by an unknown load
    c.execute("INSERT INTO audit_log(at, action, entity_type, entity_id, details_json) VALUES (?,?,?,?,?)",
              (sync_at, "D1_SYNC", "scope", SCOPE, json.dumps({"synced_at": sync_at, "window_start": "2026-10-15",
                                                              "window_end": "2026-10-16", "mode": "FALLBACK_TIME_VERIFIED",
                                                              "backfill": False, "load_id": "D1L-1", "watermark": "2026-10-16"})))
    c.commit()
    before = c.execute("SELECT * FROM demand_daily ORDER BY rowid").fetchall()
    c.close()
    dr = af.dry_run(src)                                                  # works before 0009: measures the impact
    sc = dr["scopes"][0]
    assert (sc["to_finalize"], sc["to_block"], sc["blocked_by_reason"]) == (2, 1, {"NO_VERIFIED_D1_LOAD": 1})
    assert {i["business_date"]: i["d1_load_id"] for i in sc["items"]} == {"2026-10-14": None, "2026-10-15": "D1L-1",
                                                                          "2026-10-16": "D1L-1"}
    res = db.migrate(src)
    assert res["applied"] == ["0009_actual_finalization.sql"] and res["to"] == 9
    c = sqlite3.connect(src)
    assert c.execute("SELECT * FROM demand_daily ORDER BY rowid").fetchall() == before    # nothing changed
    assert [c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in FINAL_TABLES] == [0, 0]   # no status invented
    assert c.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    reverse_0009(c)                                                       # the header's reverse statements (copy only)
    c.commit()
    assert not c.execute("SELECT 1 FROM sqlite_master WHERE name IN ('actual_finals', 'actual_final_values') "
                         "OR name LIKE 'trg_af%'").fetchall()
    c.close()
    assert db.migrate(src)["applied"] == ["0009_actual_finalization.sql"]
    assert db.migrate(src)["applied"] == []


def test_the_23_59_close_keeps_its_own_rule_and_the_d1_hook_never_fails_the_load(frozen, monkeypatch):
    from app import actual_close as ac
    fh.capture(_run("2026-10-17 09:00", origin="2026-10-17", keys=KEYS))
    _snapshot("2026-10-18", {("S1", "HN"): 3}, "2026-10-18 18:00")
    _at("2026-10-18 23:59")
    rev1 = ac.close_day(SCOPE, "BIGQUERY", "2026-10-18")
    _at("2026-10-19 06:00")
    _load({"2026-10-18": {("S1", "HN"): 4}}, "2026-10-19 05:00")
    with monkeypatch.context() as m:                                      # the hook swallows any failure
        m.setattr(af, "finalize", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
        assert af.finalize_quietly(SCOPE, "BIGQUERY", "test") is None
    assert _count("SELECT COUNT(*) FROM alerts WHERE code='ACTUAL_FINALIZE_FAILED'") == 1
    af.finalize_quietly(SCOPE, "BIGQUERY", "test")
    out = ac.reconcile(SCOPE, "BIGQUERY")                                 # the close still records the late data
    assert [(r["revision"], r["reason"], r["status"]) for r in out["revised"]] == [(2, "D1_LOADED", "CLOSED_UNVERIFIED")]
    k, daily, _rows = _q(run_from="2026-10-17", run_to="2026-10-17")
    assert (daily[0]["actual_status"], daily[0]["actual"], daily[0]["actual_close"]["current"]) == ("COMPLETE", 4, None)
    assert _count("SELECT COUNT(*) FROM actual_closes WHERE id=?", (rev1["id"],)) == 1


def test_an_unresolved_multi_item_day_still_has_its_quantity_finalized_and_still_blocks_training(frozen):
    """The MULTI_ITEM rule is the owner's decision: the open hours of such a day stay unknown and keep blocking the
    training (OPEN_HOURS_INVALID), but the actual QUANTITY of that day is finalized like any other (the block never
    spreads to unrelated work)."""
    from app.data import open_hours
    _load({"2026-10-18": {("S1", "HN"): 2, ("S2", "HN"): 3}}, "2026-10-19 05:00")
    with db.tx() as c:
        c.execute("UPDATE demand_daily SET open_hours = NULL, open_hours_status = 'MULTI_ITEM' WHERE sku_id = 'S1'")
        c.execute("UPDATE demand_daily SET open_hours = 24, open_hours_status = 'OK' WHERE sku_id = 'S2'")
    out = af.finalize(SCOPE, "BIGQUERY")
    assert [x["business_date"] for x in out["finalized"]] == ["2026-10-18"]
    assert _final("2026-10-18")[1] == {("S1", "HN"): 2.0, ("S2", "HN"): 3.0}
    with db.read() as c:
        chk = open_hours.window_check(c, SCOPE, "2026-10-18", "2026-10-18")
        assert c.execute("SELECT open_hours_status FROM demand_daily WHERE sku_id='S1'").fetchone()[0] == "MULTI_ITEM"
    assert not chk["ok"] and set(chk["not_ok"]) == {"MULTI_ITEM"}       # the training gate is unchanged
