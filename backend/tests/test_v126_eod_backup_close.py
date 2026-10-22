"""1.2.6: end-of-day forecast backup (next 7 days), backup for run dates without an official session, actual close at
23:59 with late-data revisions, WAPE unchanged, day-end scheduler slots, migration 0005 on a v4 database.

Kind of evidence: integration on an isolated database (conftest: temporary DF_DATA_DIR + process-wide fallback; the
first fixture asserts the database is not the live one). Published runs are synthetic rows inserted directly (exact
times, as in test_v125); time is frozen per step."""
import hashlib
import io
import json
import sqlite3
import threading
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import actual_close as ac, data_checks, db, forecast_backup as fb, forecast_history as fh, paths, timeutil
from app import settings_service as ss
from app.jobs import runner
from app.jobs.scheduler import Scheduler, plan
from app.ml import metrics as mt
from app.sync import Wait

import test_v125_forecast_history as t125
from test_v125_forecast_history import _utc

SCOPE = "BIGQUERY:test-project-1.d.t"          # = bigquery_scope_key of the settings made by _bq_active()
LIVE_DATA = Path(r"D:\Buymed\Lightgbm AI\data")


def _run(*a, **k):
    return t125._run(*a, scope=SCOPE, **k)


def _snapshot(*a, **k):
    return t125._snapshot(*a, scope=SCOPE, **k)


def _demand(*a, **k):
    return t125._demand(*a, scope=SCOPE, **k)


def _d1_audit(*a, **k):
    return t125._d1_audit(*a, scope=SCOPE, **k)
KEYS = (("S1", "HN"), ("S2", "HN"), ("S1", "BD"))
VN = timeutil.VN_TZ


def _at(vn: str):
    t = datetime.fromisoformat(vn).replace(tzinfo=VN)
    timeutil.set_now_override(lambda: t.astimezone(timezone.utc))


@pytest.fixture(autouse=True)
def _isolated():
    p = paths.db_path().resolve()
    assert not p.is_relative_to(LIVE_DATA.resolve()) and ("pytest-" in str(p) or "df-tests-" in str(p)), p
    _at("2026-10-20 10:00")
    yield
    timeutil.set_now_override(None)


def _bq_active():
    ss.update("bigquery", {"project_id": "test-project-1", "dataset_id": "d", "source_table": "t"})
    ss.update("general", {"active_source": "BIGQUERY"})


def _day_end_on():
    """Both day-end flows are OFF by default (owner decision 18/10/2026, test_v126_retention_defaults.py); the tests
    of their scheduled behaviour switch them on first, as the operator does after the post-deployment check."""
    ss.update("schedule", {"eod_enabled": True, "close_enabled": True})


def _fp() -> dict:
    """Row count + sha256 of every forecast table (published values, D+1 history, official pointers, backups)."""
    out = {}
    with db.read() as c:
        for t, cols in (("forecast_runs", "id, created_at, n_series"),
                        ("forecast_values", "run_id, sku_id, fc_id, target_date, horizon_day, yhat, method"),
                        ("forecast_history_values", "forecast_run_id, sku_id, fc_id, target_date, yhat, method"),
                        ("forecast_history_days", "scope_key, run_date, forecast_run_id"),
                        ("forecast_eod_snapshot_values", "snapshot_id, target_date, sku_id, fc_id, yhat, method")):
            lines = sorted("|".join(map(repr, r)) for r in c.execute(f"SELECT {cols} FROM {t}"))
            out[t] = (len(lines), hashlib.sha256("\n".join(lines).encode()).hexdigest())
    return out


def _jobs() -> dict:
    with db.read() as c:
        return {r[0]: r[1] for r in c.execute("SELECT job_type, COUNT(*) FROM job_runs GROUP BY job_type")} | {
            "training_runs": c.execute("SELECT COUNT(*) FROM training_runs").fetchone()[0]}


def _one(sql, p=()):
    with db.read() as c:
        return db.one(c, sql, p)


def _kpi(**kw):
    with db.read() as c:
        qy = fh.Query(c, kw.pop("scope", SCOPE), **kw)
        daily = qy.daily()
        return qy.kpi(daily), daily, list(qy.rows(limit=None))


# ------------------------------------------------------------------ A. WAPE after each D0 update, formula unchanged
def test_wape_is_recomputed_after_each_d0_update_with_the_unchanged_formula():
    rid = _run("2026-10-18 19:00", origin="2026-10-18", horizon=15, keys=KEYS)
    fh.capture(rid)                                         # D+1 = 19/10: 10.5, 11.5, 12.5 -> shown 11, 12, 13 (1.2.9)
    f0, j0 = _fp(), _jobs()
    seen = []
    for at, vals in (("2026-10-19 08:05", {}), ("2026-10-19 10:05", {("S1", "HN"): 3, ("S1", "BD"): 2}),
                     ("2026-10-19 12:05", {("S1", "HN"): 6, ("S2", "HN"): 1, ("S1", "BD"): 4})):
        _at(at)
        _snapshot("2026-10-19", vals, at[:-1] + "0")
        k, _d, rows = _kpi(run_from="2026-10-18", run_to="2026-10-18")
        f = [r["forecast_qty"] for r in rows]                              # 1.2.9: the whole number shown
        a = [r["actual"] for r in rows]
        # 1.2.9 (seventh batch): intraday actual is PROVISIONAL - reported in its own group, never as the finalized KPI
        kp = k["provisional"]
        assert k["status"] == fh.NO_ACTUAL and k["actual_total"] is None and kp["actual_total"] == sum(vals.values())
        exp = mt.compute(a, f)["wape"] if sum(a) > 0 else None             # ml/metrics.py, the same formula
        assert kp["wape"] == exp and (exp is None or abs(exp - sum(abs(x - y) for x, y in zip(f, a)) / sum(a)) < 1e-12)
        assert k["wape"] is None
        seen.append(kp["wape"])
        assert _fp() == f0 and _jobs() == j0                                # actual moved, forecast and jobs did not
    assert seen[0] is None and seen[1] == pytest.approx((8 + 12 + 11) / 5) and seen[2] == pytest.approx((5 + 11 + 9) / 11)


def test_a_d0_job_re_entered_with_the_same_job_id_writes_one_snapshot(small_demo):
    from app import sync
    timeutil.set_now_override(None)
    r1 = sync.run_d0_snapshot("JOB-SAME")
    r2 = sync.run_d0_snapshot("JOB-SAME")
    assert r2["already_written"] is True and r2["snapshot_id"] == r1["snapshot_id"]
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM d0_snapshots WHERE job_id='JOB-SAME'").fetchone()[0] == 1


# ------------------------------------------------------------------ B. end-of-day snapshot
def test_eod_snapshot_copies_the_next_7_days_of_the_newest_official_session():
    old = _run("2026-10-16 09:00", origin="2026-10-16", horizon=15, keys=KEYS, base=5)
    morning = _run("2026-10-18 09:00", origin="2026-10-18", horizon=15, keys=KEYS, base=20)
    evening = _run("2026-10-18 19:00", origin="2026-10-18", horizon=15, keys=KEYS, base=30)
    later_day = _run("2026-10-19 09:00", origin="2026-10-19", horizon=15, keys=KEYS, base=90)
    for r in (old, morning, evening, later_day):
        fh.capture(r)
    f0, j0 = _fp(), _jobs()
    _at("2026-10-18 23:55")
    s = fb.create(SCOPE, "BIGQUERY", "2026-10-18", "JOB-X")
    assert (s["status"], s["source_run_id"], s["source_run_date"]) == ("COMPLETE", evening, "2026-10-18")
    assert (s["first_target_date"], s["last_target_date"], s["n_days"], s["days_covered"]) == ("2026-10-19", "2026-10-25", 7, 7)
    assert s["n_rows"] == 7 * 3 and s["source_n_series"] == 3 and s["timezone"] == "Asia/Ho_Chi_Minh"
    assert (s["source_model_version"], s["job_id"], s["missing_dates_json"]) == (_one(
        "SELECT model_version FROM forecast_runs WHERE id=?", (evening,))["model_version"], "JOB-X", None)
    with db.read() as c:
        got = c.execute("SELECT target_date, sku_id, fc_id, horizon_day, yhat, method FROM forecast_eod_snapshot_values "
                        "WHERE snapshot_id=? ORDER BY 1, 2, 3", (s["id"],)).fetchall()
        src = c.execute("SELECT target_date, sku_id, fc_id, horizon_day, yhat, method FROM forecast_values WHERE run_id=? "
                        "AND target_date BETWEEN '2026-10-19' AND '2026-10-25' ORDER BY 1, 2, 3", (evening,)).fetchall()
    assert got == src and {g[0] for g in got} == {f"2026-10-{d}" for d in range(19, 26)}   # grain SKU x FC x date
    assert s["values_sha256"] == fb.fingerprint(src) and s["total_yhat"] == pytest.approx(sum(x[4] for x in src))
    assert got[0][3] == 2                                                   # D+1 of a run of 18/10 = its horizon 2
    f1 = _fp()
    assert {k: f1[k] for k in f1 if k != "forecast_eod_snapshot_values"} == {k: f0[k] for k in f0 if k != "forecast_eod_snapshot_values"}
    assert _jobs() == j0                                                    # no training / forecast / job started


def test_rerun_and_concurrent_runs_never_duplicate_or_change_a_complete_snapshot():
    fh.capture(_run("2026-10-17 09:00", origin="2026-10-17", horizon=15, keys=KEYS))
    _at("2026-10-18 23:55")
    first = fb.create(SCOPE, "BIGQUERY", "2026-10-18")
    f0 = _fp()
    again = fb.create(SCOPE, "BIGQUERY", "2026-10-18")
    assert again["already_complete"] is True and again["id"] == first["id"] and again["attempts"] == 1
    assert _fp() == f0
    res, errs = [], []

    def go():
        try:
            res.append(fb.create(SCOPE, "BIGQUERY", "2026-10-17"))
        except Exception as e:  # noqa: BLE001
            errs.append(e)
    ts = [threading.Thread(target=go) for _ in range(4)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert not errs and len({r["id"] for r in res}) == 1 and sum(not r["already_complete"] for r in res) == 1
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM forecast_eod_snapshots").fetchone()[0] == 2
    with pytest.raises(sqlite3.IntegrityError):                            # a COMPLETE snapshot never changes
        with db.tx() as c:
            c.execute("UPDATE forecast_eod_snapshots SET total_yhat=0 WHERE id=?", (first["id"],))
    for stmt in ("UPDATE forecast_eod_snapshot_values SET yhat=0", "DELETE FROM forecast_eod_snapshot_values",
                 "DELETE FROM forecast_eod_snapshots"):
        with pytest.raises(sqlite3.IntegrityError):
            with db.tx() as c:
                c.execute(stmt)
    assert _fp()["forecast_eod_snapshot_values"][0] == 2 * 21


def test_without_a_valid_forecast_no_snapshot_is_faked(client):
    _bq_active()
    _day_end_on()
    _at("2026-10-18 23:56")
    j0 = _jobs()
    r = runner.execute("EOD_FORECAST_SNAPSHOT", "EOD_FORECAST_SNAPSHOT:2026-10-18", "SCHEDULE", None,
                       {"business_date": "2026-10-18"})
    assert (r["status"], r["reason_code"]) == ("SKIPPED", "EOD_NO_VALID_FORECAST")
    s = _one("SELECT * FROM forecast_eod_snapshots")
    assert (s["status"], s["n_rows"], s["source_run_id"], s["total_yhat"]) == ("NO_SOURCE", 0, None, None)
    assert _fp()["forecast_eod_snapshot_values"][0] == 0 and _fp()["forecast_values"][0] == 0
    assert _one("SELECT level, code FROM alerts WHERE code='EOD_NO_VALID_FORECAST'")["level"] == "WARNING"
    j1 = _jobs()
    assert j1.pop("EOD_FORECAST_SNAPSHOT") == 1 and j1 == j0                # nothing else was started
    with db.read() as c:
        assert fh.fallback_days(c, SCOPE, timeutil.today_vn()) == []


def test_a_short_source_is_incomplete_never_used_and_completed_later_by_a_valid_session():
    fh.capture(_run("2026-10-16 09:00", origin="2026-10-16", horizon=7, keys=KEYS))   # covers 16..22 only
    _at("2026-10-18 23:55")
    s = fb.create(SCOPE, "BIGQUERY", "2026-10-18")
    assert (s["status"], s["status_reason"], s["days_covered"], s["n_rows"]) == ("INCOMPLETE", "SOURCE_COVERS_4_OF_7_DAYS", 4, 0)
    assert json.loads(s["missing_dates_json"]) == ["2026-10-23", "2026-10-24", "2026-10-25"]
    assert _one("SELECT code FROM alerts WHERE code='EOD_SNAPSHOT_INCOMPLETE'")
    with pytest.raises(sqlite3.IntegrityError):                            # an incomplete snapshot can hold no value
        with db.tx() as c:
            c.execute("INSERT INTO forecast_eod_snapshot_values VALUES (?, '2026-10-19', 'S1', 'HN', 1, 1.0, 'X')", (s["id"],))
    _at("2026-10-19 10:00")
    with db.read() as c:
        assert all(d["snapshot_id"] != s["id"] for d in fh.fallback_days(c, SCOPE))      # never used as backup
    # the session of 18/10 whose recording had failed is recorded later: the next attempt completes the same record
    late = _run("2026-10-18 20:00", origin="2026-10-18", horizon=15, keys=KEYS)
    fh.reconcile()
    s2 = fb.create(SCOPE, "BIGQUERY", "2026-10-18")
    assert (s2["id"], s2["status"], s2["source_run_id"], s2["attempts"], s2["n_rows"]) == (s["id"], "COMPLETE", late, 2, 21)
    s3 = fb.create(SCOPE, "BIGQUERY", "2026-10-18")
    assert s3["already_complete"] and s3["attempts"] == 2


def test_the_job_waits_for_a_running_forecast_job_and_keeps_the_slot_business_date():
    _bq_active()
    _day_end_on()
    fh.capture(_run("2026-10-18 09:00", origin="2026-10-18", horizon=15, keys=KEYS))
    with db.tx() as c:
        c.execute("INSERT INTO job_runs(id, job_type, idempotency_key, trigger, status, created_at, started_at) "
                  "VALUES ('JOB-TR', 'TRAIN_PUBLISH', 'k-tr', 'SCHEDULE', 'RUNNING', ?, ?)",
                  (_utc("2026-10-18 23:40"), _utc("2026-10-18 23:40")))
    _at("2026-10-18 23:55")
    with pytest.raises(Wait):
        fb.run_job("JOB-E", {"business_date": "2026-10-18"})
    r = runner.execute("EOD_FORECAST_SNAPSHOT", "EOD_FORECAST_SNAPSHOT:2026-10-18", "SCHEDULE", None,
                       {"business_date": "2026-10-18"})
    assert (r["status"], r["reason_code"], r["attempt"]) == ("RETRY_WAIT", "WAITING_FOR_FORECAST_JOB", 0)
    # 19/10 06:30 Vietnam = 18/10 23:30 UTC: past the waiting limit, the snapshot is still the one of 18/10
    _at("2026-10-19 06:30")
    out = fb.run_job("JOB-E", {"business_date": "2026-10-18"})
    assert (out["business_date"], out["snapshot_status"], out["first_target_date"]) == ("2026-10-18", "COMPLETE", "2026-10-19")
    assert out["forecast_job_still_running"] == "JOB-TR"
    # a manual run before today's slot snapshots yesterday, never a day that has not ended
    _at("2026-10-19 15:00")
    assert fb.latest_due_day(ss.get("schedule").eod_time) == "2026-10-18"
    _at("2026-10-19 23:56")
    assert fb.latest_due_day(ss.get("schedule").eod_time) == "2026-10-19"


# ------------------------------------------------------------------ C. backup for days without an official session
def _two_sources_and_snapshots():
    p1 = _run("2026-10-15 09:00", origin="2026-10-15", horizon=15, keys=KEYS, base=10)
    p2 = _run("2026-10-16 09:00", origin="2026-10-16", horizon=15, keys=KEYS[:2], base=50)    # no S1|BD
    for r in (p1, p2):
        fh.capture(r)
    snaps = {}
    for d in ("2026-10-15", "2026-10-16", "2026-10-17", "2026-10-18"):
        _at(f"{d} 23:55")
        snaps[d] = fb.create(SCOPE, "BIGQUERY", d)
        assert snaps[d]["status"] == "COMPLETE"
    return p1, p2, snaps


def test_backup_fills_only_days_without_a_session_with_the_newest_covering_snapshot(client):
    p1, p2, snaps = _two_sources_and_snapshots()
    _at("2026-10-20 10:00")
    with db.read() as c:
        fbd = {d["run_date"]: d for d in fh.fallback_days(c, SCOPE)}
    assert sorted(fbd) == ["2026-10-17", "2026-10-18", "2026-10-19", "2026-10-20"]       # 15/10, 16/10 have sessions
    assert fbd["2026-10-17"]["snapshot_id"] == snaps["2026-10-17"]["id"]                   # newest, not the 15/10 one
    assert fbd["2026-10-20"]["snapshot_id"] == snaps["2026-10-18"]["id"] and fbd["2026-10-20"]["target_date"] == "2026-10-21"
    k, daily, rows = _kpi(run_from="2026-10-15", run_to="2026-10-20")
    src = {d["run_date"]: d["forecast_source"] for d in daily}
    assert src == {"2026-10-15": "PRIMARY", "2026-10-16": "PRIMARY", "2026-10-17": "SNAPSHOT", "2026-10-18": "SNAPSHOT",
                   "2026-10-19": "SNAPSHOT", "2026-10-20": "SNAPSHOT"}
    assert (k["days_primary"], k["days_backup"], k["days_with_history"]) == (2, 4, 6)
    r17 = sorted((r["sku_id"], r["fc_id"], r["yhat"], r["forecast_run_id"], r["snapshot_id"]) for r in rows
                 if r["run_date"] == "2026-10-17")
    # values of the 17/10 snapshot (source p2, D+1 = 18/10 = horizon 3): 51, 52; S1|BD absent -> no row, never 0
    assert r17 == [("S1", "HN", 51.0, p2, snaps["2026-10-17"]["id"]), ("S2", "HN", 52.0, p2, snaps["2026-10-17"]["id"])]
    assert all(r["yhat"] > 0 for r in rows) and {r["run_date"]: r["snapshot_business_date"] for r in rows}["2026-10-19"] == "2026-10-18"
    # a snapshot no longer covering a day leaves it empty: 25/10 and 26/10 have no forecast at all
    _at("2026-10-27 10:00")
    k, daily, rows = _kpi(run_from="2026-10-23", run_to="2026-10-26")
    assert [d["forecast_source"] for d in daily] == ["SNAPSHOT", "SNAPSHOT", None, None]
    assert [d["has_history"] for d in daily] == [True, True, False, False]
    assert [d["forecast"] for d in daily][2:] == [None, None] and {r["run_date"] for r in rows} == {"2026-10-23", "2026-10-24"}


def test_a_new_official_session_takes_over_and_reading_writes_nothing(client):
    p1, p2, snaps = _two_sources_and_snapshots()
    _at("2026-10-20 10:00")
    _demand("2026-10-18", {("S1", "HN"): 40, ("S2", "HN"): 60})
    _d1_audit("2026-10-18", "FALLBACK_TIME_VERIFIED")
    tables = ("job_runs", "training_runs", "forecast_runs", "forecast_values", "forecast_history_sessions",
              "forecast_history_values", "forecast_history_days", "forecast_eod_snapshots", "forecast_eod_snapshot_values",
              "actual_closes", "actual_close_values", "audit_log", "alerts", "app_settings")
    with db.read() as c:
        before = {t: c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in tables}
    p = {"scope_key": SCOPE, "run_from": "2026-10-15", "run_to": "2026-10-20"}
    sm = client.get("/api/history/summary", params=p).json()
    assert [d["run_date"] for d in sm["backup_days"]] == ["2026-10-20", "2026-10-19", "2026-10-18", "2026-10-17"]
    rows = client.get("/api/history/rows", params={**p, "limit": 1000}).json()
    assert {r["forecast_source"] for r in rows["rows"]} == {"PRIMARY", "SNAPSHOT"}
    only = client.get("/api/history/rows", params={**p, "limit": 1000, "forecast_source": "PRIMARY"}).json()
    assert {r["run_date"] for r in only["rows"]} == {"2026-10-15", "2026-10-16"}
    bk = client.get("/api/history/summary", params={**p, "forecast_source": "SNAPSHOT"}).json()
    # 1.2.9 (seventh batch): those days are not finalized -> the provisional group of the KPI
    assert bk["kpi"]["n_rows"] == 4 * 2 and bk["kpi"]["provisional"]["actual_total"] == 100   # 17/10 -> 18/10: 51 vs 40, 52 vs 60
    assert bk["kpi"]["provisional"]["wape"] == pytest.approx((11 + 8) / 100) and bk["kpi"]["wape"] is None
    assert client.get("/api/history/summary", params={**p, "forecast_source": "OTHER"}).status_code == 422
    x = client.get("/api/history/export.xlsx", params={**p, "lang": "vi"})
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Thông tin", "Phiên forecast", "Theo ngày", "Chi tiết SKU x FC"]
    det = list(wb["Chi tiết SKU x FC"].iter_rows(values_only=True))
    assert det[0][15:] == ("Nguồn forecast", "Mã snapshot dự phòng", "Ngày nghiệp vụ của snapshot", "Snapshot tạo lúc")
    assert {r[15] for r in det[1:]} == {"Forecast chính", "Snapshot dự phòng"}
    assert all(r[16] and r[17] for r in det[1:] if r[15] == "Snapshot dự phòng") and all(r[16] is None for r in det[1:] if r[15] == "Forecast chính")
    info = {r[0]: r[1] for r in wb["Thông tin"].iter_rows(min_row=2, values_only=True)}
    assert info["Phiên bản định dạng file"] == "H4" and info["Số ngày dùng snapshot dự phòng"] == 4     # 1.2.9: H4
    with db.read() as c:
        assert {t: c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in tables} == before     # GET writes nothing
    f0 = _fp()["forecast_eod_snapshot_values"]
    # a valid session of 17/10 appears (published late): the official forecast wins again, the snapshot is untouched
    new = _run("2026-10-17 22:00", origin="2026-10-17", horizon=15, keys=KEYS, base=70)
    fh.capture(new)
    k, daily, rows = _kpi(run_from="2026-10-17", run_to="2026-10-17")
    assert daily[0]["forecast_source"] == "PRIMARY" and {r["forecast_run_id"] for r in rows} == {new} and len(rows) == 3
    assert _fp()["forecast_eod_snapshot_values"] == f0


# ------------------------------------------------------------------ D. actual close at 23:59
def test_the_close_uses_the_vietnam_business_date_and_refuses_an_early_close():
    _bq_active()
    fh.capture(_run("2026-10-17 19:00", origin="2026-10-17", horizon=15, keys=KEYS))
    _snapshot("2026-10-18", {("S1", "HN"): 1}, "2026-10-18 10:00")
    s18 = _snapshot("2026-10-18", {("S1", "HN"): 3, ("S2", "HN"): 2}, "2026-10-18 18:00")
    _at("2026-10-18 20:00")
    with pytest.raises(Exception) as e:
        ac.run_job(None, {"business_date": "2026-10-18"})
    assert getattr(e.value, "code", None) == "ACTUAL_CLOSE_TOO_EARLY"
    # 19/10 03:00 Vietnam = 18/10 20:00 UTC: a manual close closes the Vietnam day 18/10 (UTC would say 17/10)
    _at("2026-10-19 03:00")
    assert ac.latest_due_day(ss.get("schedule").close_time) == "2026-10-18"
    out = ac.run_job(None, {})
    assert (out["business_date"], out["source"], out["source_ref"], out["rows"], out["total_qty"]) == (
        "2026-10-18", "D0", s18, 2, 5.0)
    assert out["cutoff_at"] == "2026-10-18T16:59:00+00:00" and out["close_status"] == "CLOSED_UNVERIFIED"
    c1 = _one("SELECT * FROM actual_closes WHERE business_date='2026-10-18'")
    assert (c1["revision"], c1["kind"], c1["status_reason"], c1["timezone"]) == (1, "CLOSE", "D0_NO_COMPLETION_SIGNAL",
                                                                                 "Asia/Ho_Chi_Minh")


def test_close_retries_and_concurrent_closes_write_one_revision():
    _bq_active()
    _day_end_on()
    _snapshot("2026-10-18", {("S1", "HN"): 3}, "2026-10-18 18:00")
    _at("2026-10-18 23:59")
    r1 = runner.execute("ACTUAL_CLOSE", "ACTUAL_CLOSE:2026-10-18", "SCHEDULE", None, {"business_date": "2026-10-18"})
    r2 = runner.execute("ACTUAL_CLOSE", "ACTUAL_CLOSE:2026-10-18", "RETRY")
    assert r1["status"] == "SUCCEEDED" and r2["note"] == "ALREADY_SUCCEEDED"
    res, errs = [], []

    def go():
        try:
            res.append(ac.close_day(SCOPE, "BIGQUERY", "2026-10-18"))
        except Exception as e:  # noqa: BLE001
            errs.append(e)
    ts = [threading.Thread(target=go) for _ in range(4)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert not errs and all(r["already_closed"] for r in res)
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM actual_closes").fetchone()[0] == 1
    with pytest.raises(sqlite3.IntegrityError):              # a close is never edited: changes are new revisions
        with db.tx() as c:
            c.execute("UPDATE actual_closes SET total_qty=99")
    with pytest.raises(sqlite3.IntegrityError):              # and a revision must follow the previous one
        with db.tx() as c:
            c.execute("INSERT INTO actual_closes(id, scope_key, data_source, business_date, revision, kind, reason, n_rows, "
                      "status, cutoff_at, closed_at, timezone) VALUES ('X', ?, 'BIGQUERY', '2026-10-18', 3, "
                      "'RECONCILIATION', 'X', 0, 'CLOSED_NO_DATA', 'x', 'x', 'x')", (SCOPE,))


def test_a_missing_actual_is_never_zero_and_without_a_signal_a_close_is_never_complete():
    rid = _run("2026-10-17 19:00", origin="2026-10-17", horizon=15, keys=KEYS)
    rid2 = _run("2026-10-18 19:00", origin="2026-10-18", horizon=15, keys=KEYS)
    fh.capture(rid)
    fh.capture(rid2)
    _at("2026-10-18 23:59")
    none = ac.close_day(SCOPE, "BIGQUERY", "2026-10-18")                  # no D0 read, no D-1 load that day
    assert (none["status"], none["n_rows"], none["total_qty"], none["source"]) == ("CLOSED_NO_DATA", 0, None, None)
    assert _one("SELECT level FROM alerts WHERE code='ACTUAL_CLOSE_NO_DATA'")["level"] == "WARNING"
    _snapshot("2026-10-19", {("S1", "HN"): 4, ("S2", "HN"): 0}, "2026-10-19 18:00")   # S1|BD absent from the read
    _at("2026-10-19 23:59")
    c19 = ac.close_day(SCOPE, "BIGQUERY", "2026-10-19")
    assert (c19["status"], c19["n_rows"], c19["total_qty"]) == ("CLOSED_UNVERIFIED", 2, 4.0)
    with db.read() as c:
        assert {(r[0], r[1]) for r in c.execute("SELECT sku_id, fc_id FROM actual_close_values WHERE close_id=?",
                                                (c19["id"],))} == {("S1", "HN"), ("S2", "HN")}
    _at("2026-10-20 10:00")
    k, daily, _rows = _kpi(run_from="2026-10-17", run_to="2026-10-18")
    by = {d["target_date"]: d for d in daily}
    assert by["2026-10-18"]["actual"] is None and by["2026-10-18"]["actual_status"] == "NO_ACTUAL"
    assert by["2026-10-18"]["actual_close"]["status"] == "CLOSED_NO_DATA"
    assert (by["2026-10-19"]["actual_status"], by["2026-10-19"]["actual_reason"]) == ("CLOSED", "CLOSED_D0")
    assert by["2026-10-19"]["actual_close"]["current"] is True
    # 1.2.9 (seventh batch): a closed day that is not finalized is reported in the `closed` group, not as final
    assert k["status"] == "NO_ACTUAL" and k["actual_total"] is None
    assert k["closed"]["actual_total"] == 4.0                            # the shown figure is the source's, unchanged


def test_late_data_adds_controlled_revisions_and_never_touches_the_forecast():
    _bq_active()
    rid = _run("2026-10-17 19:00", origin="2026-10-17", horizon=15, keys=KEYS)
    fh.capture(rid)
    _snapshot("2026-10-18", {("S1", "HN"): 3, ("S2", "HN"): 2}, "2026-10-18 18:00")
    _at("2026-10-18 23:59")
    rev1 = ac.close_day(SCOPE, "BIGQUERY", "2026-10-18")
    f0, j0 = _fp(), _jobs()
    _at("2026-10-19 13:00")
    assert ac.reconcile(SCOPE, "BIGQUERY")["revised"] == []                 # nothing new: nothing written
    _demand("2026-10-18", {("S1", "HN"): 9, ("S2", "HN"): 4, ("S1", "BD"): 6}, loaded="2026-10-19T05:00:00+00:00")
    _d1_audit("2026-10-18", "FALLBACK_TIME_VERIFIED")
    out = ac.reconcile_quietly(SCOPE, "BIGQUERY", "test D-1")           # what runs after each D-1 load
    assert [(r["revision"], r["reason"], r["previous_total"], r["total_qty"], r["status"]) for r in out["revised"]] == [
        (2, "D1_LOADED", 5.0, 19.0, "CLOSED_UNVERIFIED")]
    rev2 = _one("SELECT * FROM actual_closes WHERE revision=2")
    assert (rev2["n_changed"], rev2["n_added"], rev2["n_removed"], rev2["previous_id"]) == (2, 1, 0, rev1["id"])
    assert ac.reconcile(SCOPE, "BIGQUERY")["revised"] == []                 # idempotent
    _demand("2026-10-18", {("S1", "HN"): 10, ("S2", "HN"): 4}, loaded="2026-10-20T05:00:00+00:00")   # a reload
    _at("2026-10-20 13:00")
    assert [r["reason"] for r in ac.reconcile(SCOPE, "BIGQUERY")["revised"]] == ["D1_RELOADED"]
    rev3 = _one("SELECT * FROM actual_closes WHERE revision=3")
    assert (rev3["n_changed"], rev3["n_added"], rev3["n_removed"], rev3["total_qty"]) == (1, 0, 1, 14.0)
    _d1_audit("2026-10-18", "SIGNAL")                                       # the source confirms the day
    assert [(r["reason"], r["status"]) for r in ac.reconcile(SCOPE, "BIGQUERY")["revised"]] == [
        ("COMPLETION_SIGNAL_CONFIRMED", "CLOSED_VERIFIED")]
    k, daily, _rows = _kpi(run_from="2026-10-17", run_to="2026-10-17")
    # 1.2.9 (seventh batch): the close keeps its own rule (revision 4 CLOSED_VERIFIED on the signal), but the history
    # shows COMPLETE only for a FINALIZED day - this day was not finalized (no verified D-1 load), so it stays CLOSED
    assert (daily[0]["actual_status"], daily[0]["actual_close"]["revision"], k["status"]) == ("CLOSED", 4, "NO_ACTUAL")
    assert daily[0]["actual_close"]["status"] == "CLOSED_VERIFIED"
    r1 = _one("SELECT * FROM actual_closes WHERE id=?", (rev1["id"],))
    assert (r1["total_qty"], r1["source"], r1["status"]) == (5.0, "D0", "CLOSED_UNVERIFIED")      # revision 1 kept
    assert _fp() == f0 and _jobs() == j0                                   # forecast and jobs untouched
    dc = data_checks.run(paths.db_path())
    assert {x["check"]: x["status"] for x in dc["checks"]}["actual_close_revisions"] == "OK"
    assert dc["row_counts"]["actual_closes"] == 4


# ------------------------------------------------------------------ E. scheduler
def _first_seen(at: str):
    with db.tx() as c:
        c.execute("INSERT OR REPLACE INTO sync_state(name, value_json, updated_at) VALUES ('flows_first_seen', ?, 'x')",
                  (json.dumps({"eod": at, "close": at}),))


def test_day_end_slots_carry_their_business_date_and_run_once():
    base = ss.ScheduleSettings(train_enabled=False, eod_enabled=True, close_enabled=True)
    seen = {"eod": "2026-10-17T00:00:00+07:00", "close": "2026-10-17T00:00:00+07:00"}
    with db.read() as c:
        assert plan(datetime(2026, 10, 18, 23, 56, tzinfo=VN), base, "BIGQUERY", None, c) == []     # never seen: nothing
        acts = plan(datetime(2026, 10, 18, 23, 56, tzinfo=VN), base, "BIGQUERY", None, c, seen)
        k = {(a["key"], a["trigger"], a["params"]["business_date"]) for a in acts}
        assert k == {("EOD_FORECAST_SNAPSHOT:2026-10-18", "SCHEDULE", "2026-10-18"),
                     ("EOD_FORECAST_SNAPSHOT:2026-10-17", "CATCH_UP", "2026-10-17"),     # yesterday's missed slot, once
                     ("ACTUAL_CLOSE:2026-10-17", "CATCH_UP", "2026-10-17")}
        # 00:10 on 19/10 is within the grace of 18/10 23:59: still the close of 18/10 (the slot's day, not 'today')
        acts = plan(datetime(2026, 10, 19, 0, 10, tzinfo=VN), base, "BIGQUERY", None, c,
                    {"eod": "2026-10-18T23:00:00+07:00", "close": "2026-10-18T23:00:00+07:00"})
        assert {(a["key"], a["trigger"], a["params"]["business_date"]) for a in acts} == {
            ("EOD_FORECAST_SNAPSHOT:2026-10-18", "SCHEDULE", "2026-10-18"), ("ACTUAL_CLOSE:2026-10-18", "SCHEDULE", "2026-10-18")}
        off = base.model_copy(update={"eod_enabled": False, "close_enabled": False})
        assert plan(datetime(2026, 10, 19, 0, 10, tzinfo=VN), off, "BIGQUERY", None, c, seen) == []
    # the existing D0 / D-1 / training planning is unchanged by the new flows
    with db.read() as c:
        assert plan(datetime(2026, 10, 15, 10, 5, tzinfo=VN), base, "BIGQUERY", None, c) == []


def test_a_restart_never_doubles_a_slot_and_installing_never_back_dates(small_demo):
    _day_end_on()
    t = datetime(2026, 10, 18, 23, 56, tzinfo=VN)
    timeutil.set_now_override(lambda: t.astimezone(timezone.utc))
    first = Scheduler()
    acts = first.tick(t, wait=True)                             # first tick of this version: 23:55 slot is before it
    assert not [a for a in acts if a["job_type"] in ("EOD_FORECAST_SNAPSHOT", "ACTUAL_CLOSE")]
    seen = json.loads(_one("SELECT value_json FROM sync_state WHERE name='flows_first_seen'")["value_json"])
    assert seen == {"eod": t.isoformat(), "close": t.isoformat()}
    t2 = datetime(2026, 10, 19, 0, 5, tzinfo=VN)                # the 23:59 close slot comes after the first tick
    timeutil.set_now_override(lambda: t2.astimezone(timezone.utc))
    a1 = Scheduler().tick(t2, wait=True)                        # a restarted process ...
    a2 = Scheduler().tick(t2 + timedelta(seconds=30), wait=True)   # ... and another one
    assert [a["key"] for a in a1 if a["job_type"] == "ACTUAL_CLOSE"] == ["ACTUAL_CLOSE:2026-10-18"]
    assert not [a for a in a2 if a["job_type"] == "ACTUAL_CLOSE"]
    with db.read() as c:
        rows = db.rows(c, "SELECT idempotency_key, status FROM job_runs WHERE job_type IN ('EOD_FORECAST_SNAPSHOT','ACTUAL_CLOSE')")
    assert [r["idempotency_key"] for r in rows] == ["ACTUAL_CLOSE:2026-10-18"]
    assert json.loads(_one("SELECT value_json FROM sync_state WHERE name='flows_first_seen'")["value_json"]) == seen


def test_whole_missed_days_are_recorded_and_only_yesterday_is_caught_up(small_demo):
    now = datetime(2026, 10, 20, 8, 0, tzinfo=VN)
    _first_seen("2026-10-01T00:00:00+07:00")
    sched = ss.ScheduleSettings(train_enabled=False, eod_enabled=True, close_enabled=True)
    with db.read() as c:
        acts = plan(now, sched, "DEMO", now - timedelta(days=4), c,
                    {"eod": "2026-10-01T00:00:00+07:00", "close": "2026-10-01T00:00:00+07:00"})
    runs = {a["key"] for a in acts if a["action"] == "run"}
    skips = {a["key"] for a in acts if a["action"] == "skip"}
    assert runs == {"EOD_FORECAST_SNAPSHOT:2026-10-19", "ACTUAL_CLOSE:2026-10-19"}
    assert {"EOD_FORECAST_SNAPSHOT:2026-10-17", "ACTUAL_CLOSE:2026-10-16", "ACTUAL_CLOSE:2026-10-18"} <= skips
    assert not (runs & skips)


def test_schedules_api_shows_both_new_flows_and_their_switches(client, small_demo, monkeypatch):
    from app.jobs import scheduler as sched_mod
    from test_v126_schedule_config import RunningScheduler
    ov = client.get("/api/schedules").json()
    assert {"d1", "d0", "train", "eod", "close"} <= set(ov["flows"])
    vn = "Asia/Ho_Chi_Minh"
    assert ov["flows"]["eod"]["config"] == {"time": "23:55", "timezone": vn}
    assert ov["flows"]["close"]["config"] == {"time": "23:59", "timezone": vn}
    assert ov["flows"]["eod"]["readiness"]["ready"] and ov["flows"]["close"]["enabled"] is False    # default OFF
    # 1.2.6: switching a day-end flow ON needs a running scheduler that confirms it (this test app runs none)
    r = client.put("/api/schedules/close", json={"enabled": True})
    assert r.status_code == 409 and r.json()["code"] == "SCHEDULE_NOT_REGISTERED"
    assert ss.get("schedule").close_enabled is False
    with monkeypatch.context() as m:          # never monkeypatch.undo(): conftest's isolation must stay in place
        m.setattr(sched_mod, "SCHEDULER", RunningScheduler())
        r = client.put("/api/schedules/close", json={"enabled": True})
        assert r.status_code == 200 and ss.get("schedule").close_enabled is True
        r = client.put("/api/schedules/close", json={"enabled": False})
        assert r.status_code == 200 and ss.get("schedule").close_enabled is False
    r = client.put("/api/settings/schedule", json={**ss.get("schedule").model_dump(), "close_enabled": True})
    assert r.status_code == 409 and r.json()["code"] == "SCHEDULE_SWITCH_VIA_FLOW_API"
    assert client.put("/api/schedules/close", json={"time": "24:00"}).status_code == 422
    # the backup and the close stand for the END of the day: an earlier time is refused, nothing is saved
    for flow in ("eod", "close"):
        assert client.put(f"/api/schedules/{flow}", json={"time": "08:00"}).status_code == 422
    assert (ss.get("schedule").eod_time, ss.get("schedule").close_time) == ("23:55", "23:59")
    assert client.put("/api/schedules/eod", json={"time": "23:30"}).json()["config"] == {"time": "23:30", "timezone": vn}
    ss.update("general", {"active_source": "MANUAL_IMPORT"})
    assert client.get("/api/schedules/close").json()["readiness"]["reasons"][0]["code"] == "CLOSE_SOURCE_MANUAL_IMPORT"


# ------------------------------------------------------------------ F. migration 0005 on a v4 database
def test_migration_0005_on_a_v4_database_keeps_every_row(tmp_path, monkeypatch):
    target = tmp_path / "v4.db"
    conn = db.connect(target, create=True)
    db.current_version(conn)
    for v, p in db.available_migrations():
        if v > 4:
            break
        for stmt in db._split_sql(p.read_text(encoding="utf-8")):
            conn.execute(stmt)
        conn.execute("INSERT INTO schema_migrations(version, name, applied_at) VALUES (?,?,?)", (v, p.name, "x"))
    conn.close()
    with monkeypatch.context() as m:          # never monkeypatch.undo(): conftest's isolation must stay in place
        m.setattr(paths, "db_path", lambda: target)
        fh.capture(_run("2026-10-18 09:00", origin="2026-10-18", horizon=15, keys=KEYS))
        tabs = ("training_runs", "forecast_runs", "forecast_values", "forecast_history_sessions", "forecast_history_values",
                "forecast_history_days", "job_runs", "audit_log")
        with db.read() as c:
            before = {t: c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in tabs}
            vals0 = c.execute("SELECT COUNT(*), SUM(yhat) FROM forecast_values").fetchone()
        res = db.migrate(target)
        assert (res["from"], res["to"], res["applied"]) == (4, 10, ["0005_eod_backup_actual_close.sql",
                                                                   "0006_eod_snapshot_retention.sql",
                                                                   "0007_open_hours.sql",       # 1.2.9: additive
                                                                   "0008_customer_aggregates.sql",  # 1.2.9b: additive
                                                                   "0009_actual_finalization.sql",  # 1.2.9: additive
                                                                   "0010_history_sync.sql"])  # 1.2.13: additive
        assert db.migrate(target)["applied"] == []
        with db.read() as c:
            assert {t: c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in tabs} == before
            assert c.execute("SELECT COUNT(*), SUM(yhat) FROM forecast_values").fetchone() == vals0
        _at("2026-10-18 23:55")
        assert fb.create(SCOPE, "BIGQUERY", "2026-10-18")["status"] == "COMPLETE"
        assert data_checks.run(target)["status"] in ("OK", "WARN")
