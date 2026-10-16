"""v1.1.3: exact forecast formula, feature catalogue, dated history + Excel, quality metrics, training explanation."""
import io
import math
from datetime import date, datetime, timedelta

import numpy as np
import pytest
from openpyxl import load_workbook

from app import db
from app import settings_service as ss
from app import timeutil
from app.ml.features import feature_names
from app.timeutil import VN_TZ


def _job(client, jid):
    return client.get(f"/api/jobs/{jid}").json()


def _train(client):
    j = _job(client, client.post("/api/training/run", json={"force": True}).json()["id"])
    if j["reason_code"] == "PROMOTION_REVIEW_REQUIRED":    # 1.1.7: a model in use is replaced only by explicit promotion
        p = client.post(f"/api/training-runs/{j['result']['result']['training_run_id']}/promote", json={"confirm": True})
        assert p.status_code == 200, p.text
        return p.json()
    assert j["reason_code"] == "FORECAST_PUBLISHED", j
    return j["result"]["result"]


@pytest.fixture
def demo_run(client, fast_model):
    ss.update("horizons", {"enabled": [7]})
    assert client.post("/api/demo/generate", json={"days": 150, "n_skus": 8, "seed": 3}).status_code == 200
    return _train(client)["forecast_run_id"]


def _rows(run_id, method=None):
    with db.read() as c:
        sql = "SELECT sku_id, fc_id, target_date, horizon_day, yhat, method FROM forecast_values WHERE run_id=?"
        p = [run_id]
        if method:
            sql += " AND method=?"
            p.append(method)
        return db.rows(c, sql + " ORDER BY sku_id, fc_id, target_date", p)


def _explain(client, run_id, r):
    return client.get(f"/api/forecast-runs/{run_id}/explain",
                      params={"sku_id": r["sku_id"], "fc_id": r["fc_id"], "target_date": r["target_date"]}).json()


# ------------------------------------------------------------------ G1 formula
def test_lightgbm_formula_chain_is_exact_for_the_published_value(client, demo_run):
    rows = _rows(demo_run, "LIGHTGBM")
    for r in (rows[0], rows[6], rows[-1]):                                  # day 1, day 7 (recursive), another series
        e = _explain(client, demo_run, r)
        L = e["lightgbm"]
        assert L["status"] == "OK" and L["objective"] == "tweedie" and L["link"] == "exp"
        assert L["attribution_method"].startswith("TREE_SHAP")
        contribs = [f["contribution"] for f in L["features"]]
        assert L["contrib_sum"] == pytest.approx(sum(contribs), abs=1e-9)          # sum of the per-feature values
        assert L["bias"] + L["contrib_sum"] == pytest.approx(L["raw_score"], abs=1e-9)
        assert math.exp(L["raw_score"]) == pytest.approx(L["model_output"], rel=1e-9)   # link = exp, LightGBM's own output
        assert L["model_output_matches_link"] is True
        assert L["clipped"] == max(L["model_output"], 0)
        assert round(L["clipped"], 4) == pytest.approx(L["reproduced_yhat"]) == pytest.approx(r["yhat"])   # = published
        assert L["matches_published"] is True
        assert L["n_features"] == len(feature_names(ss.get("model"))) == len(L["features"])
        for f in L["features"]:
            assert f["effect"] == ("NEGLIGIBLE" if abs(f["contribution"]) < 0.005 else "UP" if f["contribution"] > 0 else "DOWN")
        # per-row contributions are not the model-wide importance
        assert {g["feature"] for g in L.get("global_importance", [])} <= {f["name"] for f in L["features"]}
        assert "gain" in L["global_importance"][0] and "contribution" not in L["global_importance"][0]


def test_identity_link_objective_is_explained_without_exp(client, fast_model):
    ss.update("model", {"objective": "regression"})
    ss.update("horizons", {"enabled": [7]})
    client.post("/api/demo/generate", json={"days": 150, "n_skus": 8, "seed": 3})
    run = _train(client)["forecast_run_id"]
    r = _rows(run, "LIGHTGBM")[0]
    L = _explain(client, run, r)["lightgbm"]
    assert L["objective"] == "regression" and L["link"] == "identity"
    assert L["bias"] + L["contrib_sum"] == pytest.approx(L["raw_score"], abs=1e-9)
    assert L["model_output"] == pytest.approx(L["raw_score"], abs=1e-9)          # no transformation
    assert round(max(L["model_output"], 0), 4) == pytest.approx(r["yhat"])


# ------------------------------------------------------------------ G2 features: values come from the source data
def test_feature_values_windows_and_no_future_data(client, demo_run):
    run = client.get(f"/api/forecast-runs/{demo_run}").json()
    r = next(x for x in _rows(demo_run, "LIGHTGBM") if x["horizon_day"] == 1)
    e = _explain(client, demo_run, r)
    f = {x["name"]: x for x in e["lightgbm"]["features"]}
    t = date.fromisoformat(r["target_date"])
    with db.read() as c:
        dem = {d: q for d, q in c.execute("SELECT order_date, qty FROM demand_daily WHERE scope_key=? AND sku_id=? AND fc_id=?",
                                           (run["scope_key"], r["sku_id"], r["fc_id"]))}
        cal = {d for (d,) in c.execute("SELECT DISTINCT order_date FROM demand_daily WHERE scope_key=?", (run["scope_key"],))}
    def q(d):                       # value of a day as the model sees it: unknown if no data at all that day
        ds = str(d)
        return None if ds not in cal else float(dem.get(ds, 0.0))
    for k in (1, 2, 7):
        want = q(t - timedelta(days=k))
        got = f[f"lag_{k}"]["value"]
        assert (got is None and want is None) or got == pytest.approx(want)
        assert f[f"lag_{k}"]["window"] == {"from": str(t - timedelta(days=k)), "to": str(t - timedelta(days=k))}
    win = [q(t - timedelta(days=k)) for k in range(1, 8)]
    known = [v for v in win if v is not None]
    assert f["roll_mean_7"]["value"] == pytest.approx(sum(known) / len(known))            # denominator = known days
    w28 = [v for v in (q(t - timedelta(days=k)) for k in range(1, 29)) if v is not None]
    assert f["zero_share_28"]["value"] == pytest.approx(sum(1 for v in w28 if v == 0) / len(w28))
    assert f["roll_std_28"]["value"] == pytest.approx(float(np.std(w28)), rel=1e-5)       # population std
    assert f["dow"]["value"] == t.weekday() and f["dow"]["window"] == {"from": str(t), "to": str(t)}
    for x in f.values():                                                                    # nothing after data_as_of
        if x["window"] and not x["includes_target_day"]:
            assert x["window"]["to"] <= run["data_as_of"]
        assert x["catalog"]["vi"]["formula"] and x["catalog"]["en"]["formula"]


def test_feature_catalogue_covers_every_model_feature(client):
    cat = client.get("/api/features/catalog", params={"target_date": "2026-10-14"}).json()
    names = [x["name"] for x in cat["features"]]
    assert names == feature_names(ss.get("model"))
    assert all(x["group"] != "OTHER" and x["vi"]["formula"] and x["vi"]["missing"] is not None for x in cat["features"])
    rm = next(x for x in cat["features"] if x["name"] == "roll_mean_28")
    assert rm["window"] == {"from": "2026-09-16", "to": "2026-10-13"} and "ĐÃ BIẾT" in rm["vi"]["formula"]


# ------------------------------------------------------------------ G4 history: weekday, Excel, no double counting
def test_history_weekday_and_excel_export(client, demo_run):
    for method in ("LIGHTGBM", "FALLBACK_NEW_SERIES"):
        r = next(x for x in _rows(demo_run, method) if x["horizon_day"] == 1)
        e = _explain(client, demo_run, r)
        h = e["history"]
        for d in h["days"]:
            assert d["weekday"] == date.fromisoformat(d["date"]).weekday()
        if method == "LIGHTGBM":
            day1 = next(d for d in h["days"] if d["date"] == e["identity"]["data_as_of"])
            assert {"lag_1", "roll_mean_7", "days_since_last_sale"} <= set(day1["used_by"])
        resp = client.get(f"/api/forecast-runs/{demo_run}/explain/history.xlsx",
                          params={"sku_id": r["sku_id"], "fc_id": r["fc_id"], "target_date": r["target_date"], "lang": "vi"})
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        assert wb.sheetnames == ["Thông tin", "Lịch sử demand"]
        ws = wb["Lịch sử demand"]
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == len(h["days"])
        names = ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật"]
        for (d_, wd, qty, status, used, _by, _note), d in zip(rows, h["days"]):
            assert isinstance(d_, datetime) and d_.date().isoformat() == d["date"] <= e["identity"]["data_as_of"]
            assert wd == names[d_.weekday()]
            if d["qty"] is None:
                assert qty is None                                  # unknown stays empty, never 0
            else:
                assert isinstance(qty, (int, float)) and qty == d["qty"]
            assert (used == "Có") == d["in_calc"]
        assert all(ws.cell(row=i, column=1).number_format == "DD/MM/YYYY" for i in range(2, ws.max_row + 1))
        info = {k: v for k, v in wb["Thông tin"].iter_rows(min_row=2, values_only=True)}
        assert info["Mã lần dự báo (run ID)"] == demo_run and info["SKU"] == r["sku_id"] and info["FC/kho"] == r["fc_id"]
        if h.get("calc"):
            assert info["Tổng demand các ngày dùng"] == pytest.approx(sum(x[2] for x in rows if x[4] == "Có" and x[2] is not None))


def test_history_is_daily_demand_and_ignores_intraday_snapshots(client, demo_run):
    r = next(x for x in _rows(demo_run, "LIGHTGBM") if x["horizon_day"] == 1)
    before = _explain(client, demo_run, r)["history"]["days"]
    now = datetime.now(VN_TZ).replace(hour=10, minute=5, second=0, microsecond=0)
    for hh in (0, 2):                                                # two snapshots of the same day
        timeutil.set_now_override(lambda x=hh: now + timedelta(hours=x))
        client.post("/api/jobs/run", json={"job_type": "D0_SYNC"})
    after = _explain(client, demo_run, r)["history"]["days"]
    assert before == after                                           # snapshots never enter (or double) the history


# ------------------------------------------------------------------ G5 metrics
def test_quality_block_defines_wape_bias_points_and_gate(client, demo_run):
    ss.update("model", {"backtest_horizon_days": 14})                # default fold length: 14 days, step 7 -> overlap
    demo_run = _train(client)["forecast_run_id"]
    r = next(x for x in _rows(demo_run, "LIGHTGBM"))
    q = _explain(client, demo_run, r)["quality"]
    for m in (q["backtest_pooled"], q["backtest_this_method"]):
        assert m["sum_abs_error"] / m["sum_actual"] == pytest.approx(m["wape"])
        assert (m["sum_forecast"] - m["sum_actual"]) / m["sum_actual"] == pytest.approx(m["bias"])
    gate = q["quality_gate"]
    qs = ss.get("quality")
    th = {c["check"]: c for c in gate["checks"]}
    assert th["WAPE_MAX"]["threshold"] == qs.wape_max and th["ABS_BIAS_MAX"]["threshold"] == qs.abs_bias_max
    assert th["WAPE_MAX"]["passed"] == (q["backtest_pooled"]["wape"] <= qs.wape_max), (th, q["backtest_pooled"]["wape"], qs.wape_max)   # raw, not rounded
    assert th["MIN_BACKTEST_POINTS"]["value"] == q["backtest_pooled"]["n_points"]
    assert gate["thresholds_approved_by_business"] is False
    assert q["backtest_config"]["horizon_days"] == 14 and q["backtest_config"]["step_days"] == 7
    assert q["folds_overlap"] is True and len(q["backtest_folds"]) == 2
    assert sum(f["n_points"] for f in q["backtest_folds"]) == q["backtest_pooled"]["n_points"]


def test_backtest_points_reproduce_the_stored_metrics_without_writing(client, demo_run):
    run = client.get(f"/api/forecast-runs/{demo_run}").json()
    with db.read() as c:
        before = {t: c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in ("training_runs", "backtest_runs", "forecast_runs", "job_runs")}
    resp = client.get(f"/api/training-runs/{run['training_run_id']}/backtest-points.xlsx?lang=en")
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    pts = list(wb["Evaluation points"].iter_rows(min_row=2, values_only=True))
    summ = {k: v for k, v in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    a = sum(p[9] for p in pts)
    f = sum(p[10] for p in pts)
    err = sum(p[11] for p in pts)
    assert len(pts) == summ["Evaluation points"]
    assert err / a == pytest.approx(summ["WAPE = Σ|A−F| ÷ ΣA"]) and (f - a) / a == pytest.approx(summ["Bias = (ΣF − ΣA) ÷ ΣA (positive = over-forecast)"])
    tr = client.get(f"/api/training-runs/{run['training_run_id']}").json()
    assert len(pts) == tr["metrics"]["n_points"] and err / a == pytest.approx(tr["metrics"]["wape"])
    assert any(p[9] == 0 for p in pts)                                 # zero actuals are evaluation points too
    from collections import Counter
    dup = Counter((p[4], p[5], p[2]) for p in pts)
    assert summ["SKU×FC×dates in ≥ 2 folds"] == sum(1 for v in dup.values() if v > 1)   # stated, not assumed
    assert summ["Distinct SKU×FC×dates"] == len(dup)
    with db.read() as c:
        after = {t: c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in before}
    assert after == before                                              # read-only


def test_backtest_points_refused_when_data_changed(client, demo_run):
    run = client.get(f"/api/forecast-runs/{demo_run}").json()
    client.post("/api/demo/generate", json={"days": 150, "n_skus": 8, "seed": 9})   # new data for... a NEW batch
    with db.tx() as c:                                                 # simulate late data in the SAME scope
        c.execute("UPDATE demand_daily SET qty = qty + 1 WHERE scope_key=? AND rowid IN (SELECT rowid FROM demand_daily WHERE scope_key=? LIMIT 1)",
                  (run["scope_key"], run["scope_key"]))
    r = client.get(f"/api/training-runs/{run['training_run_id']}/backtest-points.xlsx")
    assert r.status_code == 409 and r.json()["code"] == "DATA_CHANGED_SINCE_TRAINING"


# ------------------------------------------------------------------ G6 training explanation
def test_training_explanation_matches_the_database_and_has_no_side_effects(client, demo_run):
    run = client.get(f"/api/forecast-runs/{demo_run}").json()
    with db.read() as c:
        before = {t: c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in ("training_runs", "backtest_runs", "forecast_runs", "job_runs")}
        tr = db.one(c, "SELECT * FROM training_runs WHERE id=?", (run["training_run_id"],))
    o = client.get("/api/model/training-explain").json()
    t = o["training_run"]
    assert t["id"] == tr["id"] and t["n_rows"] == tr["n_rows"] > 0 and t["train_end"] == run["data_as_of"]
    assert o["params"]["objective"] == "tweedie" and o["params"]["tweedie_variance_power"] == 1.2
    assert o["link"] == "exp" and o["early_stopping"] is False and o["validation_set"] is False
    assert o["artifact"]["exists"] and o["artifact"]["sha256_matches"] is True
    assert o["artifact"]["n_trees"] == ss.get("model").num_boost_round
    assert [f["name"] for f in o["features"]] == feature_names(ss.get("model"))
    assert o["forecast_runs"][0]["id"] == demo_run and o["backtest"]["id"]
    assert sum(m["n_series"] for m in o["method_counts"]) == t["n_series"]
    assert o["active"]["active_source"] == "DEMO" and o["active"]["has_dataset"] is True
    with db.read() as c:
        assert {k: c.execute(f"SELECT COUNT(*) FROM {k}").fetchone()[0] for k in before} == before


def test_training_explanation_without_any_training(client):
    o = client.get("/api/model/training-explain").json()
    assert o["training_run"] is None
    ss.update("general", {"active_source": "MANUAL_IMPORT"})
    o = client.get("/api/model/training-explain").json()
    assert o["active"]["has_dataset"] is False and o["active"]["reason_code"] == "NO_ACTIVE_DATASET"
