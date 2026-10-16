import io
import zipfile
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook

from app import db

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(rows, headers=("order_id", "created_at", "sku_id", "fc_id", "order_qty"), sheet="Data"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows(n_days=70, sku="S1", qty=3, start=datetime(2026, 1, 1, 9, 0), prefix="O"):
    return [(f"{prefix}{i}", start + timedelta(days=i), sku, "F1", qty) for i in range(n_days)]


def _upload(client, content, name="data.xlsx", replace=None):
    data = {"replace_batch_id": replace} if replace else {}
    return client.post("/api/imports/upload", files={"file": (name, content, XLSX)}, data=data)


def _import(client, content, replace=None, accept=False):
    up = _upload(client, content, replace=replace)
    assert up.status_code == 200, up.text
    u = up.json()
    v = client.post(f"/api/imports/uploads/{u['upload_id']}/validate",
                    json={"sheet": u["suggested_sheet"], "mapping": u["suggested_mapping"]})
    assert v.status_code == 200, v.text
    b = v.json()
    c = client.post(f"/api/imports/{b['id']}/confirm", json={"accept_row_errors": accept})
    return b, c


def test_template_download_has_instructions_and_data(client):
    r = client.get("/api/templates/import-template.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Instructions", "Fields", "Data"} <= set(wb.sheetnames)
    hdr = [c.value for c in wb["Data"][1]]
    assert hdr[:6] == ["order_id", "created_at", "order_date", "sku_id", "fc_id", "order_qty"]


def test_template_itself_imports(client):
    r = client.get("/api/templates/import-template.xlsx")
    b, c = _import(client, r.content)
    assert b["valid_rows"] == 21 and b["error_rows"] == 0
    assert c.status_code == 200 and c.json()["status"] == "ACTIVE"


def test_sample_files_import(client):
    r = client.get("/api/templates/sample-history.xlsx")
    b, c = _import(client, r.content)
    assert c.status_code == 200 and b["valid_rows"] > 1000
    r = client.get("/api/templates/sample-with-errors.xlsx")
    up = _upload(client, r.content).json()
    v = client.post(f"/api/imports/uploads/{up['upload_id']}/validate",
                    json={"sheet": "Data", "mapping": up["suggested_mapping"]}).json()
    assert v["error_rows"] >= 8 and v["duplicate_rows"] == 1 and v["requires_error_ack"]
    rep = client.get(f"/api/imports/{v['id']}/errors.xlsx")
    wb = load_workbook(io.BytesIO(rep.content))
    assert wb["Lỗi"].max_row > 8                                     # default = Vietnamese report
    wb_en = load_workbook(io.BytesIO(client.get(f"/api/imports/{v['id']}/errors.xlsx?lang=en").content))
    assert wb_en["Errors"].max_row == wb["Lỗi"].max_row


def test_csv_and_bad_files_rejected(client):
    r = _upload(client, b"order_id,sku\n1,2\n", name="data.csv")
    assert r.status_code == 415 and r.json()["code"] == "NOT_XLSX"
    r = _upload(client, b"this is not a zip", name="data.xlsx")
    assert r.status_code == 400 and r.json()["code"] == "CORRUPT_FILE"
    r = _upload(client, b"", name="data.xlsx")
    assert r.json()["code"] == "EMPTY_FILE"
    # macro container disguised as .xlsx
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_xlsx(_rows(2)))) as src, zipfile.ZipFile(buf, "w") as dst:
        for i in src.infolist():
            dst.writestr(i, src.read(i.filename))
        dst.writestr("xl/vbaProject.bin", b"x")
    r = _upload(client, buf.getvalue())
    assert r.json()["code"] == "MACRO_NOT_ALLOWED"


def test_missing_required_column_and_wrong_types(client):
    content = _xlsx([("O1", "2026-01-01 10:00", "S1", 3)], headers=("order_id", "created_at", "sku_id", "order_qty"))
    u = _upload(client, content).json()
    r = client.post(f"/api/imports/uploads/{u['upload_id']}/validate",
                    json={"sheet": "Data", "mapping": u["suggested_mapping"]})
    assert r.status_code == 422 and r.json()["code"] == "MISSING_REQUIRED_COLUMNS"
    content = _xlsx([("O1", "hôm qua", "S1", "F1", "nhiều")])
    u = _upload(client, content).json()
    b = client.post(f"/api/imports/uploads/{u['upload_id']}/validate",
                    json={"sheet": "Data", "mapping": u["suggested_mapping"]}).json()
    assert b["valid_rows"] == 0 and b["error_rows"] == 1
    r = client.post(f"/api/imports/{b['id']}/confirm", json={"accept_row_errors": True})
    assert r.status_code == 422 and r.json()["code"] == "NO_VALID_ROWS"


def test_empty_sheet(client):
    u = _upload(client, _xlsx([])).json()
    r = client.post(f"/api/imports/uploads/{u['upload_id']}/validate",
                    json={"sheet": "Data", "mapping": u["suggested_mapping"]})
    assert r.json()["code"] == "EMPTY_FILE"


def test_column_mapping_for_renamed_columns_and_sheet_choice(client):
    wb = Workbook()
    wb.active.title = "Notes"
    ws = wb.create_sheet("Orders2026")
    ws.append(["Mã đơn", "Thời gian", "Mã hàng", "Kho", "SL"])
    ws.append(["X1", "2026-01-01 10:00", "S", "F", 2])
    buf = io.BytesIO()
    wb.save(buf)
    u = _upload(client, buf.getvalue()).json()
    assert u["suggested_sheet"] is None                      # ambiguous -> user must choose
    m = {"order_id": "Mã đơn", "created_at": "Thời gian", "sku_id": "Mã hàng", "fc_id": "Kho", "order_qty": "SL"}
    b = client.post(f"/api/imports/uploads/{u['upload_id']}/validate", json={"sheet": "Orders2026", "mapping": m}).json()
    assert b["valid_rows"] == 1 and b["status"] == "PENDING_REVIEW"


def test_errors_require_acknowledgement(client):
    content = _xlsx(_rows(5) + [("BAD", "2026-01-01 10:00", None, "F1", 1)])
    u = _upload(client, content).json()
    b = client.post(f"/api/imports/uploads/{u['upload_id']}/validate",
                    json={"sheet": "Data", "mapping": u["suggested_mapping"]}).json()
    r = client.post(f"/api/imports/{b['id']}/confirm", json={"accept_row_errors": False})
    assert r.status_code == 422 and r.json()["code"] == "ROW_ERRORS_NOT_ACKNOWLEDGED"
    r = client.post(f"/api/imports/{b['id']}/confirm", json={"accept_row_errors": True})
    assert r.status_code == 200 and r.json()["valid_rows"] == 5


def test_cancel_writes_nothing(client):
    u = _upload(client, _xlsx(_rows(5))).json()
    b = client.post(f"/api/imports/uploads/{u['upload_id']}/validate",
                    json={"sheet": "Data", "mapping": u["suggested_mapping"]}).json()
    r = client.post(f"/api/imports/{b['id']}/cancel")
    assert r.json()["status"] == "CANCELLED"
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM order_lines WHERE batch_id=?", (b["id"],)).fetchone()[0] == 0
        assert c.execute("SELECT COUNT(*) FROM demand_daily WHERE scope_key=?", (b["id"],)).fetchone()[0] == 0


def test_reimport_same_file_creates_independent_batch_without_double_counting(client):
    content = _xlsx(_rows(10))
    a, _ = _import(client, content)
    b, _ = _import(client, content)
    assert a["id"] != b["id"]
    with db.read() as c:
        qa = c.execute("SELECT SUM(qty) FROM demand_daily WHERE scope_key=?", (a["id"],)).fetchone()[0]
        qb = c.execute("SELECT SUM(qty) FROM demand_daily WHERE scope_key=?", (b["id"],)).fetchone()[0]
    assert qa == qb == 30   # each batch is its own dataset; nothing merged


def test_replace_batch_only_affects_selected_batch(client, small_demo):
    a, _ = _import(client, _xlsx(_rows(10, qty=3, prefix="A")))
    b, _ = _import(client, _xlsx(_rows(10, qty=7, prefix="B")))
    with db.read() as c:
        demo_before = c.execute("SELECT COUNT(*), SUM(qty) FROM demand_daily WHERE scope_key=?",
                                (small_demo["batch_id"],)).fetchone()
    new, conf = _import(client, _xlsx(_rows(10, qty=5, prefix="N")), replace=a["id"])
    assert conf.status_code == 200, conf.text
    detail_a = client.get(f"/api/imports/{a['id']}").json()
    assert detail_a["status"] == "REPLACED" and detail_a["replaced_by_batch_id"] == new["id"]
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM order_lines WHERE batch_id=?", (a["id"],)).fetchone()[0] == 0
        assert c.execute("SELECT SUM(qty) FROM demand_daily WHERE scope_key=?", (b["id"],)).fetchone()[0] == 70
        assert c.execute("SELECT SUM(qty) FROM demand_daily WHERE scope_key=?", (new["id"],)).fetchone()[0] == 50
        demo_after = c.execute("SELECT COUNT(*), SUM(qty) FROM demand_daily WHERE scope_key=?",
                               (small_demo["batch_id"],)).fetchone()
    assert tuple(demo_before) == tuple(demo_after)


def test_replace_and_delete_keep_other_batches_with_identical_dates_sku_fc(client):
    # same order ids, dates, SKU and FC in every file: isolation must come from the batch, not from the keys
    a, _ = _import(client, _xlsx(_rows(10, qty=3)))
    b, _ = _import(client, _xlsx(_rows(10, qty=7)))

    def snapshot(batch_id):
        with db.read() as c:
            lines = c.execute("SELECT order_id, created_at, order_date, sku_id, fc_id, order_qty FROM order_lines "
                              "WHERE batch_id=? ORDER BY id", (batch_id,)).fetchall()
            daily = c.execute("SELECT order_date, sku_id, fc_id, qty, n_orders FROM demand_daily WHERE scope_key=? "
                              "ORDER BY order_date", (batch_id,)).fetchall()
        return [tuple(r) for r in lines], [tuple(r) for r in daily]

    b_before = snapshot(b["id"])
    assert len(b_before[0]) == 10 and sum(r[3] for r in b_before[1]) == 70
    new, conf = _import(client, _xlsx(_rows(10, qty=5)), replace=a["id"])
    assert conf.status_code == 200, conf.text
    assert snapshot(a["id"]) == ([], [])
    assert snapshot(b["id"]) == b_before
    assert sum(r[3] for r in snapshot(new["id"])[1]) == 50
    r = client.post(f"/api/imports/{new['id']}/delete", json={"confirm_text": new["id"]})
    assert r.json()["status"] == "DELETED"
    assert snapshot(b["id"]) == b_before
    assert client.get(f"/api/imports/{b['id']}").json()["status"] == "ACTIVE"


def test_invalid_replacement_keeps_old_batch(client):
    a, _ = _import(client, _xlsx(_rows(10, qty=3)))
    bad = _xlsx([("Z", "nope", None, None, -1)])
    u = _upload(client, bad, replace=a["id"]).json()
    v = client.post(f"/api/imports/uploads/{u['upload_id']}/validate",
                    json={"sheet": "Data", "mapping": u["suggested_mapping"]}).json()
    r = client.post(f"/api/imports/{v['id']}/confirm", json={"accept_row_errors": True})
    assert r.status_code == 422
    assert client.get(f"/api/imports/{a['id']}").json()["status"] == "ACTIVE"
    with db.read() as c:
        assert c.execute("SELECT SUM(qty) FROM demand_daily WHERE scope_key=?", (a["id"],)).fetchone()[0] == 30
    # replacing a batch that is no longer active is refused at upload time
    client.post(f"/api/imports/{a['id']}/delete", json={"confirm_text": a["id"]})
    r = _upload(client, _xlsx(_rows(3)), replace=a["id"])
    assert r.status_code == 409 and r.json()["code"] == "REPLACE_TARGET_INVALID"


def test_delete_requires_exact_confirmation_and_keeps_metadata(client):
    a, _ = _import(client, _xlsx(_rows(5)))
    r = client.post(f"/api/imports/{a['id']}/delete", json={"confirm_text": "yes"})
    assert r.status_code == 422
    r = client.post(f"/api/imports/{a['id']}/delete", json={"confirm_text": a["id"]})
    assert r.json()["status"] == "DELETED"
    hist = client.get("/api/imports").json()
    assert any(h["id"] == a["id"] and h["status"] == "DELETED" for h in hist)
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM audit_log WHERE entity_id=?", (a["id"],)).fetchone()[0] >= 2
