"""1.2.5: Forecast & actual history - D+1 only, the newest valid session of each run date, immutable forecasts, actual
read independently (never 0 for a missing actual, PROVISIONAL unless the completion signal confirmed the day), KPI /
WAPE, filters, paging, Excel of all filtered rows, read-only API, migration 0004 on a v3 database.

Kind of evidence: integration on an isolated database (conftest: temporary DF_DATA_DIR + process-wide fallback).
Selection tests use synthetic published runs inserted directly (exact times); the end-to-end tests publish real runs
with the pipeline (DEMO data, fake BigQuery client)."""
import io
import json
import sqlite3
import threading
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from app import data_checks, db, forecast_history as fh, paths, timeutil
from app.bq import product_names
from app.ml import metrics as mt

SCOPE = "BIGQUERY:p.d.t"
NOW = datetime(2026, 10, 20, 3, 0, tzinfo=timezone.utc)          # 20/10/2026 10:00 Vietnam time


@pytest.fixture(autouse=True)
def _frozen():
    timeutil.set_now_override(lambda: NOW)
    product_names._cache.update(updated_at=None, names={})
    yield
    timeutil.set_now_override(None)
    product_names._cache.update(updated_at=None, names={})


def _utc(vn: str) -> str:
    """'2026-10-18 08:00' Vietnam time -> UTC ISO string as stored."""
    t = datetime.fromisoformat(vn).replace(tzinfo=timeutil.VN_TZ)
    return t.astimezone(timezone.utc).isoformat(timespec="seconds")


_seq = {"n": 0}


def _run(created_vn: str, *, job_started_vn: str | None = None, scope=SCOPE, source="BIGQUERY", origin="2026-10-18",
         horizon=3, keys=(("S1", "HN"), ("S2", "HN"), ("S1", "BD")), base=10.0, job_status="SUCCEEDED") -> str:
    """One published forecast run (+ its accepted training run and, optionally, the job that produced it)."""
    _seq["n"] += 1
    n = _seq["n"]
    rid, tid = f"FCR-T{n:03d}", f"TRN-T{n:03d}"
    jid = f"JOB-T{n:03d}" if job_started_vn else None
    d0 = datetime.fromisoformat(origin).date()
    with db.tx() as c:
        c.execute("INSERT INTO training_runs(id, model_version, scope_key, data_source, status, started_at, quality_gate_json) "
                  "VALUES (?,?,?,?, 'ACCEPTED', ?, ?)", (tid, f"MV-{n}", scope, source, _utc(created_vn),
                                                        json.dumps({"assessment": {"usage": {"status": "USABLE_WITH_WARNING"}}})))
        if jid:
            c.execute("INSERT INTO job_runs(id, job_type, idempotency_key, trigger, status, created_at, started_at) "
                      "VALUES (?, 'TRAIN_PUBLISH', ?, 'MANUAL', ?, ?, ?)",
                      (jid, f"k-{jid}", job_status, _utc(job_started_vn), _utc(job_started_vn)))
        c.execute("INSERT INTO forecast_runs(id, training_run_id, model_version, scope_key, data_source, forecast_origin, "
                  "data_as_of, horizons_json, max_horizon, n_series, status, created_at, job_id) "
                  "VALUES (?,?,?,?,?,?,?, '[3]', ?, ?, 'PUBLISHED', ?, ?)",
                  (rid, tid, f"MV-{n}", scope, source, origin, str(d0 - timedelta(days=1)), horizon, len(keys),
                   _utc(created_vn), jid))
        c.executemany("INSERT INTO forecast_values(run_id, sku_id, fc_id, target_date, horizon_day, yhat, method) "
                      "VALUES (?,?,?,?,?,?, 'LIGHTGBM')",
                      [(rid, s, f, str(d0 + timedelta(days=h)), h + 1, base + i + 0.5 * h)
                       for i, (s, f) in enumerate(keys) for h in range(horizon)])
    return rid


def _day(run_date, scope=SCOPE):
    with db.read() as c:
        return db.one(c, "SELECT * FROM forecast_history_days WHERE scope_key=? AND run_date=?", (scope, run_date))


def _sessions(run_date=None):
    with db.read() as c:
        return {r["forecast_run_id"]: r for r in db.rows(
            c, "SELECT * FROM forecast_history_sessions" + (" WHERE run_date=?" if run_date else ""),
            (run_date,) if run_date else ())}


def _count(table):
    with db.read() as c:
        return c.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]


def _demand(day, rows, scope=SCOPE, source="BIGQUERY", loaded="2026-10-20T00:00:00+00:00"):
    with db.tx() as c:
        c.execute("DELETE FROM demand_daily WHERE scope_key=? AND order_date=?", (scope, day))
        c.executemany("INSERT INTO demand_daily(scope_key, data_source, order_date, sku_id, fc_id, qty, n_orders, loaded_at) "
                      "VALUES (?,?,?,?,?,?,1,?)", [(scope, source, day, s, f, q, loaded) for (s, f), q in rows.items()])


def _d1_audit(watermark, mode, scope=SCOPE):
    with db.tx() as c:
        db.audit(c, "D1_SYNC", "scope", scope, {"mode": mode, "watermark": watermark, "backfill": False})
        c.execute("INSERT OR REPLACE INTO sync_state(name, value_json, updated_at) VALUES (?,?,?)",
                  (f"last_d1_sync:{scope}", json.dumps({"watermark": watermark, "mode": mode}), timeutil.now_iso()))


def _snapshot(day, values, at_vn, scope=SCOPE, source="BIGQUERY"):
    sid = f"SNP-{day}-{at_vn[-5:].replace(':', '')}"
    with db.tx() as c:
        c.execute("INSERT INTO d0_snapshots(id, data_source, scope_key, order_date, snapshot_at, sync_status, rows_read, "
                  "rows_valid, rows_error, created_at) VALUES (?,?,?,?,?, 'SUCCEEDED', 1, 1, 0, ?)",
                  (sid, source, scope, day, _utc(at_vn), _utc(at_vn)))
        c.executemany("INSERT INTO d0_snapshot_values(snapshot_id, sku_id, fc_id, cumulative_qty) VALUES (?,?,?,?)",
                      [(sid, s, f, q) for (s, f), q in values.items()])
    return sid


def _q(**kw):
    with db.read() as c:
        qy = fh.Query(c, kw.pop("scope", SCOPE), **kw)
        daily = qy.daily()
        return qy.kpi(daily), daily, list(qy.rows(limit=None)), qy.count()


# ------------------------------------------------------------------ D+1 only, session choice
def test_only_the_d1_forecast_is_kept_and_forecast_values_are_untouched():
    rid = _run("2026-10-18 08:00", horizon=15)
    before = _count("forecast_values")
    res = fh.capture(rid)
    assert res["decision"] == fh.SELECTED and res["run_date"] == "2026-10-18" and res["target_date"] == "2026-10-19"
    with db.read() as c:
        vals = db.rows(c, "SELECT target_date, run_date, COUNT(*) AS n FROM forecast_history_values GROUP BY 1, 2")
        src = db.rows(c, "SELECT sku_id, fc_id, yhat, method FROM forecast_values WHERE run_id=? AND target_date='2026-10-19' "
                         "ORDER BY 1, 2", (rid,))
        kept = db.rows(c, "SELECT sku_id, fc_id, yhat, method FROM forecast_history_values ORDER BY 1, 2")
    assert vals == [{"target_date": "2026-10-19", "run_date": "2026-10-18", "n": 3}]   # never D+7 / D+15 / D+30
    assert kept == src
    assert _count("forecast_values") == before == 3 * 15                                 # multi-day forecast unchanged
    s = _sessions()[rid]
    assert (s["horizon_day"], s["n_rows"], s["capture_mode"], s["usage_status"]) == (2, 3, "AT_PUBLISH", "USABLE_WITH_WARNING")


def test_the_newest_valid_session_of_the_day_is_official_and_older_ones_are_traced():
    a = _run("2026-10-18 08:00")
    b = _run("2026-10-18 12:00", base=20)
    c_ = _run("2026-10-18 18:00", base=30)
    other_day = _run("2026-10-19 09:00", origin="2026-10-19", base=40)
    for r in (a, b, c_, other_day):
        fh.capture(r)
    assert _day("2026-10-18")["forecast_run_id"] == c_
    assert _day("2026-10-19")["forecast_run_id"] == other_day          # another day never touches 18/10
    s = _sessions("2026-10-18")
    assert [s[x]["decision"] for x in (a, b, c_)] == ["SELECTED", "SELECTED", "SELECTED"]   # each was newest when recorded
    with db.read() as c:
        aud = [json.loads(r[0]) for r in c.execute("SELECT details_json FROM audit_log WHERE action='FORECAST_HISTORY_SELECTED' "
                                                   "AND entity_id=? ORDER BY id", (f"{SCOPE}|2026-10-18",))]
    assert [(x["forecast_run_id"], x["replaced"]) for x in aud] == [(a, None), (b, a), (c_, b)]
    _k, _d, rows, _n = _q(run_from="2026-10-18", run_to="2026-10-18")
    assert {r["forecast_run_id"] for r in rows} == {c_} and len(rows) == 3


def test_a_failed_or_invalid_latest_session_keeps_the_previous_valid_one():
    good = _run("2026-10-18 08:00")
    fh.capture(good)
    # a later training that FAILED leaves no published run at all (nothing to record), the official one stays
    with db.tx() as c:
        c.execute("INSERT INTO job_runs(id, job_type, idempotency_key, trigger, status, created_at, started_at) "
                  "VALUES ('JOB-FAIL', 'TRAIN_PUBLISH', 'k-fail', 'MANUAL', 'FAILED', ?, ?)",
                  (_utc("2026-10-18 20:00"), _utc("2026-10-18 20:00")))
    assert fh.reconcile()["missing"] == 0
    # a later published run without a value for D+1 (data too old for the horizon) is NOT_ELIGIBLE and replaces nothing
    stale = _run("2026-10-18 21:00", origin="2026-10-10", horizon=3)
    res = fh.capture(stale)
    assert (res["decision"], res["decision_reason"]) == ("NOT_ELIGIBLE", "NO_D1_VALUES")
    assert _day("2026-10-18")["forecast_run_id"] == good
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM forecast_history_values WHERE forecast_run_id=?", (stale,)).fetchone()[0] == 0
    with pytest.raises(sqlite3.IntegrityError):         # the database refuses a pointer to an ineligible session
        with db.tx() as c:
            c.execute("UPDATE forecast_history_days SET forecast_run_id=? WHERE run_date='2026-10-18'", (stale,))


def test_a_session_that_started_earlier_but_finished_later_never_replaces_a_newer_one():
    newer = _run("2026-10-18 10:30", job_started_vn="2026-10-18 10:00")
    late_old = _run("2026-10-18 11:00", job_started_vn="2026-10-18 09:00", base=99)   # started first, published last
    fh.capture(newer)
    res = fh.capture(late_old)
    assert res["decision"] == fh.OLDER and newer in res["decision_reason"]
    assert _day("2026-10-18")["forecast_run_id"] == newer
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM forecast_history_values WHERE forecast_run_id=?", (late_old,)).fetchone()[0] == 0
    # recorded in the other order the result is the same
    assert data_checks.run(paths.db_path())["status"] in ("OK", "WARN")


def test_session_time_is_the_job_start_and_the_run_date_is_the_vietnam_date():
    # job started 23:50 VN on 17/10, published 00:20 VN on 18/10 -> run date 17/10, D+1 = 18/10 (horizon 1)
    rid = _run("2026-10-18 00:20", job_started_vn="2026-10-17 23:50", origin="2026-10-18")
    res = fh.capture(rid)
    assert (res["run_date"], res["target_date"]) == ("2026-10-17", "2026-10-18")
    assert _sessions()[rid]["session_started_at"] == _utc("2026-10-17 23:50")


def test_concurrent_recording_picks_the_newest_and_never_duplicates():
    ids = [_run(f"2026-10-18 {h:02d}:00", base=h) for h in (8, 10, 12, 14, 16)]
    errors = []

    def go(r):
        try:
            fh.capture(r)
            fh.capture(r)               # a retry racing with the first attempt
        except Exception as e:          # noqa: BLE001
            errors.append(e)
    th = [threading.Thread(target=go, args=(r,)) for r in reversed(ids)] + [threading.Thread(target=fh.reconcile)]
    for t in th:
        t.start()
    for t in th:
        t.join()
    assert not errors
    assert _day("2026-10-18")["forecast_run_id"] == ids[-1]
    assert len(_sessions("2026-10-18")) == 5
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM forecast_history_days").fetchone()[0] == 1
        dup = c.execute("SELECT COUNT(*) FROM (SELECT forecast_run_id, sku_id, fc_id FROM forecast_history_values "
                        "GROUP BY 1, 2, 3 HAVING COUNT(*) > 1)").fetchone()[0]
    assert dup == 0
    assert {x["check"]: x["status"] for x in data_checks.run(paths.db_path())["checks"]}["forecast_history_official_is_newest"] == "OK"


def test_retry_and_reconcile_are_idempotent():
    rid = _run("2026-10-18 08:00")
    first = fh.reconcile()
    assert first["missing"] == 1 and first["recorded"][0]["decision"] == "SELECTED"
    n = (_count("forecast_history_sessions"), _count("forecast_history_values"), _count("audit_log"))
    assert fh.reconcile() == {"missing": 0, "recorded": [], "errors": []}
    assert fh.capture(rid)["already_recorded"] is True
    assert (_count("forecast_history_sessions"), _count("forecast_history_values"), _count("audit_log")) == n
    assert _sessions()[rid]["capture_mode"] == "RECOVERED"


def test_a_recording_failure_never_loses_the_forecast_and_is_retried(monkeypatch):
    rid = _run("2026-10-18 08:00")
    with monkeypatch.context() as m:
        m.setattr(fh, "_record", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("disk full")))
        assert fh.capture_after_publish(rid) is None
    with db.read() as c:
        assert db.one(c, "SELECT id FROM forecast_runs WHERE id=?", (rid,))          # forecast still published
        al = db.one(c, "SELECT level, params_json FROM alerts WHERE code='FORECAST_HISTORY_CAPTURE_FAILED'")
    assert al["level"] == "ERROR" and "disk full" in al["params_json"]
    assert _count("forecast_history_sessions") == 0
    # the next publish records its own run AND the one that failed
    nxt = _run("2026-10-19 08:00", origin="2026-10-19")
    assert fh.capture_after_publish(nxt)["decision"] == "SELECTED"
    s = _sessions()
    assert s[rid]["capture_mode"] == "RECOVERED" and s[nxt]["capture_mode"] == "AT_PUBLISH"


def test_history_rows_are_immutable():
    rid = _run("2026-10-18 08:00")
    fh.capture(rid)
    for sql in ("UPDATE forecast_history_values SET yhat = 0", "DELETE FROM forecast_history_values",
                "UPDATE forecast_history_sessions SET decision='X'", "DELETE FROM forecast_history_sessions",
                "DELETE FROM forecast_history_days", "UPDATE forecast_history_days SET target_date='2026-10-30'"):
        with pytest.raises(sqlite3.IntegrityError):
            with db.tx() as c:
                c.execute(sql)


# ------------------------------------------------------------------ actual
def test_actual_join_missing_is_never_zero_and_status_follows_the_real_signal():
    keys = (("S1", "HN"), ("S2", "HN"), ("S1", "BD"))
    r17 = _run("2026-10-17 09:00", origin="2026-10-17", keys=keys)        # D+1 = 18/10: loaded by D-1 with a signal
    r18 = _run("2026-10-18 09:00", origin="2026-10-18", keys=keys)        # D+1 = 19/10: loaded, fallback only
    r19 = _run("2026-10-19 09:00", origin="2026-10-19", keys=keys)        # D+1 = 20/10 = today: D0 snapshot
    r20 = _run("2026-10-20 09:00", origin="2026-10-20", keys=keys)        # D+1 = 21/10: future
    for r in (r17, r18, r19, r20):
        fh.capture(r)
    _demand("2026-10-18", {("S1", "HN"): 5, ("S2", "HN"): 7})            # S1 x BD had no order -> 0 on a loaded day
    _d1_audit("2026-10-18", "SIGNAL")
    _demand("2026-10-19", {("S1", "HN"): 4})
    _d1_audit("2026-10-19", "FALLBACK_TIME_VERIFIED")
    _snapshot("2026-10-20", {("S2", "HN"): 3}, "2026-10-20 08:00")
    _k, daily, rows, _n = _q(run_from="2026-10-17", run_to="2026-10-20")
    by = {(r["target_date"], r["sku_id"], r["fc_id"]): r for r in rows}
    assert by[("2026-10-18", "S1", "HN")]["actual"] == 5 and by[("2026-10-18", "S1", "BD")]["actual"] == 0
    assert by[("2026-10-18", "S1", "BD")]["actual_status"] == "COMPLETE"
    assert {by[("2026-10-19", s, f)]["actual_status"] for s, f in keys} == {"PROVISIONAL"}
    assert by[("2026-10-19", "S1", "HN")]["actual_reason"] == "D1_NO_COMPLETION_SIGNAL"
    assert by[("2026-10-20", "S2", "HN")]["actual"] == 3 and by[("2026-10-20", "S1", "HN")]["actual"] == 0
    assert by[("2026-10-20", "S1", "HN")]["actual_reason"] == "D0_INTRADAY"
    fut = [by[("2026-10-21", s, f)] for s, f in keys]
    assert all(r["actual"] is None and r["actual_status"] == "NO_ACTUAL" and r["deviation"] is None for r in fut)
    d = {x["target_date"]: x for x in daily}
    assert d["2026-10-21"]["actual"] is None and d["2026-10-18"]["actual"] == 12
    assert len(rows) == 12                                                  # 4 days x 3 series: no row duplicated


def test_a_day_not_loaded_yet_uses_the_last_d0_read_and_a_loaded_day_without_rows_is_unknown():
    r = _run("2026-10-18 09:00")
    fh.capture(r)
    _k, _d, rows, _n = _q(run_from="2026-10-18", run_to="2026-10-18")
    assert {x["actual_reason"] for x in rows} == {"NOT_LOADED_YET"} and {x["actual"] for x in rows} == {None}
    _snapshot("2026-10-19", {("S1", "HN"): 2}, "2026-10-19 18:00")
    _k, _d, rows, _n = _q(run_from="2026-10-18", run_to="2026-10-18")
    assert {x["actual_reason"] for x in rows} == {"D0_LAST_READ"} and sorted(x["actual"] for x in rows) == [0, 0, 2]
    _d1_audit("2026-10-19", "FALLBACK_TIME_VERIFIED")              # D-1 says 19/10 is loaded but it holds no row
    _demand("2026-10-17", {("S1", "HN"): 1})
    _k, _d, rows, _n = _q(run_from="2026-10-18", run_to="2026-10-18")
    assert {x["actual_reason"] for x in rows} == {"DAY_WITHOUT_DATA"} and {x["actual"] for x in rows} == {None}


def test_actual_updates_never_change_the_stored_forecast_nor_create_a_session():
    rid = _run("2026-10-18 09:00")
    fh.capture(rid)
    with db.read() as c:
        kept = db.rows(c, "SELECT * FROM forecast_history_values ORDER BY sku_id, fc_id")
    _demand("2026-10-19", {("S1", "HN"): 5})
    _k1, _d, rows1, _n = _q(run_from="2026-10-18", run_to="2026-10-18")
    _demand("2026-10-19", {("S1", "HN"): 9, ("S2", "HN"): 1}, loaded="2026-10-20T02:00:00+00:00")   # retroactive change
    _k2, _d, rows2, _n = _q(run_from="2026-10-18", run_to="2026-10-18")
    with db.read() as c:
        assert db.rows(c, "SELECT * FROM forecast_history_values ORDER BY sku_id, fc_id") == kept
    assert [r["yhat"] for r in rows1] == [r["yhat"] for r in rows2]
    assert _k1["actual_total"] == 5 and _k2["actual_total"] == 10
    assert {r["actual_as_of"] for r in rows2} == {"2026-10-20T02:00:00+00:00"}
    assert (_count("forecast_history_sessions"), _count("forecast_runs")) == (1, 1)


# ------------------------------------------------------------------ KPI / WAPE
def test_kpi_and_wape_use_only_rows_with_an_actual_and_match_the_metric_module():
    a = _run("2026-10-17 09:00", origin="2026-10-17")
    b = _run("2026-10-18 09:00", origin="2026-10-18")
    for r in (a, b):
        fh.capture(r)
    _demand("2026-10-18", {("S1", "HN"): 5, ("S2", "HN"): 30, ("S1", "BD"): 2})
    _d1_audit("2026-10-18", "FALLBACK_TIME_VERIFIED")      # 19/10 (D+1 of b) has no actual yet
    k, daily, rows, n = _q(run_from="2026-10-17", run_to="2026-10-18")
    comp = [r for r in rows if r["actual"] is not None]
    ref = mt.compute([r["actual"] for r in comp], [r["yhat"] for r in comp])
    assert k["wape"] == pytest.approx(ref["wape"]) and k["bias"] == pytest.approx(ref["bias"])
    assert k["actual_total"] == 37 and k["n_rows_compared"] == 3 and k["n_rows"] == n == 6
    assert k["forecast_total"] == pytest.approx(sum(r["yhat"] for r in rows))
    assert k["deviation"] == pytest.approx(sum(r["yhat"] for r in comp) - 37)
    assert k["status"] == "PROVISIONAL" and k["days_compared"] == 1
    row = next(r for r in rows if (r["sku_id"], r["fc_id"], r["target_date"]) == ("S2", "HN", "2026-10-18"))
    assert row["row_wape"] == pytest.approx(abs(row["yhat"] - 30) / 30)


def test_wape_with_zero_actual_and_without_actual():
    a = _run("2026-10-18 09:00")
    fh.capture(a)
    k, *_ = _q(run_from="2026-10-18", run_to="2026-10-18")
    assert (k["status"], k["wape"], k["actual_total"], k["notes"]) == ("NO_ACTUAL", None, None, ["NO_ACTUAL"])
    _demand("2026-10-19", {})              # loaded day ...
    _demand("2026-10-18", {("S9", "HN"): 1})
    _d1_audit("2026-10-19", "SIGNAL")
    _demand("2026-10-19", {("OTHER", "HN"): 0.0})     # ... with only a zero row of another series
    k, _d, rows, _n = _q(run_from="2026-10-18", run_to="2026-10-18")
    assert k["actual_total"] == 0 and k["wape"] is None and k["notes"] == ["SUM_ACTUAL_ZERO"]
    assert k["status"] == "COMPLETE" and all(r["row_wape"] is None for r in rows)


# ------------------------------------------------------------------ filters, names, paging, export, API
def _named():
    with db.tx() as c:
        c.execute("INSERT OR REPLACE INTO sync_state(name, value_json, updated_at) VALUES ('product_names', ?, 'v1')",
                  (json.dumps({"names": {"S1": "Thuốc Hạ Sốt Paracetamol"}}, ensure_ascii=False),))


def test_filters_combine_search_is_accent_insensitive_and_names_never_duplicate_rows():
    _named()
    for d in ("2026-10-17", "2026-10-18"):
        fh.capture(_run(f"{d} 09:00", origin=d))
    _demand("2026-10-18", {("S1", "HN"): 5})
    _d1_audit("2026-10-18", "FALLBACK_TIME_VERIFIED")
    _k, _d, rows, n = _q(run_from="2026-10-17", run_to="2026-10-18")
    assert n == 6 and len({(r["run_date"], r["sku_id"], r["fc_id"]) for r in rows}) == 6
    assert {r["product_name"] for r in rows if r["sku_id"] == "S1"} == {"Thuốc Hạ Sốt Paracetamol"}
    assert {r["product_name"] for r in rows if r["sku_id"] == "S2"} == {None}
    _k, _d, rows, n = _q(run_from="2026-10-17", run_to="2026-10-18", q="thuoc ha sot", fc_id="HN")
    assert n == 2 and {(r["sku_id"], r["fc_id"]) for r in rows} == {("S1", "HN")}
    _k, _d, rows, n = _q(run_from="2026-10-17", run_to="2026-10-18", actual_status="PROVISIONAL")
    assert n == 3 and {r["target_date"] for r in rows} == {"2026-10-18"}
    k, _d, rows, n = _q(run_from="2026-10-17", run_to="2026-10-18", actual_status="NO_ACTUAL", fc_id="BD")
    assert n == 1 and k["actual_total"] is None and k["n_rows"] == 1
    with pytest.raises(Exception):
        _q(run_from="2026-10-17", run_to="2026-10-18", actual_status="DONE")


def test_api_pages_sorts_and_exports_every_filtered_row_without_writing(client):
    _named()
    keys = tuple((f"S{i}", fc) for i in range(1, 40) for fc in ("HN", "BD"))
    for d in ("2026-10-16", "2026-10-17", "2026-10-18"):
        fh.capture(_run(f"{d} 09:00", origin=d, keys=keys))
    _demand("2026-10-18", {("S1", "HN"): 500, ("S2", "HN"): 1})
    _d1_audit("2026-10-18", "FALLBACK_TIME_VERIFIED")
    tables = ("job_runs", "training_runs", "forecast_runs", "forecast_values", "forecast_history_sessions",
              "forecast_history_values", "forecast_history_days", "audit_log", "alerts", "app_settings")
    before = {t: _count(t) for t in tables}
    p = {"scope_key": SCOPE, "run_from": "2026-10-16", "run_to": "2026-10-18", "fc_id": "HN"}
    ov = client.get("/api/history/overview", params={"scope_key": SCOPE}).json()
    assert [d["run_date"] for d in ov["days"]] == ["2026-10-18", "2026-10-17", "2026-10-16"] and ov["fcs"] == ["BD", "HN"]
    sm = client.get("/api/history/summary", params=p).json()
    assert sm["kpi"]["n_rows"] == 3 * 39 and len(sm["daily"]) == 3
    page1 = client.get("/api/history/rows", params={**p, "limit": 50, "offset": 0, "sort": "abs_error"}).json()
    page3 = client.get("/api/history/rows", params={**p, "limit": 50, "offset": 100, "sort": "abs_error"}).json()
    assert page1["total"] == 117 and len(page1["rows"]) == 50 and len(page3["rows"]) == 17
    assert page1["rows"][0]["sku_id"] == "S1" and page1["rows"][0]["target_date"] == "2026-10-18"   # largest |F - A|
    x = client.get("/api/history/export.xlsx", params={**p, "lang": "vi"})
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Thông tin", "Phiên forecast", "Theo ngày", "Chi tiết SKU x FC"]
    det = list(wb["Chi tiết SKU x FC"].iter_rows(values_only=True))
    assert len(det) - 1 == 117                                            # all filtered rows, not the 50 on screen
    head = det[0]
    assert head[:8] == ("Ngày chạy forecast", "Ngày được dự báo", "Mã phiên (run ID)", "Phiên bản mô hình",
                        "Thời điểm phiên được chọn", "SKU", "Tên sản phẩm", "FC/kho")
    assert {r[7] for r in det[1:]} == {"HN"}                              # nothing outside the filter
    assert {r[12] for r in det[1:]} == {"Tạm tính", "Chưa có dữ liệu thực tế"}
    s1 = [r for r in det[1:] if r[5] == "S1"]
    assert {r[6] for r in s1} == {"Thuốc Hạ Sốt Paracetamol"}
    assert {r[6] for r in det[1:] if r[5] == "S3"} == {"Chưa có tên sản phẩm"}
    assert all(r[9] is None for r in det[1:] if r[1].date().isoformat() != "2026-10-18")   # no actual = empty, not 0
    en = load_workbook(io.BytesIO(client.get("/api/history/export.xlsx", params={**p, "lang": "en"}).content))
    assert en.sheetnames[-1] == "SKU x FC detail"
    assert {t: _count(t) for t in tables} == before                       # reading never writes, starts no job


def test_default_scope_is_the_active_source_even_without_history(client, small_demo):
    fh.capture(_run("2026-10-18 09:00"))                  # BIGQUERY history exists, active source is DEMO
    ov = client.get("/api/history/overview").json()
    assert ov["data_source"] == "DEMO" and ov["days"] == [] and ov["scope_key"] == ov["active_scope_key"]
    assert [s["scope_key"] for s in ov["scopes"]] == [SCOPE]
    sm = client.get("/api/history/summary").json()
    assert sm["kpi"]["forecast_total"] is None and all(not d["has_history"] for d in sm["daily"])
    assert client.get("/api/history/summary", params={"scope_key": "nope"}).status_code == 422


# ------------------------------------------------------------------ end to end with the real pipeline
def test_a_real_publish_records_its_d1_history(client, small_demo):
    timeutil.set_now_override(None)
    r = client.post("/api/training/run", json={"force": True})
    assert r.status_code == 200, r.text
    with db.read() as c:
        fr = db.one(c, "SELECT * FROM forecast_runs ORDER BY rowid DESC LIMIT 1")
        job = db.one(c, "SELECT * FROM job_runs WHERE id=?", (fr["job_id"],))
        s = db.one(c, "SELECT * FROM forecast_history_sessions WHERE forecast_run_id=?", (fr["id"],))
        n = c.execute("SELECT COUNT(*) FROM forecast_history_values WHERE forecast_run_id=?", (fr["id"],)).fetchone()[0]
    assert s and s["capture_mode"] == "AT_PUBLISH" and s["decision"] == "SELECTED"
    assert s["session_started_at"] == job["started_at"]
    run_date = fh._vn_date(job["started_at"])
    assert s["run_date"] == str(run_date) and s["target_date"] == str(run_date + timedelta(days=1))
    assert n == fr["n_series"] == s["n_rows"]
    assert data_checks.run(paths.db_path())["status"] in ("OK", "WARN")
    assert {x["check"]: x["status"] for x in data_checks.run(paths.db_path())["checks"]}["forecast_history_all_runs_recorded"] == "OK"


def test_migration_0004_on_a_v3_database_keeps_every_row_and_records_old_runs(tmp_path, monkeypatch):
    target = tmp_path / "v3.db"
    conn = db.connect(target, create=True)
    db.current_version(conn)
    for v, p in db.available_migrations():
        if v > 3:
            break
        for stmt in db._split_sql(p.read_text(encoding="utf-8")):
            conn.execute(stmt)
        conn.execute("INSERT INTO schema_migrations(version, name, applied_at) VALUES (?,?,?)", (v, p.name, "x"))
    conn.close()
    with monkeypatch.context() as m:          # never monkeypatch.undo(): conftest's isolation must stay in place
        m.setattr(paths, "db_path", lambda: target)
        _run("2026-10-18 09:00")
        before = {t: _count(t) for t in ("training_runs", "forecast_runs", "forecast_values", "job_runs")}
        res = db.migrate(target)
        assert res["from"] == 3 and res["to"] == 4 and res["applied"] == ["0004_forecast_history.sql"]
        assert {t: _count(t) for t in before} == before
        assert fh.reconcile()["recorded"][0]["decision"] == "SELECTED"
        assert data_checks.run(target)["status"] in ("OK", "WARN")


def test_remote_proxy_allows_the_history_reads_only():
    import importlib.util
    from pathlib import Path
    spec = importlib.util.spec_from_file_location(
        "rp125", Path(__file__).resolve().parents[2] / "scripts" / "remote" / "readonly_proxy.py")
    rp = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(rp)
    for p in ("/api/history/overview", "/api/history/summary?run_from=2026-10-01", "/api/history/rows?limit=10",
              "/api/history/export.xlsx?lang=vi"):
        assert rp.classify("GET", p) is None, p
    assert rp.classify("POST", "/api/history/summary") == (405, "REMOTE_READ_ONLY")
    assert rp.classify("GET", "/api/history/reconcile") == (403, "REMOTE_FORBIDDEN")
    assert rp.classify("GET", "/api/history") == (403, "REMOTE_FORBIDDEN")
