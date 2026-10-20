"""1.2.9: day of week (one convention, persisted switch, target-day value), SKU open hours as a MANDATORY model input of
the BigQuery D-1 source (roles of every D-1 column, validation, training gate, feature matrix), weekday names in the
daily forecast table.

All on isolated databases (conftest) with synthetic data; BigQuery is the fake client here - the real-BigQuery check of
the aggregation SQL is test_v129_real_bq_open_hours.py (opt-in)."""
import json
import sqlite3
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from app import db, paths
from app import settings_service as ss
from app import weekday as wd
from app.bq import connector
from app.data import identity, open_hours
from app.messages import AppError
from app.ml import feature_registry as registry
from app.ml.features import (OPEN_HOURS, Panel, build_panel, calendar_names, calendar_row, direct_extra, dow_value,
                             feature_names, features_at, open_hours_on, origin_features, static_codes)
from fake_bq import FakeBQ

WEEK = ["2026-10-19", "2026-10-20", "2026-10-21", "2026-10-22", "2026-10-23", "2026-10-24", "2026-10-25"]   # Mon..Sun
SCHEMA = [("so_code", "STRING"), ("created_date", "DATE"), ("sku_code", "STRING"), ("fc_code", "STRING"),
          ("demand_quantity", "INT64"), ("sales_price", "FLOAT64"), ("product_id", "INT64"), ("seller_class", "STRING"),
          ("refresh_time", "STRING"), ("total_open_hours", "FLOAT64"), ("total_open_days", "FLOAT64")]
FM = {"order_id": {"column": "so_code", "usage": "OPTIONAL"}, "order_date": {"column": "created_date", "usage": "REQUIRED"},
      "created_at": {"column": None, "usage": "UNUSED"},
      "sku_id": {"column": "sku_code", "usage": "REQUIRED"}, "fc_id": {"column": "fc_code", "usage": "REQUIRED"},
      "order_qty": {"column": "demand_quantity", "usage": "REQUIRED"},
      "unit_price": {"column": "sales_price", "usage": "OPTIONAL"}}
REF = {"product_id": {"column": "product_id", "usage": "OPTIONAL"}, "sales_price": {"column": "sales_price", "usage": "OPTIONAL"},
       "open_hours": {"column": "total_open_hours", "usage": "OPTIONAL"},
       "open_days": {"column": "total_open_days", "usage": "OPTIONAL"}}
BQ_SCOPE = "BIGQUERY:lakehouse-prod-394907.sales.orders"


def _fm(d):
    return {k: ss.FieldMapping(**v) for k, v in d.items()}


def _cols():
    return [{"name": n, "type": t} for n, t in SCHEMA]


# ================================================================== A. day of week
def test_iso_weekday_monday_1_to_sunday_7_and_labels():
    assert [wd.iso_weekday(d) for d in WEEK] == [1, 2, 3, 4, 5, 6, 7]
    assert [wd.label(wd.iso_weekday(d)) for d in WEEK] == ["Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "Chủ nhật"]
    assert wd.label(7, "en") == "Sunday"
    # every input type of the code base gives the weekday of its own calendar day
    assert wd.iso_weekday(date(2026, 10, 25)) == wd.iso_weekday(datetime(2026, 10, 25, 23, 59)) == \
        wd.iso_weekday(pd.Timestamp("2026-10-25")) == wd.iso_weekday(np.datetime64("2026-10-25")) == 7
    # year end, leap day, month end
    assert (wd.iso_weekday("2026-12-31"), wd.iso_weekday("2027-01-01"), wd.iso_weekday("2028-02-29")) == (4, 5, 2)


@pytest.mark.parametrize("bad", [None, 0, 8, -1, 2.5, "x", "", True, False, float("nan")])
def test_weekday_label_refuses_missing_or_invalid_values(bad):
    assert wd.label(bad) is None


def test_weekday_of_a_business_date_never_shifts_with_a_time_zone():
    # 2026-10-19 23:30 in Vietnam is still Monday; the business date string is what the app stores
    vn = pd.Timestamp("2026-10-19 23:30", tz="Asia/Ho_Chi_Minh")
    assert wd.iso_weekday(vn) == 1 and wd.iso_weekday(str(vn.date())) == 1
    with pytest.raises(ValueError):
        wd.iso_weekday("2026-02-30")                  # an impossible date is an error, never a guessed weekday


def test_new_models_encode_dow_1_to_7_old_models_keep_0_to_6():
    new = ss.ModelSettings(forecast_strategy="DIRECT", weekday_encoding="ISO")
    old = ss.ModelSettings(forecast_strategy="DIRECT")                       # stored before 1.2.9: no encoding
    for d, n in zip(WEEK, range(1, 8)):
        ts = pd.Timestamp(d)
        assert dow_value(ts, new) == n and dow_value(ts, old) == n - 1
        vn, vo = dict(zip(calendar_names(new), calendar_row(ts, new))), dict(zip(calendar_names(old), calendar_row(ts, old)))
        assert vn["is_weekend"] == vo["is_weekend"] == (1.0 if n >= 6 else 0.0)
        assert direct_extra(new, ts, 2)[0] == n and direct_extra(old, ts, 2)[0] == n - 1
    assert feature_names(new) == feature_names(old)                           # same columns, only the code differs


def test_weekday_switch_is_saved_and_survives_reload_and_restart(client):
    for value in (False, True, False):
        r = client.put("/api/settings/model", json={**client.get("/api/settings").json()["model"], "feature_weekday": value})
        assert r.status_code == 200 and r.json()["feature_weekday"] is value
        assert client.get("/api/settings").json()["model"]["feature_weekday"] is value          # page reload
        with db.read() as c:                                                                  # what a restart reads
            assert json.loads(c.execute("SELECT value_json FROM app_settings WHERE section='model'").fetchone()[0])[
                "feature_weekday"] is value
    from fastapi.testclient import TestClient
    from app.main import create_app
    with TestClient(create_app(), headers={"X-DF-Client": "web"}) as c2:                          # a new app instance
        assert c2.get("/api/settings").json()["model"]["feature_weekday"] is False
    names = feature_names(ss.get("model"))
    assert "dow" not in names and "is_weekend" not in names


def test_model_settings_refuse_switching_open_hours_off_and_never_store_the_dow_encoding(client):
    m = client.get("/api/settings").json()["model"]
    assert m["feature_open_hours"] is None and m["weekday_encoding"] is None
    r = client.put("/api/settings/model", json={**m, "feature_open_hours": False})
    assert r.status_code == 422 and r.json()["code"] == "OPEN_HOURS_MANDATORY"
    r = client.put("/api/settings/model", json={**m, "weekday_encoding": "ISO"})
    assert r.status_code == 200 and r.json()["weekday_encoding"] is None


def test_demo_training_uses_target_day_weekday_1_to_7_and_inference_matches(client, fast_model):
    assert client.post("/api/demo/generate", json={"days": 150, "n_skus": 8, "seed": 3}).status_code == 200
    j = client.post("/api/training/run", json={}).json()
    assert j["status"] == "SUCCEEDED", j
    run = client.get("/api/system/status").json()["latest_forecast_run"]
    with db.read() as c:
        tr = db.one(c, "SELECT * FROM training_runs WHERE id=?", (run["training_run_id"],))
    fc = json.loads(tr["feature_config_json"])
    assert fc["config"]["weekday_encoding"] == "ISO" and "dow" in fc["features"]
    assert fc["config"]["feature_open_hours"] is False and not set(OPEN_HOURS) & set(fc["features"])   # DEMO: no column
    assert fc["data_identity"]["columns"] == list(identity.CONTENT_COLUMNS)
    import lightgbm as lgb
    assert lgb.Booster(model_file=str(paths.models_dir() / tr["model_path"])).feature_name() == fc["features"]
    vals = client.get(f"/api/forecast-runs/{run['id']}/values?limit=5000").json()["rows"]
    x = next(v for v in vals if v["method"] == "LIGHTGBM")
    ex = client.get(f"/api/forecast-runs/{run['id']}/explain",
                    params={"sku_id": x["sku_id"], "fc_id": x["fc_id"], "target_date": x["target_date"]}).json()
    assert ex["reproduction"]["status"] == "OK" and ex["lightgbm"]["matches_published"]
    feats = {f["name"]: f["value"] for f in ex["lightgbm"]["features"]}
    assert feats["dow"] == wd.iso_weekday(x["target_date"])                  # the TARGET day, 1 = Monday ... 7 = Sunday
    assert {d["weekday"] for d in ex["history"]["days"]} <= set(range(1, 8))
    d0 = ex["history"]["days"][0]
    assert d0["weekday"] == wd.iso_weekday(d0["date"])


# ================================================================== B. open hours: rules
@pytest.mark.parametrize("nulls, lo, hi, out", [
    (0, 24, 24, (24.0, "OK")), (0, 0, 0, (0.0, "OK")), (0, 13.85, 13.85, (13.85, "OK")),
    (1, 10, 10, (None, "MISSING")), (3, None, None, (None, "MISSING")), (0, None, None, (None, "MISSING")),
    (0, 5, 9.63, (None, "CONFLICT")), (0, -1, -1, (None, "OUT_OF_RANGE")), (0, 24.5, 24.5, (None, "OUT_OF_RANGE")),
    (0, 3, 30, (None, "OUT_OF_RANGE")), (None, None, None, (None, None))])
def test_open_hours_cell_rules(nulls, lo, hi, out):
    assert open_hours.cell(nulls, lo, hi) == out


# ================================================================== B. mapping roles
def _bq(**extra):
    ss.update("bigquery", {"dataset_id": "sales", "source_table": "orders", "location": "asia-southeast1",
                           "partition_column": "", "field_mapping": FM, "reference_mapping": REF, **extra})


def test_open_hours_must_be_mapped_and_cannot_be_unused():
    _bq()
    assert connector.validate_mapping(_cols())["ok"]
    for bad in ({**REF, "open_hours": {"column": None, "usage": "OPTIONAL"}},
                {**REF, "open_hours": {"column": "total_open_hours", "usage": "UNUSED"}},
                {k: v for k, v in REF.items() if k != "open_hours"}):
        res = connector.validate_mapping(_cols(), _fm(FM), _fm(bad))
        assert not res["ok"]
        assert {"field": "open_hours", "code": "REQUIRED_NOT_MAPPED", "reference": True, "role": "REQUIRED_FEATURE"} in res["errors"]
    gone = [c for c in _cols() if c["name"] != "total_open_hours"]                 # the column left the source
    assert ("open_hours", "COLUMN_NOT_FOUND") in {(e["field"], e["code"]) for e in connector.validate_mapping(gone)["errors"]}
    text = [c if c["name"] != "total_open_hours" else {"name": c["name"], "type": "STRING"} for c in _cols()]
    assert ("open_hours", "TYPE_INCOMPATIBLE") in {(e["field"], e["code"]) for e in connector.validate_mapping(text)["errors"]}


def test_one_column_cannot_be_key_or_target_and_a_model_input():
    _bq()
    leak = connector.validate_mapping(_cols(), _fm(FM), _fm({**REF, "open_hours": {"column": "demand_quantity", "usage": "OPTIONAL"}}))
    assert any(e["code"] == "COLUMN_ROLE_CONFLICT" and e["column"] == "demand_quantity" and
               set(e["roles"]) == {"TARGET", "REQUIRED_FEATURE"} for e in leak["errors"])
    idf = connector.validate_mapping(_cols(), _fm({**FM, "unit_price": {"column": "sku_code", "usage": "OPTIONAL"}}))
    assert any(e["code"] == "COLUMN_ROLE_CONFLICT" and set(e["roles"]) == {"KEY", "OPTIONAL_FEATURE"} for e in idf["errors"])
    two_keys = connector.validate_mapping(_cols(), _fm({**FM, "fc_id": {"column": "sku_code", "usage": "REQUIRED"}}))
    assert any(e["code"] == "SAME_COLUMN_FOR_SEVERAL_KEYS" for e in two_keys["errors"])        # unchanged 1.1.x rule
    # a REFERENCE field may read the column of an input (sales_price: unit price input + reference) - a warning only
    ok = connector.validate_mapping(_cols())
    assert ok["ok"] and any(w["code"] == "COLUMN_USED_TWICE" and w["column"] == "sales_price" for w in ok["warnings"])


def test_every_d1_column_is_listed_with_its_role_and_nothing_is_classified_by_itself():
    _bq()
    roles = {r["name"]: r for r in connector.column_roles(_cols())}
    assert set(roles) == {n for n, _ in SCHEMA}
    assert roles["sku_code"]["role"] == roles["fc_code"]["role"] == roles["created_date"]["role"] == "KEY"
    assert roles["so_code"]["role"] == "KEY" and roles["demand_quantity"]["role"] == "TARGET"
    assert roles["total_open_hours"]["role"] == "REQUIRED_FEATURE" and roles["total_open_hours"]["model_input"]
    assert roles["sales_price"]["role"] == "OPTIONAL_FEATURE"                 # input (unit_price) wins over reference
    assert {f["field"] for f in roles["sales_price"]["fields"]} == {"unit_price", "sales_price"}
    assert roles["product_id"]["role"] == "REFERENCE" and not roles["product_id"]["model_input"]
    for n in ("seller_class", "refresh_time"):                             # in the source, used by no field
        assert roles[n]["status"] == "UNCLASSIFIED" and roles[n]["role"] is None and not roles[n]["model_input"]
    _bq(extra_reference_columns=["seller_class"])
    roles = {r["name"]: r for r in connector.column_roles(_cols())}
    assert roles["seller_class"]["status"] == "EXTRA_REFERENCE" and roles["seller_class"]["role"] == "REFERENCE"
    assert roles["refresh_time"]["status"] == "UNCLASSIFIED"
    missing = connector.column_roles([c for c in _cols() if c["name"] != "product_id"])
    assert {"name": "product_id", "status": "MISSING_IN_SOURCE"}.items() <= next(r for r in missing if r["name"] == "product_id").items()


def test_extra_reference_columns_are_validated():
    _bq()
    res = connector.validate_mapping(_cols(), extra_reference_columns=["sku_code", "seller_class", "not_there"])
    assert {"field": "extra_reference_columns", "code": "EXTRA_REFERENCE_ALREADY_MAPPED", "column": "sku_code",
            "fields": ["sku_id"]} in res["errors"]
    assert any(w["code"] == "EXTRA_REFERENCE_NOT_IN_SOURCE" and w["column"] == "not_there" for w in res["warnings"])
    assert not any(e.get("column") == "seller_class" for e in res["errors"])
    with pytest.raises(Exception):
        ss.update("bigquery", {"extra_reference_columns": ["a; DROP TABLE x"]})


def test_a_mapping_saved_before_1_2_9_reads_unchanged_and_keeps_its_hash():
    """Production saved its mapping on 1.2.8 (open hours as a reference column, no extra_reference_columns): it must
    load as is and keep its mapping hash - otherwise the validated production state (D-1, training) would be lost."""
    _bq()
    raw = json.loads(ss.get("bigquery").model_dump_json())
    raw.pop("extra_reference_columns")
    with db.tx() as c:
        c.execute("UPDATE app_settings SET value_json=? WHERE section='bigquery'", (json.dumps(raw),))
    bq = ss.get("bigquery")
    assert bq.extra_reference_columns == [] and bq.reference_mapping["open_hours"].column == "total_open_hours"
    import hashlib
    key = {"table": f"{bq.project_id}.{bq.dataset_id}.{bq.source_table}",            # the 1.2.8 formula, verbatim
           "field_mapping": {k: v.model_dump() for k, v in sorted(bq.field_mapping.items())},
           "reference_mapping": {k: v.model_dump() for k, v in sorted(bq.reference_mapping.items())},
           "created_at_is_utc": bq.created_at_is_utc, "partition_column": bq.partition_column,
           "history_start_date": bq.history_start_date, "exact_duplicate_policy": bq.exact_duplicate_policy,
           "history_source": connector._hist_identity(), "window_start_date": ss.get("data_policy").window_start_date}
    assert connector.mapping_hash(bq) == hashlib.sha256(json.dumps(key, sort_keys=True).encode()).hexdigest()[:16]
    ss.update("bigquery", {"extra_reference_columns": ["seller_class"]})
    assert connector.mapping_hash(ss.get("bigquery")) == connector.mapping_hash(bq)    # classification changes no read


def test_mapping_api_saves_extra_reference_columns_and_reports_roles(client):
    _bq()
    connector.set_client_factory(FakeBQ(schema=SCHEMA))
    r = client.post("/api/bigquery/mapping", json={"field_mapping": FM, "reference_mapping": REF,
                                                    "extra_reference_columns": ["seller_class"]})
    assert r.status_code == 200, r.json()
    st = client.get("/api/bigquery/status").json()
    assert st["settings"]["extra_reference_columns"] == ["seller_class"] and st["field_roles"]["open_hours"] == "REQUIRED_FEATURE"
    assert st["mandatory_reference_fields"] == ["open_hours"]
    roles = {x["name"]: x["status"] for x in st["column_roles"]}
    assert roles["seller_class"] == "EXTRA_REFERENCE" and roles["refresh_time"] == "UNCLASSIFIED"
    r = client.post("/api/bigquery/mapping", json={"field_mapping": FM, "reference_mapping": {**REF, "open_hours": {"column": None, "usage": "OPTIONAL"}}})
    assert r.status_code == 422 and r.json()["code"] == "BQ_MAPPING_INVALID" and "REQUIRED_NOT_MAPPED" in r.json()["detail"]


# ================================================================== B. D-1 aggregation, storage, validation
def test_daily_sql_reads_line_counts_and_min_max_of_open_hours():
    _bq()
    sql, _ = connector._daily_sql(ss.get("bigquery"), date(2026, 10, 1), date(2026, 10, 2))
    assert "`total_open_hours` AS ref_open_hours" in sql
    assert "COUNTIF(ref_open_hours IS NULL) AS open_hours_null_rows" in sql
    assert "MIN(ref_open_hours) AS open_hours_min" in sql and "MAX(ref_open_hours) AS open_hours_max" in sql
    assert "AVG(ref_open_hours)" not in sql and "COALESCE(ref_open_hours" not in sql      # no invented aggregate / fill
    # every UNION ALL branch has the same number of columns as the first one (top-level commas of the select list)
    def n_cols(sel: str) -> int:
        depth, n = 0, 1
        for ch in sel:
            depth += ch == "("
            depth -= ch == ")"
            n += ch == "," and depth == 0
        return n
    parts = sql.split("UNION ALL")
    assert len(parts) == 4
    first = n_cols(parts[0].split("SELECT 'VALID'", 1)[1].split("FROM checked", 1)[0])
    assert first == 20          # 1.2.9b: + open_hours_items / open_hours_item_values / open_hours_no_item_rows (MULTI_ITEM)
    for p in parts[1:]:
        assert n_cols(p.split("SELECT", 1)[1].rsplit("FROM", 1)[0]) == first, p


def test_d1_sync_stores_value_and_status_per_day_sku_fc(client):
    from test_bigquery import PASSING_VALIDATION
    y = date.today() - timedelta(days=1)
    base = {"n_orders": 1, "avg_unit_price": 10.0, "avg_discount_rate": None, "promo_share": None, "voucher_share": None,
            "max_source_updated_at": None, "max_batch_id": None}
    rows = [{**base, "order_date": y, "sku_id": "OK", "fc_id": "F1", "qty": 5, "n_rows": 2,
             "open_hours_null_rows": 0, "open_hours_min": 24.0, "open_hours_max": 24.0},
            {**base, "order_date": y, "sku_id": "ZERO", "fc_id": "F1", "qty": 1, "n_rows": 1,
             "open_hours_null_rows": 0, "open_hours_min": 0.0, "open_hours_max": 0.0},
            {**base, "order_date": y, "sku_id": "NULL", "fc_id": "F1", "qty": 2, "n_rows": 2,
             "open_hours_null_rows": 2, "open_hours_min": None, "open_hours_max": None},
            {**base, "order_date": y, "sku_id": "DIFF", "fc_id": "F1", "qty": 3, "n_rows": 2,
             "open_hours_null_rows": 0, "open_hours_min": 19.37, "open_hours_max": 24.0},
            {**base, "order_date": y, "sku_id": "BIG", "fc_id": "F1", "qty": 4, "n_rows": 1,
             "open_hours_null_rows": 0, "open_hours_min": 25.0, "open_hours_max": 25.0}]
    fake = FakeBQ(schema=SCHEMA, responses={"daily": rows, "completion": [{"n": 1}],
                                            "profile": [{"n_rows": 100, "min_date": date(2026, 1, 1), "max_date": y}],
                                            **PASSING_VALIDATION})
    connector.set_client_factory(fake)
    _bq(completion_table="done", completion_date_column="d")
    connector.test_connection()
    ss.update("bigquery", {"schema_snapshot": connector.read_schema()["columns"]})
    connector.profile(7)
    assert connector.validate_source()["passed"]
    ss.update("bigquery", {"production_enabled": True})
    ss.update("general", {"active_source": "BIGQUERY"})
    from app.sync import run_d1_sync
    res = run_d1_sync("J1")
    assert res["open_hours_status_rows"] == {"OK": 2, "MISSING": 1, "OUT_OF_RANGE": 1, "CONFLICT": 1, "MULTI_ITEM": 0, "NOT_READ": 0}
    with db.read() as c:
        got = {r[0]: (r[1], r[2], r[3]) for r in c.execute(
            "SELECT sku_id, open_hours, open_hours_status, qty FROM demand_daily WHERE scope_key=?", (BQ_SCOPE,))}
    assert got == {"OK": (24.0, "OK", 5.0), "ZERO": (0.0, "OK", 1.0), "NULL": (None, "MISSING", 2.0),
                   "DIFF": (None, "CONFLICT", 3.0), "BIG": (None, "OUT_OF_RANGE", 4.0)}   # demand itself always written
    from app.sync import d1_dry_run                                          # the dry run says what a write would store
    agg = next(s for s in d1_dry_run()["steps"] if s["step"] == "aggregate")
    assert agg["open_hours_status_rows"] == {"OK": 2, "MISSING": 1, "OUT_OF_RANGE": 1, "CONFLICT": 1, "MULTI_ITEM": 0, "NOT_READ": 0}
    from app import data_checks
    chk = {x["check"]: x for x in data_checks.run(paths.db_path())["checks"]}
    for k in ("open_hours_status_values", "open_hours_value_matches_status", "open_hours_source_scope"):
        assert chk[k]["status"] == "OK", chk[k]
    assert chk["open_hours_window_ready"]["status"] == "WARN" and chk["open_hours_window_ready"]["count"] == 3


def test_full_period_validation_reports_missing_and_out_of_range_open_hours():
    _bq()
    connector.set_client_factory(FakeBQ(schema=SCHEMA))
    bq = ss.get("bigquery")
    s = {"n_rows": 100, "rows_valid": 100, "qty_read": 10, "qty_valid": 10, "ref_open_hours_non_null": 70,
         "ref_open_hours_negative": 1, "ref_open_hours_above_max": 2, "ref_open_hours_min": -1.0, "ref_open_hours_max": 30.0}
    sqls = connector._validation_sqls(bq, date(2026, 1, 1), date(2026, 1, 2))
    assert "COUNTIF(ref_open_hours > 24) AS ref_open_hours_above_max" in sqls["validate_summary"]
    v = connector._analyse(bq, date(2026, 1, 1), date(2026, 1, 2), s, [], [], 0, None, {"ok": True})
    w = {x["code"]: x["detail"] for x in v["warnings"]}
    assert w["OPEN_HOURS_MISSING"] == {"rows": 30, "share": 0.3}
    assert w["OPEN_HOURS_OUT_OF_RANGE"] == {"rows": 3, "min": -1.0, "max": 30.0}
    assert v["reference_completeness"]["open_hours"]["max"] == 30.0
    assert not any(b["code"].startswith("OPEN_HOURS") for b in v["blocking_issues"])   # the D-1 load is not blocked


# ================================================================== B. features: no look-ahead
def _panel(T=20, S=2):
    dates = pd.date_range("2026-09-01", periods=T, freq="D")
    Y = np.ones((S, T))
    hours = np.full((S, T), np.nan)
    hours[0, :] = np.arange(T, dtype=float)               # series 0: hours = day index (0..19, all within 0..24)
    hours[1, 10:] = 24.0                                  # series 1: unknown before day 10
    p = Panel([("A", "F"), ("B", "F")], dates, Y, np.full((S, T), 10.0), np.full((S, T), np.nan), np.zeros((S, T)),
              np.zeros(S, dtype=int), np.zeros(T, dtype=bool), {"A": 0, "B": 1}, {"F": 0}, hours=hours)
    return p


def test_open_hours_features_read_only_days_before_the_target_or_origin():
    p = _panel()
    cfg = ss.ModelSettings(feature_open_hours=True)
    names = feature_names(cfg)
    assert names[-2:] == OPEN_HOURS and open_hours_on(cfg)
    sc, fcc = static_codes(p)
    t = 12
    x = features_at(p, p.Y, p.price, p.disc, p.promo, p.dates, t, cfg, sc, fcc)
    last, mean7 = x[:, names.index("open_hours_last")], x[:, names.index("open_hours_mean_7")]
    assert last[0] == 11.0 and mean7[0] == pytest.approx(np.mean(np.arange(5, 12)))
    assert last[1] == 24.0 and mean7[1] == 24.0                                  # known days 10, 11 only
    p.hours[:, t:] = 99.0                                                       # the target day and later never read
    x2 = features_at(p, p.Y, p.price, p.disc, p.promo, p.dates, t, cfg, sc, fcc)
    assert np.array_equal(x, x2, equal_nan=True)
    x9 = features_at(p, p.Y, p.price, p.disc, p.promo, p.dates, 9, cfg, sc, fcc)
    assert np.isnan(x9[1, names.index("open_hours_last")]) and np.isnan(x9[1, names.index("open_hours_mean_7")])
    # DIRECT: features at the origin t0 (data < t0); RECURSIVE after the origin: nothing known, never filled
    p = _panel()
    d = ss.ModelSettings(feature_open_hours=True, forecast_strategy="DIRECT")
    dn = feature_names(d)
    xo = origin_features(p, p.Y, p.price, p.disc, p.promo, p.dates, 20, d, sc, fcc)
    assert dn.index("open_hours_last") < dn.index("horizon") and xo[0, dn.index("open_hours_last")] == 19.0
    S, h = p.Y.shape[0], 5                                  # RECURSIVE: arrays extended past the origin, as model.forecast
    ext = lambda a, fill: np.concatenate([a, np.full((S, h), fill)], axis=1)  # noqa: E731
    dates = p.dates.append(pd.date_range(p.dates[-1] + pd.Timedelta(days=1), periods=h, freq="D"))
    xf = features_at(p, ext(p.Y, np.nan), ext(p.price, 10.0), ext(p.disc, np.nan), ext(p.promo, 0.0), dates, 22, cfg, sc, fcc)
    assert np.isnan(xf[0, names.index("open_hours_last")])                     # day 21: after the origin, unknown
    assert xf[0, names.index("open_hours_mean_7")] == pytest.approx(np.mean([15, 16, 17, 18, 19]))   # known days only
    # truncation (backtest folds) keeps the hours of the kept days only
    assert p.truncate(8).hours.shape == (2, 8)


def test_a_panel_without_open_hours_is_never_passed_off_as_the_input():
    p = _panel()
    p.hours = None
    cfg = ss.ModelSettings(feature_open_hours=True)
    sc, fcc = static_codes(p)
    with pytest.raises(ValueError, match="OPEN_HOURS_NOT_IN_PANEL"):
        features_at(p, p.Y, p.price, p.disc, p.promo, p.dates, 12, cfg, sc, fcc)
    features_at(p, p.Y, p.price, p.disc, p.promo, p.dates, 12, ss.ModelSettings(), sc, fcc)     # input not used: fine


def test_data_checks_report_every_broken_open_hours_invariant():
    from app import data_checks
    with db.tx() as c:
        rows = [("BIGQUERY:p.d.t", "BIGQUERY", "2026-10-01", "A", 24.0, "OK"),          # valid
                ("BIGQUERY:p.d.t", "BIGQUERY", "2026-10-02", "B", None, "OK"),          # OK without a value
                ("BIGQUERY:p.d.t", "BIGQUERY", "2026-10-03", "C", 30.0, "OK"),          # outside 0..24
                ("BIGQUERY:p.d.t", "BIGQUERY", "2026-10-04", "D", 12.0, "MISSING"),     # a value without OK
                ("BIGQUERY:p.d.t", "BIGQUERY", "2026-10-05", "E", None, "GUESSED"),     # unknown status
                ("DEMO-X", "DEMO", "2026-10-05", "F", None, "OK")]                      # not a BigQuery row
        c.executemany("INSERT INTO demand_daily(scope_key, data_source, order_date, sku_id, fc_id, qty, n_orders, loaded_at, "
                      "open_hours, open_hours_status) VALUES (?,?,?,?, 'F1', 1, 1, 'x', ?, ?)", rows)
    chk = {x["check"]: x for x in data_checks.run(paths.db_path())["checks"]}
    assert chk["open_hours_status_values"]["status"] == "FAIL" and chk["open_hours_status_values"]["count"] == 1
    assert chk["open_hours_value_matches_status"]["status"] == "FAIL" and chk["open_hours_value_matches_status"]["count"] == 4
    assert chk["open_hours_source_scope"]["status"] == "FAIL" and chk["open_hours_source_scope"]["count"] == 1


def test_a_d1_load_between_the_check_and_the_read_never_trains_on_missing_hours(client, fast_model, monkeypatch):
    from app.ml import pipeline
    with monkeypatch.context() as m:
        _bq_scope_from_demo(m)
        with db.tx() as c:
            c.execute("UPDATE demand_daily SET open_hours=24, open_hours_status='OK' WHERE scope_key=?", (BQ_SCOPE,))
        real = pipeline.load_demand

        def racing_load(*a, **k):        # a D-1 reload made one day MISSING after the gate passed
            d = real(*a, **k)
            d.loc[d.index[0], "open_hours"] = None
            return d
        m.setattr(pipeline, "load_demand", racing_load)
        res = pipeline.train_and_publish(force=True, source="BIGQUERY")
    assert res["status"] == "FAILED" and res["reason_code"] == "OPEN_HOURS_INVALID" and res["failed_stage"] == "PREPARING_DATA"
    with db.read() as c:
        tr = db.one(c, "SELECT status, reject_reason FROM training_runs WHERE id=?", (res["training_run_id"],))
        assert c.execute("SELECT COUNT(*) FROM forecast_runs WHERE scope_key=?", (BQ_SCOPE,)).fetchone()[0] == 0
    assert (tr["status"], tr["reject_reason"]) == ("FAILED", "OPEN_HOURS_INVALID")


def test_panel_takes_open_hours_from_the_demand_rows_only():
    d = pd.DataFrame({"order_date": ["2026-09-01", "2026-09-02", "2026-09-03"], "sku_id": ["A"] * 3, "fc_id": ["F"] * 3,
                      "qty": [1.0, 2.0, 3.0], "open_hours": [24.0, None, 0.0]})
    p = build_panel(d)
    assert p.hours[0, 0] == 24.0 and np.isnan(p.hours[0, 1]) and p.hours[0, 2] == 0.0      # 0 stays 0, NULL stays NaN
    assert np.isnan(build_panel(d.drop(columns=["open_hours"])).hours).all()


def test_models_before_1_2_9_keep_their_feature_list():
    for cfg in (ss.ModelSettings(), ss.ModelSettings(forecast_strategy="DIRECT"),
                ss.ModelSettings.model_validate({k: v for k, v in ss.ModelSettings().model_dump().items()
                                                 if k not in ("feature_open_hours", "weekday_encoding")})):
        assert cfg.feature_open_hours is None and not open_hours_on(cfg)
        assert not set(OPEN_HOURS) & set(feature_names(cfg))


# ================================================================== B. training gate + feature matrix (BIGQUERY scope)
def _bq_scope_from_demo(monkeypatch):
    from app.data.demo import create_demo_batch
    demo = create_demo_batch(days=150, n_skus=8, seed=3)
    _bq()
    with db.tx() as c:
        c.execute("INSERT INTO demand_daily(scope_key, data_source, order_date, sku_id, fc_id, qty, n_orders, avg_unit_price, "
                  "avg_discount_rate, promo_share, voucher_share, loaded_at) SELECT ?, 'BIGQUERY', order_date, sku_id, fc_id, qty, "
                  "n_orders, avg_unit_price, avg_discount_rate, promo_share, voucher_share, loaded_at FROM demand_daily "
                  "WHERE scope_key=?", (BQ_SCOPE, demo["batch_id"]))
    from app.jobs import flows
    monkeypatch.setattr(flows, "_bq_ready", lambda c, blockers: None)      # production state is not under test here


def test_training_on_bigquery_stops_until_every_row_of_the_window_has_valid_open_hours(client, fast_model, monkeypatch):
    from app.ml import pipeline
    with monkeypatch.context() as m:
        _bq_scope_from_demo(m)
        with pytest.raises(AppError) as e:                                     # loaded before 1.2.9: NOT_LOADED
            pipeline.train_and_publish(force=True, source="BIGQUERY")
        assert e.value.code == "OPEN_HOURS_INVALID"
        det = json.loads(e.value.detail)
        assert det["ok_rows"] == 0 and det["not_ok"]["NOT_LOADED"]["rows"] == det["rows"] > 0
        with db.read() as c:
            assert c.execute("SELECT COUNT(*) FROM training_runs WHERE scope_key=?", (BQ_SCOPE,)).fetchone()[0] == 0
        with db.tx() as c:                                                     # a D-1 full reload with the column
            c.execute("UPDATE demand_daily SET open_hours = 8 + (julianday(order_date) % 17), open_hours_status='OK' "
                      "WHERE scope_key=?", (BQ_SCOPE,))
            day = c.execute("SELECT MAX(order_date) FROM demand_daily WHERE scope_key=?", (BQ_SCOPE,)).fetchone()[0]
            c.execute("UPDATE demand_daily SET open_hours=NULL, open_hours_status='CONFLICT' WHERE scope_key=? AND "
                      "order_date=? AND rowid=(SELECT MIN(rowid) FROM demand_daily WHERE scope_key=? AND order_date=?)",
                      (BQ_SCOPE, day, BQ_SCOPE, day))
        with pytest.raises(AppError) as e:
            pipeline.train_and_publish(force=True, source="BIGQUERY")
        det = json.loads(e.value.detail)
        assert det["not_ok"] == {"CONFLICT": {"rows": 1, "first": day, "last": day}}
        with db.tx() as c:
            c.execute("UPDATE demand_daily SET open_hours=24, open_hours_status='OK' WHERE scope_key=? AND "
                      "open_hours_status='CONFLICT'", (BQ_SCOPE,))
        res = pipeline.train_and_publish(force=True, source="BIGQUERY")
        assert res["status"] == "PUBLISHED", res
        with db.read() as c:
            tr = db.one(c, "SELECT * FROM training_runs WHERE id=?", (res["training_run_id"],))
            hours = {(r[0], r[1], r[2]): r[3] for r in c.execute(
                "SELECT sku_id, fc_id, order_date, open_hours FROM demand_daily WHERE scope_key=?", (BQ_SCOPE,))}
        fc = json.loads(tr["feature_config_json"])
        assert fc["config"]["feature_open_hours"] is True and set(OPEN_HOURS) <= set(fc["features"])
        assert fc["feature_groups"]["open_hours"] == {"mandatory": True, "used": True}
        assert fc["feature_decisions"]["open_hours"]["window_check"]["ok"]
        assert fc["data_identity"]["columns"] == list(identity.OPEN_HOURS_COLUMNS)
        import lightgbm as lgb
        booster = lgb.Booster(model_file=str(paths.models_dir() / tr["model_path"]))
        assert booster.feature_name() == fc["features"]                        # train and inference: one schema
        imp = dict(zip(booster.feature_name(), booster.feature_importance("split")))
        assert "open_hours_last" in imp                                         # a real column of the matrix
        with db.read() as c:
            run = db.one(c, "SELECT * FROM forecast_runs WHERE id=?", (res["forecast_run_id"],))
        vals = client.get(f"/api/forecast-runs/{run['id']}/values?limit=5000").json()["rows"]
        x = next(v for v in vals if v["method"] == "LIGHTGBM")
        ex = client.get(f"/api/forecast-runs/{run['id']}/explain",
                        params={"sku_id": x["sku_id"], "fc_id": x["fc_id"], "target_date": x["target_date"]}).json()
        assert ex["reproduction"]["status"] == "OK" and ex["lightgbm"]["matches_published"]
        feats = {f["name"]: f["value"] for f in ex["lightgbm"]["features"]}
        origin_prev = str(date.fromisoformat(run["forecast_origin"]) - timedelta(days=1))
        assert feats["open_hours_last"] == hours.get((x["sku_id"], x["fc_id"], origin_prev))   # value of the day BEFORE


def test_registry_shows_open_hours_mandatory_for_bigquery_and_not_applicable_elsewhere(client, fast_model, monkeypatch):
    client.post("/api/demo/generate", json={"days": 90, "n_skus": 4, "seed": 1})
    demo = {e["feature_key"]: e for e in client.get("/api/features/registry").json()["entries"]}
    assert demo["open_hours_last"]["availability_status"] == "UNAVAILABLE" and demo["open_hours_last"]["reason"] == "NOT_IN_SOURCE"
    assert demo["open_hours_last"]["required"] is False and demo["open_hours_last"]["kind"] == "MODEL"
    assert "open_hours" not in {e["feature_key"] for e in demo.values() if e["kind"] == "PLANNED"}
    with monkeypatch.context() as m:
        _bq_scope_from_demo(m)
        v = client.get("/api/features/registry", params={"source": "BIGQUERY"}).json()
        reg = {e["feature_key"]: e for e in v["entries"]}
        assert reg["open_hours_last"]["availability_status"] == "UNAVAILABLE"
        assert reg["open_hours_last"]["reason"] == "MANDATORY_NOT_READY" and reg["open_hours_last"]["required"] is True
        assert v["next_training"]["blocked_by"] == "OPEN_HOURS_INVALID" and "open_hours_last" in v["next_training"]["features"]
        with db.tx() as c:
            c.execute("UPDATE demand_daily SET open_hours=24, open_hours_status='OK' WHERE scope_key=?", (BQ_SCOPE,))
        v = client.get("/api/features/registry", params={"source": "BIGQUERY"}).json()
        assert {e["feature_key"]: e for e in v["entries"]}["open_hours_last"]["reason"] == "MANDATORY_DATA_OK"
        assert v["next_training"]["blocked_by"] is None


def test_feature_catalogue_includes_the_open_hours_only_for_the_bigquery_source(client):
    names = [f["name"] for f in client.get("/api/features/catalog").json()["features"]]
    assert not set(OPEN_HOURS) & set(names)                                   # DEMO (default source)
    ss.update("general", {"active_source": "BIGQUERY"})
    cat = {f["name"]: f for f in client.get("/api/features/catalog").json()["features"]}
    assert set(OPEN_HOURS) <= set(cat) and cat["open_hours_last"]["group"] == "AVAILABILITY"
    assert "0–24" in cat["open_hours_last"]["vi"]["formula"] and "không thay bằng 0" in cat["open_hours_last"]["vi"]["missing"]
    assert "1 = Thứ 2" in cat["dow"]["vi"]["formula"]


def test_an_identity_recorded_before_1_2_9_survives_a_reload_that_adds_open_hours(client, fast_model):
    assert client.post("/api/demo/generate", json={"days": 150, "n_skus": 8, "seed": 3}).status_code == 200
    assert client.post("/api/training/run", json={}).json()["status"] == "SUCCEEDED"
    run = client.get("/api/system/status").json()["latest_forecast_run"]
    with db.read() as c:
        tr = db.one(c, "SELECT * FROM training_runs WHERE id=?", (run["training_run_id"],))
    with db.tx() as c:          # same values rewritten (new loaded_at) + open hours added: what a D-1 full reload does
        c.execute("UPDATE demand_daily SET open_hours=24, open_hours_status='OK', loaded_at='2099-01-01T00:00:00+00:00' "
                  "WHERE scope_key=?", (run["scope_key"],))
    with db.read() as c:
        st = identity.data_state(tr, run["scope_key"], c)
    assert st["status"] == "SAME_CONTENT", st
    with db.read() as c:
        assert identity.window_identity(run["scope_key"], "2000-01-01", "2100-01-01", c,
                                        columns=identity.OPEN_HOURS_COLUMNS)["columns"][-1] == "open_hours"
        with pytest.raises(ValueError):
            identity.window_identity(run["scope_key"], "2000-01-01", "2100-01-01", c, columns=["qty; DROP TABLE x"])


# ================================================================== B. migration 0007 (additive) on a database copy
def test_migration_0007_is_additive_and_reversible_on_a_copy(tmp_path):
    src = tmp_path / "v6.db"
    c = sqlite3.connect(src)
    c.execute("CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY, name TEXT NOT NULL, applied_at TEXT NOT NULL)")
    for v, p in db.available_migrations():
        if v > 6:
            continue
        for stmt in db._split_sql(p.read_text(encoding="utf-8")):
            c.execute(stmt)
        c.execute("INSERT INTO schema_migrations VALUES (?,?,?)", (v, p.name, "2026-10-19T00:00:00+00:00"))
    c.executemany("INSERT INTO demand_daily(scope_key, data_source, order_date, sku_id, fc_id, qty, n_orders, avg_unit_price, "
                  "loaded_at) VALUES (?,?,?,?,?,?,?,?,?)",
                  [("BIGQUERY:p.d.t", "BIGQUERY", f"2026-10-{d:02d}", f"S{s}", "F1", float(d * s), 1, 9.5, "2026-10-19")
                   for d in range(1, 11) for s in range(1, 4)])
    c.commit()
    before = c.execute("SELECT * FROM demand_daily ORDER BY rowid").fetchall()
    cols6 = [r[1] for r in c.execute("PRAGMA table_info(demand_daily)")]
    c.close()
    res = db.migrate(src)
    # 1.2.9b: 0008 (customer counts, additive) follows 0007 in the same upgrade
    assert res["applied"] == ["0007_open_hours.sql", "0008_customer_aggregates.sql",
                              "0009_actual_finalization.sql"] and res["to"] == 9      # 0009: seventh batch
    c = sqlite3.connect(src)
    cols7 = [r[1] for r in c.execute("PRAGMA table_info(demand_daily)")]
    assert cols7 == cols6 + ["open_hours", "open_hours_status", "customer_count", "repeat_customer_count",
                             "repeat_interval_days_sum"]
    after = c.execute(f"SELECT {', '.join(cols6)} FROM demand_daily ORDER BY rowid").fetchall()
    assert after == before                                                   # every existing value unchanged
    assert c.execute("SELECT COUNT(*) FROM demand_daily WHERE open_hours IS NOT NULL OR open_hours_status IS NOT NULL").fetchone()[0] == 0
    assert c.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    # the reverse statements written in the migration headers (0009, then 0008) bring a COPY back to the schema-6 layout
    from test_v129g_actual_final import reverse_0009
    reverse_0009(c)
    for col in ("repeat_interval_days_sum", "repeat_customer_count", "customer_count"):
        c.execute(f"ALTER TABLE demand_daily DROP COLUMN {col}")
    c.execute("DELETE FROM schema_migrations WHERE version = 8")
    c.execute("ALTER TABLE demand_daily DROP COLUMN open_hours_status")
    c.execute("ALTER TABLE demand_daily DROP COLUMN open_hours")
    c.execute("DELETE FROM schema_migrations WHERE version = 7")
    c.commit()
    assert [r[1] for r in c.execute("PRAGMA table_info(demand_daily)")] == cols6
    assert c.execute("SELECT * FROM demand_daily ORDER BY rowid").fetchall() == before
    c.close()


# ================================================================== C. weekday names in the daily forecast table
def test_values_carry_the_weekday_of_the_target_date_and_export_names_it(client, fast_model):
    assert client.post("/api/demo/generate", json={"days": 150, "n_skus": 8, "seed": 3}).status_code == 200
    assert client.post("/api/training/run", json={}).json()["status"] == "SUCCEEDED"
    run = client.get("/api/system/status").json()["latest_forecast_run"]
    j = client.get(f"/api/forecast-runs/{run['id']}/values?limit=5000&horizon=7").json()
    assert j["weekday_rule"] == "ISO_1_MONDAY_7_SUNDAY"
    rows = j["rows"]
    assert all(r["weekday"] == wd.iso_weekday(r["target_date"]) for r in rows)
    assert {r["weekday"] for r in rows} == set(range(1, 8))                 # 7 consecutive days: every weekday once
    first = min(r["target_date"] for r in rows)
    assert {r["horizon_day"] for r in rows if r["target_date"] == first} == {1}      # horizon_day unchanged
    x = client.get(f"/api/forecast-runs/{run['id']}/export.xlsx", params={"horizon": 7, "lang": "vi"})
    wb = load_workbook(BytesIO(x.content))
    info = {r[0]: r[1] for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)}
    assert info["Phiên bản định dạng file"] == "6"
    sh = wb["Chi tiết theo ngày"]
    head = [c.value for c in sh[1]]
    assert head[head.index("Ngày") + 1] == "Thứ trong tuần" and "Ngày thứ" in head
    for r in sh.iter_rows(min_row=2, values_only=True):
        d = r[head.index("Ngày")]
        d = d.date() if isinstance(d, datetime) else d
        assert r[head.index("Thứ trong tuần")] == wd.label(wd.iso_weekday(d))
        assert isinstance(r[head.index("Ngày thứ")], int)
