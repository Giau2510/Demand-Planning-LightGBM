"""v1.1.2: display formats (DD/MM/YYYY everywhere in exports), dated explanation history, method names in Excel."""
import hashlib
import io
import logging
from datetime import datetime

import pytest
from openpyxl import load_workbook

from app import db
from app import settings_service as ss
from app import timeutil
from app.ml import methods, model
from app.timeutil import VN_TZ


def _job(client, jid):
    return client.get(f"/api/jobs/{jid}").json()


@pytest.fixture
def demo_run(client, fast_model):
    ss.update("horizons", {"enabled": [7]})
    assert client.post("/api/demo/generate", json={"days": 150, "n_skus": 8, "seed": 3}).status_code == 200
    j = _job(client, client.post("/api/training/run", json={"force": True}).json()["id"])
    assert j["reason_code"] == "FORECAST_PUBLISHED", j
    return j["result"]["result"]["forecast_run_id"]


def _xlsx(client, path):
    r = client.get(path)
    assert r.status_code == 200
    return load_workbook(io.BytesIO(r.content))


def _db_fingerprint(run_id):
    with db.read() as c:
        rows = c.execute("SELECT sku_id, fc_id, target_date, horizon_day, yhat, method FROM forecast_values WHERE run_id=? "
                         "ORDER BY sku_id, fc_id, target_date", (run_id,)).fetchall()
    return hashlib.sha256(repr([tuple(r) for r in rows]).encode()).hexdigest()


# ------------------------------------------------------------------ D. method names in Excel
def test_every_method_code_has_a_name_in_both_languages_and_unknown_codes_are_safe(caplog):
    codes = {v for k, v in vars(model).items() if k.startswith(("METHOD_", "FB_")) and isinstance(v, str)}
    for code in codes:
        vi, en = methods.label(code, "vi"), methods.label(code, "en")
        assert vi and en and vi != code and en != code                      # never a bare internal code
    assert methods.label("LIGHTGBM", "vi") == "Mô hình học máy (LightGBM)"
    assert methods.label("FALLBACK_NEW_SERIES", "vi") == "Dự phòng cho chuỗi SKU–FC mới"
    with caplog.at_level(logging.WARNING, logger="df.methods"):
        assert methods.label("FALLBACK_SOMETHING_NEW", "vi") == "Phương pháp chưa có mô tả (FALLBACK_SOMETHING_NEW)"
        assert methods.label("FALLBACK_SOMETHING_NEW", "en") == "Undocumented method (FALLBACK_SOMETHING_NEW)"
    assert any("method code without label" in r.getMessage() for r in caplog.records)


def test_excel_method_names_follow_the_language_and_export_changes_nothing(client, demo_run):
    before = _db_fingerprint(demo_run)
    vi = _xlsx(client, f"/api/forecast-runs/{demo_run}/export.xlsx?horizon=7&lang=vi")
    en = _xlsx(client, f"/api/forecast-runs/{demo_run}/export.xlsx?horizon=7&lang=en")
    assert _db_fingerprint(demo_run) == before                               # export never writes
    with db.read() as c:
        used = {r[0] for r in c.execute("SELECT DISTINCT method FROM forecast_values WHERE run_id=?", (demo_run,))}
    assert "LIGHTGBM" in used and len(used) >= 2
    for wb, lang in ((vi, "vi"), (en, "en")):
        for ws, name_col, code_col in ((wb.worksheets[1], 2, 5), (wb.worksheets[2], 6, 7)):
            head = [c.value for c in ws[1]]
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            assert {r[code_col] for r in rows} == used                    # code column: original codes, unchanged
            for r in rows:
                assert r[name_col] == methods.label(r[code_col], lang)    # name column: the language's name
                assert r[name_col] != r[code_col]
            assert head[name_col] == ("Phương pháp" if lang == "vi" else "Method")
            assert "không dịch" in head[code_col] if lang == "vi" else "not translated" in head[code_col]


# ------------------------------------------------------------------ B. dates DD/MM/YYYY in every export
def test_forecast_export_dates_are_dd_mm_yyyy_in_both_languages(client, demo_run):
    run = client.get(f"/api/forecast-runs/{demo_run}").json()
    for lang in ("vi", "en"):
        wb = _xlsx(client, f"/api/forecast-runs/{demo_run}/export.xlsx?horizon=7&lang={lang}")
        daily = wb.worksheets[2]
        assert all(daily.cell(row=r, column=3).number_format == "DD/MM/YYYY" for r in range(2, daily.max_row + 1))
        info = [r[1] for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True)]
        text = " ".join(str(v) for v in info)
        o = datetime.fromisoformat(run["forecast_origin"]).strftime("%d/%m/%Y")
        assert o in text
        # no ISO calendar date anywhere except the explicit UTC technical fields
        iso_like = [v for v in info if isinstance(v, str) and len(v) >= 10 and v[4] == "-" and v[7] == "-"]
        assert all("T" in v for v in iso_like), iso_like                  # only the *_utc timestamps stay ISO


def test_snapshot_export_and_error_report_follow_the_language_with_dd_mm_yyyy(client, demo_run):
    now = datetime.now(VN_TZ).replace(hour=10, minute=5, second=0, microsecond=0)
    timeutil.set_now_override(lambda: now)
    client.post("/api/jobs/run", json={"job_type": "D0_SYNC"})
    sid = client.get("/api/snapshots").json()[0]["id"]
    vi = _xlsx(client, f"/api/snapshots/{sid}/export.xlsx?lang=vi")
    en = _xlsx(client, f"/api/snapshots/{sid}/export.xlsx?lang=en")
    assert vi.sheetnames == ["Thông tin snapshot", "Số lượng lũy kế"] and en.sheetnames == ["Snapshot info", "Cumulative quantities"]
    iv = {r[0]: r[1] for r in vi.worksheets[0].iter_rows(min_row=2, values_only=True)}
    ie = {r[0]: r[1] for r in en.worksheets[0].iter_rows(min_row=2, values_only=True)}
    assert iv["Ngày đặt hàng"] == now.strftime("%d/%m/%Y") == ie["Order date"]
    assert iv["Thời điểm đọc nguồn (giờ Việt Nam)"] == now.strftime("%d/%m/%Y %H:%M") + " (giờ Việt Nam, UTC+7)"
    assert ie["Source read at (Vietnam time)"] == now.strftime("%d/%m/%Y %H:%M") + " (Vietnam time, UTC+7)"
    with db.read() as c:
        stored = {(r[0], r[1]): r[2] for r in c.execute("SELECT sku_id, fc_id, cumulative_qty FROM d0_snapshot_values WHERE snapshot_id=?", (sid,))}
    vals = list(vi.worksheets[1].iter_rows(min_row=2, values_only=True))
    assert [c.value for c in vi.worksheets[1][1]][:3] == ["SKU", "FC/kho", "Số lượng đặt lũy kế đến lúc đọc"]
    assert {(r[0], r[1]): r[2] for r in vals} == stored                    # numbers unchanged


def test_error_report_formats_valid_dates_and_keeps_unparseable_raw_values():
    from app.data import excel_io
    batch = {"id": "IMP-1", "file_name": "x.xlsx", "sheet_name": "Data", "status": "PENDING_REVIEW", "total_rows": 3,
             "valid_rows": 1, "error_rows": 2, "duplicate_rows": 0, "min_order_date": "2026-10-01",
             "max_order_date": "2026-10-13", "created_at": "2026-10-14T03:00:00+00:00"}
    issues = [{"row_number": 2, "field": "created_at", "code": "INVALID_DATETIME", "severity": "ERROR", "raw_value": "14/13/2026"},
              {"row_number": 3, "field": "order_date", "code": "FUTURE_DATE", "severity": "ERROR", "raw_value": "2027-01-05"},
              {"row_number": 4, "field": "sku_id", "code": "MISSING_SKU", "severity": "ERROR", "raw_value": "2026-10-14"}]
    for lang, sheets in (("vi", ("Tóm tắt", "Lỗi")), ("en", ("Summary", "Errors"))):
        wb = load_workbook(io.BytesIO(excel_io.build_error_report(batch, issues, lang)))
        assert wb.sheetnames == list(sheets)
        summ = {r[0]: r[1] for r in wb.worksheets[0].iter_rows(values_only=True)}
        assert "01/10/2026" in summ.values() and "13/10/2026" in summ.values()
        assert any(str(v).startswith("14/10/2026 10:00") for v in summ.values())          # 03:00Z = 10:00 Vietnam
        raws = [r[4] for r in wb.worksheets[1].iter_rows(min_row=2, values_only=True)]
        assert raws == ["14/13/2026", "05/01/2027", "2026-10-14"]                       # bad input kept; non-date field untouched


# ------------------------------------------------------------------ C. explanation history: date <-> demand pairing
def test_explanation_history_is_dated_chronological_and_adds_up_to_the_calculation(client, demo_run):
    run = client.get(f"/api/forecast-runs/{demo_run}").json()
    with db.read() as c:
        series = c.execute("SELECT sku_id, fc_id, method FROM forecast_values WHERE run_id=? AND horizon_day=1", (demo_run,)).fetchall()
        demand = {(r[0], r[1], r[2]): r[3] for r in c.execute("SELECT sku_id, fc_id, order_date, qty FROM demand_daily")}
    checked = set()
    for sku, fc, method in series:
        if method in checked:
            continue
        e = client.get(f"/api/forecast-runs/{demo_run}/explain", params={"sku_id": sku, "fc_id": fc, "target_date": run["forecast_origin"]}).json()
        h = e["history"]
        dates = [d["date"] for d in h["days"]]
        assert dates == sorted(dates) and len(set(dates)) == len(dates)          # chronological, one row per day
        assert dates[-1] == run["data_as_of"] and all(d <= run["data_as_of"] for d in dates)   # no future
        assert h["window_days"] >= max(56, ss.get("model").inactive_days)      # covers every rule window
        for d in h["days"]:                                                    # each demand next to ITS date
            stored = demand.get((sku, fc, d["date"]))
            if d["status"] in ("ORDERS", "ZERO_ORDERS"):
                assert d["qty"] == (stored or 0.0)
            else:
                assert d["qty"] is None                                        # unknown is not 0
        if method != "LIGHTGBM":
            used = [d for d in h["days"] if d["in_calc"]]
            assert h["calc"]["known_days"] == len(used) and h["calc"]["sum"] == pytest.approx(sum(d["qty"] for d in used))
            assert e["fallback"]["matches_history"] is True
            if "known_days" in e["fallback"]:
                assert e["fallback"]["known_days"] == len(used)
                assert e["fallback"]["sum"] == pytest.approx(sum(d["qty"] for d in used))
                assert round(e["fallback"]["sum"] / len(used), 4) == pytest.approx(e["result"]["yhat"])
        checked.add(method)
    assert "LIGHTGBM" in checked and len(checked) >= 2
