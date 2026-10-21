"""BigQuery history ingest (source after 1.1.8, 16/10/2026): exact-duplicate gate BEFORE aggregation, downloadable
duplicate report, aggregation to day x SKU x FC, safe re-runs, dry run.

Fake BigQuery client (canned query results): these tests cover the app logic (order of the steps, blocking, report
content, writes, transactions). The SQL semantics (GROUP BY on every mapped column, NULL handling, sums) are checked on
REAL BigQuery with inline data in test_bq_real_semantics.py (opt-in)."""
import json
from datetime import date, datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from app import db, sync, timeutil
from app import settings_service as ss
from app.bq import connector, dup_report
from app.data.demand import bigquery_scope_key
from app.messages import AppError


def _ready():
    from test_v116 import _bq_ready, _hist_production      # tests/ is on sys.path
    fake = _bq_ready()
    _hist_production(fake)
    return fake


def _demand(sk=None):
    sk = sk or bigquery_scope_key(ss.get("bigquery"))
    with db.read() as c:
        return tuple(c.execute("SELECT COUNT(*), COALESCE(SUM(qty), 0) FROM demand_daily WHERE scope_key=?",
                               (sk,)).fetchone())


def _group(n, **vals):
    """One duplicated group as BigQuery returns it: every compared column (alias) + its number of occurrences."""
    g = {c["alias"]: None for c in connector.dup_columns(ss.get("bigquery"))}
    g.update(vals)
    g["_occurrences"] = n
    return g


def _dup(groups, rows_checked=20, qty=100.0):
    return [{"rows_checked": rows_checked, "qty_checked": qty, "duplicate_groups": len(groups),
             "duplicate_rows": sum(g["_occurrences"] for g in groups),
             "extra_duplicate_rows": sum(g["_occurrences"] - 1 for g in groups), "_groups": groups}]


def _rows(day, *items):
    """Aggregated day x SKU x FC rows (what the BigQuery aggregation returns) + the totals rows."""
    rows = [{"_kind": "VALID", "order_date": day, "sku_id": s, "fc_id": f, "qty": q, "n_orders": n, "n_rows": n}
            for s, f, q, n in items]
    n, q = sum(i[3] for i in items), sum(i[2] for i in items)
    return rows + [{"_kind": "TOTAL_READ", "_reason": "0", "qty": q, "n_rows": n},
                   {"_kind": "TOTAL_KEPT", "_reason": "0", "qty": q, "n_rows": n}]


def _real_runs(fake, kind):
    return [i for i, (sql, cfg) in enumerate(fake.queries) if sql.startswith(f"-- df:{kind}\n") and not cfg.get("dry_run")]


TWO_GROUPS = lambda: [  # noqa: E731
    _group(2, order_id="SO-1", order_date=date(2026, 10, 1), sku_id="S1", fc_id="F1", order_qty=20.0,
           source_updated_at=datetime(2026, 10, 2, 1, 0, tzinfo=timezone.utc), ref_sales_price=12.5,
           ref_outbound_qty=20.0, ref_order_status="COMPLETED"),
    _group(3, order_id=None, order_date=date(2026, 10, 1), sku_id="S2", fc_id="F1", order_qty=3.0,
           ref_order_status="WAIT_TO_CONFIRM")]


# ------------------------------------------------------------------ 5, 6, 7, 8: duplicates block, full report
def test_duplicates_block_the_load_before_aggregation_and_the_report_lists_every_row(client):
    from app.jobs import runner
    fake = _ready()
    y = timeutil.today_vn() - timedelta(days=1)
    fake.responses["dup_check"] = _dup([])
    fake.responses["daily"] = _rows(y, ("S1", "F1", 4.0, 1))
    sync.run_d1_sync("J1", fallback_mode=True)
    assert _demand() == (1, 4.0)
    daily_runs = len(_real_runs(fake, "daily"))
    fake.responses["dup_check"] = _dup(TWO_GROUPS())
    fake.responses["daily"] = _rows(y, ("S1", "F1", 99.0, 9))
    j = runner.execute("D1_SYNC", "D1_SYNC:dups", "MANUAL", None, {"fallback_mode": True})
    assert j["status"] == "FAILED" and j["reason_code"] == "BQ_DATA_DUPLICATES"         # no retry: data problem
    assert "groups=2 rows=5 extra=3" in j["error"] and "bq-trung-lap_" in j["error"]
    assert len(_real_runs(fake, "daily")) == daily_runs                                   # aggregation never ran
    assert _demand() == (1, 4.0)                                                          # nothing written
    st = connector.dup_state()
    assert st["ingest_status"] == "BLOCKED" and st["duplicate_groups"] == 2 and st["extra_duplicate_rows"] == 3
    assert st["context"] == f"D-1 job {j['id']}" and "groups" not in st                  # rows never kept in the DB
    name = st["report"]["file"]
    wb = load_workbook(dup_report.folder() / name)
    assert wb.sheetnames == ["Summary", "Duplicate_Rows", "Duplicate_Groups"]
    rows = list(wb["Duplicate_Rows"].values)
    head = rows[0]
    assert head[:2] == ("duplicate_group_id", "duplicate_occurrence")
    assert {"so_code", "created_date", "sku_code", "fc_code", "demand_quantity", "sales_price",
            "outbound_quantity", "order_status"} <= set(head)
    body = [dict(zip(head, r)) for r in rows[1:]]
    assert [(r["duplicate_group_id"], r["duplicate_occurrence"]) for r in body] == [
        ("DUP-00001", 1), ("DUP-00001", 2), ("DUP-00002", 1), ("DUP-00002", 2), ("DUP-00002", 3)]
    assert all(r["so_code"] is None and r["demand_quantity"] == 3 for r in body[2:])      # NULL so_code kept as is
    assert body[0]["refresh_time"] == "2026-10-02T01:00:00+00:00"                        # TIMESTAMP kept exactly
    groups = [dict(zip(list(wb["Duplicate_Groups"].values)[0], r)) for r in list(wb["Duplicate_Groups"].values)[1:]]
    assert [(g["duplicate_group_id"], g["occurrences"]) for g in groups] == [("DUP-00001", 2), ("DUP-00002", 3)]
    summary = {r[0]: r[1] for r in list(wb["Summary"].values)[1:]}
    assert summary["Trạng thái ingest / Ingest status"] == "BLOCKED"
    assert summary["Số nhóm trùng hoàn toàn / duplicate_groups"] == 2
    assert summary["Tổng số dòng thuộc các nhóm trùng / duplicate_rows"] == 5
    assert summary["Số dòng dư (mỗi nhóm giữ 1) / extra_duplicate_rows"] == 3
    assert summary["Số dòng nguồn đã kiểm tra / Source rows checked"] == 20
    # downloadable from the BigQuery page; nothing else is served
    lst = client.get("/api/bigquery/duplicate-reports").json()
    assert lst[0]["name"] == name
    r = client.get(f"/api/bigquery/duplicate-reports/{name}")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert client.get("/api/bigquery/duplicate-reports/..%5C..%5Cdb%5Capp.db").status_code == 404
    assert client.get("/api/bigquery/duplicate-reports/app.db").status_code == 404
    st = client.get("/api/bigquery/status").json()
    assert st["duplicate_check"]["report"]["file"] == name and st["duplicate_reports"][0]["name"] == name


def test_reports_are_never_overwritten():
    fake = _ready()
    timeutil.set_now_override(lambda: datetime(2026, 10, 16, 6, 0, 0, tzinfo=timezone.utc))
    fake.responses["dup_check"] = _dup(TWO_GROUPS())
    names = []
    for _ in range(3):
        chk = connector.duplicate_check(date(2026, 1, 1), date(2026, 10, 15), context="test")
        names.append(chk["report"]["file"])
    assert len(set(names)) == 3 and names[0] == "bq-trung-lap_2026-01-01_2026-10-15_20261016-130000.xlsx"
    assert names[1].endswith("130000-2.xlsx") and names[2].endswith("130000-3.xlsx")


# ------------------------------------------------------------------ 13: report failure still blocks
def test_xlsx_failure_still_blocks_and_writes_a_csv_fallback(monkeypatch):
    fake = _ready()
    before = _demand()
    fake.responses["dup_check"] = _dup(TWO_GROUPS())
    monkeypatch.setattr(dup_report, "_xlsx_bytes", lambda sheets: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(AppError) as e:
        connector.fetch_daily(date(2026, 1, 1), date(2026, 1, 31))
    assert e.value.code == "BQ_DATA_DUPLICATES" and "xlsx failed (OSError: disk full)" in e.value.detail
    rep = e.value.params["duplicate_check"]["report"]
    assert rep["format"] == "csv" and len(rep["files"]) == 3
    assert all((dup_report.folder() / f).is_file() for f in rep["files"])
    assert not list(dup_report.folder().glob("*.xlsx"))                                  # no half-written workbook
    assert "Duplicate_Rows" in rep["files"][1] and _demand() == before
    # nothing can be written at all: still blocked, the failure is in the message
    monkeypatch.setattr(dup_report.pd.DataFrame, "to_csv", lambda *a, **k: (_ for _ in ()).throw(OSError("read-only")))
    with pytest.raises(AppError) as e:
        connector.fetch_daily(date(2026, 1, 1), date(2026, 1, 31))
    assert e.value.code == "BQ_DATA_DUPLICATES" and "REPORT FAILED" in e.value.detail and _demand() == before


# ------------------------------------------------------------------ 14: mapping must match the live source
def test_a_mapped_column_missing_from_the_source_stops_the_run_with_the_field_list():
    fake = _ready()
    fake.schema = [(n, t) for n, t in fake.schema if n != "sales_price"]
    n_queries = len(fake.queries)
    with pytest.raises(AppError) as e:
        connector.fetch_daily(date(2026, 1, 1), date(2026, 1, 31))
    assert e.value.code == "BQ_MAPPING_INVALID"
    err = json.loads(e.value.detail)
    assert {"field": "sales_price", "code": "COLUMN_NOT_FOUND", "column": "sales_price", "reference": True} in err
    assert len(fake.queries) == n_queries                                                 # no data query at all
    ss.update("bigquery", {"field_mapping": {**{k: v.model_dump() for k, v in ss.get("bigquery").field_mapping.items()},
                                             "sku_id": {"column": None, "usage": "REQUIRED"}}})
    with pytest.raises(AppError) as e:
        connector.fetch_daily(date(2026, 1, 1), date(2026, 1, 31))
    assert {"field": "sku_id", "code": "REQUIRED_NOT_MAPPED"} in json.loads(e.value.detail)


# ------------------------------------------------------------------ 1-4, 7, 8: what is compared, nothing filtered
def test_duplicate_query_compares_every_mapped_column_and_filters_no_row():
    _ready()
    bq = ss.get("bigquery")
    cols = connector.dup_columns(bq)
    assert {c["column"] for c in cols} >= {"so_code", "created_date", "sku_code", "fc_code", "demand_quantity",
                                           "sales_price", "outbound_quantity", "order_status", "seller_code",
                                           "product_id"}
    assert "item_code" not in {c["column"] for c in cols}        # 1.2.11: never a uniqueness criterion
    sql = connector._dup_sql(bq, {"mode": "TABLE", "key_columns": [], "key_duplicate_policy": "SUM_ALL"})
    group_by = sql.split("GROUP BY ", 1)[1].split("\n", 1)[0]
    # native values, every compared column (1.2.11: pending orders are exempt through _pending, not through the key)
    assert [x.strip() for x in group_by.split(",")] == [c["alias"] for c in cols]
    where = sql.split(" WHERE ", 1)[1].split("\n", 1)[0]
    assert "IS NOT NULL" not in where and "order_status" not in where                    # only the date window
    assert "ORDER BY " + ", ".join(c["alias"] for c in cols) in sql                       # stable group numbering


def test_key_repeats_are_summed_and_never_block_the_history_load():
    fake = _ready()
    y = timeutil.today_vn() - timedelta(days=1)
    fake.responses["dup_check"] = _dup([])
    rows = _rows(y, ("S1", "F1", 15.0, 3))
    rows[-1]["_reason"] = "2"            # 2 rows share order x SKU x FC x day with another row
    fake.responses["daily"] = rows
    res = sync.run_d1_sync("J1", fallback_mode=True)
    assert res["key_duplicate_rows"] == 2 and _demand() == (1, 15.0)
    sql = fake.queries[_real_runs(fake, "daily")[-1]][0]
    assert "SELECT DISTINCT" not in sql and "_reason IS NULL GROUP BY 3, 4, 5" in sql    # no row dropped, summed


# ------------------------------------------------------------------ 9, 12: aggregate, reconcile, re-run never doubles
def test_clean_source_is_aggregated_written_once_and_reruns_never_double():
    fake = _ready()
    y = timeutil.today_vn() - timedelta(days=1)
    d2 = y - timedelta(days=1)
    fake.responses["dup_check"] = _dup([])
    rows = _rows(y, ("S1", "F1", 15.0, 3), ("S2", "F1", 7.0, 2))
    rows.insert(2, {**rows[0], "order_date": d2, "qty": 1.0, "n_orders": 1, "n_rows": 1})     # a second day
    for t in rows[-2:]:
        t.update(qty=23.0, n_rows=6)
    fake.responses["daily"] = rows
    r1 = sync.run_d1_sync("J1", fallback_mode=True)
    assert r1["full_reload"] and r1["series_days_written"] == 3 and _demand() == (3, 23.0)
    dup_i, daily_i = _real_runs(fake, "dup_check")[-1], _real_runs(fake, "daily")[-1]
    assert dup_i < daily_i                                                                # gate BEFORE aggregation
    r2 = sync.run_d1_sync("J2", fallback_mode=True)                                       # incremental re-run
    assert not r2["full_reload"] and r2["series_days_unchanged"] >= 2 and _demand() == (3, 23.0)
    with db.tx() as c:                                                                    # forget the watermark
        c.execute("DELETE FROM sync_state WHERE name LIKE 'last_d1_sync:%'")
    r3 = sync.run_d1_sync("J3", fallback_mode=True)
    assert r3["full_reload"] and r3["series_days_new"] == 0 and _demand() == (3, 23.0)


# ------------------------------------------------------------------ 10: a failure while writing leaves nothing half done
def test_a_failure_while_writing_rolls_back_everything(monkeypatch):
    fake = _ready()
    y = timeutil.today_vn() - timedelta(days=1)
    fake.responses["dup_check"] = _dup([])
    fake.responses["daily"] = _rows(y, ("S1", "F1", 4.0, 1))
    sync.run_d1_sync("J1", fallback_mode=True)
    with db.read() as c:
        state = c.execute("SELECT value_json FROM sync_state WHERE name LIKE 'last_d1_sync:%'").fetchone()[0]
    fake.responses["daily"] = _rows(y, ("S1", "F1", 9.0, 2), ("S3", "F2", 5.0, 1))
    monkeypatch.setattr(db, "audit", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("disk I/O error")))
    with pytest.raises(RuntimeError):
        sync.run_d1_sync("J2", fallback_mode=True)
    assert _demand() == (1, 4.0)                                                          # old rows, no new rows
    with db.read() as c:
        assert c.execute("SELECT value_json FROM sync_state WHERE name LIKE 'last_d1_sync:%'").fetchone()[0] == state


@pytest.mark.parametrize("bad", [
    {"order_date": "2026-13-01"}, {"sku_id": None}, {"qty": None}, {"qty": float("nan")}, {"qty": -1.0}, "repeat"])
def test_invalid_aggregated_rows_stop_the_load_and_are_never_turned_into_zero(bad):
    fake = _ready()
    y = timeutil.today_vn() - timedelta(days=1)
    fake.responses["dup_check"] = _dup([])
    rows = _rows(y, ("S1", "F1", 4.0, 1), ("S2", "F1", 2.0, 1))
    if bad == "repeat":
        rows[1]["sku_id"] = "S1"
    else:
        rows[1].update(bad)
    fake.responses["daily"] = rows
    with pytest.raises(AppError) as e:
        sync.run_d1_sync("J1", fallback_mode=True)
    assert e.value.code in ("BQ_AGGREGATE_INVALID", "BQ_SYNC_RECONCILIATION_FAILED") and _demand() == (0, 0)


# ------------------------------------------------------------------ 11: today is never loaded (Vietnam calendar)
@pytest.mark.parametrize("utc_now,end,today", [
    (datetime(2026, 10, 16, 16, 59, tzinfo=timezone.utc), "2026-10-15", "2026-10-16"),     # 23:59 in Vietnam
    (datetime(2026, 10, 16, 17, 1, tzinfo=timezone.utc), "2026-10-16", "2026-10-17")])     # 00:01 next day
def test_d1_window_ends_yesterday_in_vietnam_time(utc_now, end, today):
    fake = _ready()
    timeutil.set_now_override(lambda: utc_now)
    fake.responses["dup_check"] = _dup([])
    fake.responses["daily"] = _rows(date.fromisoformat(end), ("S1", "F1", 1.0, 1))
    r = sync.d1_dry_run()
    w = next(s for s in r["steps"] if s["step"] == "window")
    assert (w["end"], w["today_excluded"], w["timezone"]) == (end, today, "Asia/Ho_Chi_Minh")
    sql, cfg = fake.queries[_real_runs(fake, "daily")[-1]]
    assert dict((n, v) for n, _, v in cfg["params"])["end"] == date.fromisoformat(end)


# ------------------------------------------------------------------ dry run (check + aggregate, never writes)
def test_ingest_check_runs_every_step_without_writing(client):
    fake = _ready()
    ss.update("bigquery", {"production_enabled": False})                                  # allowed before production
    y = timeutil.today_vn() - timedelta(days=1)
    fake.responses["dup_check"] = _dup([], rows_checked=6, qty=22.0)
    fake.responses["daily"] = _rows(y, ("S1", "F1", 15.0, 3), ("S2", "F1", 7.0, 3))
    r = client.post("/api/bigquery/ingest-check").json()
    assert r["status"] == "READY" and r["written"] is False and _demand() == (0, 0)
    st = {s["step"]: s for s in r["steps"]}
    assert list(st) == list(sync.D1_STEPS)
    assert all(s["status"] == "PASSED" for k, s in st.items() if k != "write") and st["write"]["status"] == "NOT_RUN"
    assert st["duplicates"]["rows_checked"] == 6 and st["duplicates"]["duplicate_groups"] == 0
    assert st["aggregate"]["rows_aggregated"] == 2 and st["aggregate"]["qty_total"] == 22.0
    assert st["compare"]["series_days_new"] == 2 and st["reconciliation"]["rows_read"] == 6
    assert st["completeness"]["completion_signal"] is True
    assert "sales_price" in st["mapping"]["compared_columns"]
    assert r["readiness"]["production_enabled"] is False
    # duplicates: BLOCKED at the duplicate step, report linked, later steps not run, still nothing written
    fake.responses["dup_check"] = _dup(TWO_GROUPS())
    r = client.post("/api/bigquery/ingest-check").json()
    st = {s["step"]: s for s in r["steps"]}
    assert r["status"] == "BLOCKED" and st["duplicates"]["status"] == "BLOCKED"
    assert st["duplicates"]["error"]["code"] == "BQ_DATA_DUPLICATES" and r["duplicate_check"]["report"]["format"] == "xlsx"
    assert all(st[k]["status"] == "NOT_RUN" for k in ("aggregate", "output_checks", "reconciliation", "compare", "write"))
    assert _demand() == (0, 0)
    assert client.get("/api/bigquery/status").json()["ingest_check"]["status"] == "BLOCKED"


def test_ingest_check_says_when_yesterday_is_not_in_the_source_yet():
    """No completion signal and yesterday has no rows: the dry run still checks everything, but says that a real D-1
    would wait / retry (the source is never assumed complete just because a query returned)."""
    fake = _ready()
    ss.update("bigquery", {"completion_table": "", "completion_date_column": ""})
    y = timeutil.today_vn() - timedelta(days=1)
    fake.responses["sufficiency"] = [{"d": y - timedelta(days=i), "n": 100} for i in range(1, 8)]     # no row for y
    fake.responses["dup_check"] = _dup([])
    fake.responses["daily"] = _rows(y - timedelta(days=1), ("S1", "F1", 3.0, 1))
    r = sync.d1_dry_run()
    st = {s["step"]: s for s in r["steps"]}
    assert r["status"] == "SOURCE_INCOMPLETE" and st["completeness"]["status"] == "WARNING"
    assert st["completeness"]["sufficiency"]["rows_for_day"] == 0 and st["completeness"]["completion_signal"] is None
    assert st["aggregate"]["status"] == "PASSED" and st["write"]["status"] == "NOT_RUN" and _demand() == (0, 0)
    with pytest.raises(sync.Retryable) as e:                                              # the real load agrees
        sync.run_d1_sync("J1", fallback_mode=True)
    assert e.value.code == "SOURCE_INCOMPLETE" and _demand() == (0, 0)


def test_validation_reports_groups_rows_and_the_report_file(client):
    fake = _ready()
    fake.responses["validate_duplicates"] = _dup(TWO_GROUPS())
    v = connector.validate_source()
    b = next(x for x in v["blocking_issues"] if x["code"] == "EXACT_DUPLICATES")
    assert b["detail"]["groups"] == 2 and b["detail"]["rows_in_groups"] == 5 and b["detail"]["rows"] == 3
    assert b["detail"]["report"]["file"].startswith("bq-trung-lap_") and not v["passed"]
    assert v["duplicate_check"]["context"] == "validate_source"
