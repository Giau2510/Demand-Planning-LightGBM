"""v1.1.4: auditable TreeSHAP. LightGBM's own outputs (pred_contrib, pred_leaf, raw_score) are the oracle."""
import csv
import io
import math

import lightgbm as lgb
import numpy as np
import pytest
from openpyxl import load_workbook

from app import db, paths
from app import settings_service as ss
from app.ml import shap_audit
from app.ml import treeshap_trace as T


def _fit(X, y, params=None, rounds=40, cat=None):
    p = {"objective": "regression", "verbose": -1, "num_leaves": 15, "min_data_in_leaf": 5, "seed": 1, "deterministic": True,
         "force_row_wise": True, **(params or {})}
    ds = lgb.Dataset(X, y, categorical_feature=cat if cat is not None else "auto", free_raw_data=False)
    return lgb.train(p, ds, num_boost_round=rounds)


def _assert_exact(b: lgb.Booster, Xq: np.ndarray):
    m = T.parse_model(b.model_to_string())
    contrib = b.predict(Xq, pred_contrib=True)
    leaves = b.predict(Xq, pred_leaf=True)
    raw = b.predict(Xq, raw_score=True)
    worst = 0.0
    for r in range(len(Xq)):
        res = T.trace_row(m, [float(v) for v in Xq[r]], collect=True)
        nf = len(res["phi"])
        d = max(abs(a - c) for a, c in zip(res["phi"] + [res["base_value"]], contrib[r]))
        worst = max(worst, d)
        assert d <= 1e-12, (r, d)
        assert list(res["leaves"]) == [int(v) for v in np.atleast_1d(leaves[r])]
        assert abs(res["base_value"] + sum(res["phi"]) - raw[r]) <= 1e-9
        per = [0.0] * nf
        for c in res["components"]:
            per[c[2]] += c[7]
        assert max(abs(a - b_) for a, b_ in zip(per, res["phi"])) <= 1e-9          # terms add up to SHAP
    return m, worst


rng = np.random.default_rng(42)


def test_single_feature_model():
    X = rng.uniform(0, 10, (400, 1))
    y = np.sin(X[:, 0]) * 3 + rng.normal(0, 0.1, 400)
    b = _fit(X, y)
    _assert_exact(b, rng.uniform(-1, 11, (30, 1)))


def test_many_features_with_interactions_and_repeated_splits():
    X = rng.uniform(0, 1, (800, 5))
    y = 4 * X[:, 0] * X[:, 1] + np.where(X[:, 2] > 0.5, X[:, 3], -X[:, 3]) + rng.normal(0, 0.05, 800)
    b = _fit(X, y, {"num_leaves": 31}, rounds=60)
    m, _ = _assert_exact(b, rng.uniform(0, 1, (40, 5)))
    # the same feature is split more than once on one root-to-leaf path (UnwindPath is exercised)
    def repeated(t, node=0, seen=()):
        if node < 0:
            return False
        f = t.split_feature[node]
        return f in seen or repeated(t, t.left_child[node], seen + (f,)) or repeated(t, t.right_child[node], seen + (f,))
    assert any(repeated(t) for t in m.trees if t.num_leaves > 1)


def test_missing_values_nan_and_zero_as_missing():
    X = rng.uniform(-1, 1, (600, 3))
    X[rng.random(600) < 0.2, 0] = np.nan
    y = np.nan_to_num(X[:, 0], nan=2.0) + X[:, 1] + rng.normal(0, 0.05, 600)
    b = _fit(X, y)
    m, _ = _assert_exact(b, np.array([[np.nan, 0.2, 0.1], [0.0, np.nan, -0.5], [np.nan, np.nan, np.nan], [0.3, 0.3, 0.3]]))
    assert any(((dt >> 2) & 3) == T.MISSING_NAN for t in m.trees for dt in t.decision_type)
    Xz = X.copy()
    Xz[np.isnan(Xz)] = 0.0
    bz = _fit(Xz, y, {"zero_as_missing": True})
    mz, _ = _assert_exact(bz, np.array([[0.0, 0.2, 0.1], [1e-36, 0.0, -0.5], [np.nan, 0.5, 0.0], [-0.0, 0.1, 0.9]]))
    assert any(((dt >> 2) & 3) == T.MISSING_ZERO for t in mz.trees for dt in t.decision_type)


def test_categorical_feature_including_unseen_negative_and_missing_categories():
    X = np.column_stack([rng.integers(0, 40, 900).astype(float), rng.uniform(0, 1, 900)])
    y = (X[:, 0] % 7) * 0.5 + X[:, 1] + rng.normal(0, 0.05, 900)
    b = _fit(X, y, {"num_leaves": 31, "min_data_per_group": 5, "cat_smooth": 1}, rounds=50, cat=[0])
    m, _ = _assert_exact(b, np.array([[3, 0.5], [39, 0.1], [77, 0.4], [-2, 0.4], [np.nan, 0.6], [5.7, 0.3]]))
    assert any(t.decision_type and any(dt & T.K_CATEGORICAL_MASK for dt in t.decision_type) for t in m.trees)


def test_feature_value_exactly_on_a_split_threshold_goes_left():
    X = rng.uniform(0, 100, (500, 2))
    y = X[:, 0] * 0.1 + rng.normal(0, 0.5, 500)
    b = _fit(X, y)
    m = T.parse_model(b.model_to_string())
    th = m.trees[0].threshold[0]
    f = m.trees[0].split_feature[0]
    row = np.array([[50.0, 50.0]])
    row[0, f] = th
    _assert_exact(b, row)
    assert m.trees[0].decision(th, 0) == m.trees[0].left_child[0]


def test_tweedie_raw_score_exp_link_and_positive_negative_tiny_contributions():
    X = rng.uniform(0, 1, (700, 4))
    y = rng.poisson(np.exp(1 + X[:, 0] - X[:, 1]))
    b = _fit(X, y, {"objective": "tweedie", "tweedie_variance_power": 1.2, "bagging_fraction": 0.9, "bagging_freq": 1}, rounds=80)
    Xq = rng.uniform(0, 1, (20, 4))
    m, _ = _assert_exact(b, Xq)
    assert m.objective == "tweedie"
    raw = b.predict(Xq, raw_score=True)
    assert np.allclose(np.exp(raw), b.predict(Xq), rtol=1e-12)                          # link = exp
    phis = np.array([T.trace_row(m, list(map(float, r)), collect=False)["phi"] for r in Xq])
    assert (phis > 0).any() and (phis < 0).any() and (np.abs(phis) < 0.01).any()
    # with bagging, the root count of a tree is the bag, not the full training set -> base value is NOT the plain mean raw score
    assert m.trees[1].internal_count[0] < 700


def test_engine_mismatch_is_detected():
    X = rng.uniform(0, 1, (300, 2))
    b = _fit(X, X[:, 0] + rng.normal(0, 0.1, 300))
    m = T.parse_model(b.model_to_string())
    m.trees[3].leaf_value[0] += 1e-6                                                   # a corrupted engine / model
    x = [0.0, 0.0]
    res = T.trace_row(m, x, collect=False)
    off = b.predict(np.array([x]), pred_contrib=True)[0]
    assert max(abs(a - c) for a, c in zip(res["phi"] + [res["base_value"]], off)) > 1e-12


def test_limits_feature_count_timeout_and_size():
    X = rng.uniform(0, 1, (300, 3))
    b = _fit(X, X[:, 0])
    m = T.parse_model(b.model_to_string())
    with pytest.raises(T.TraceError) as e:
        T.trace_row(m, [0.1, 0.2])
    assert e.value.code == "FEATURE_COUNT_MISMATCH"
    with pytest.raises(T.TraceError) as e:
        T.trace_row(m, [0.1, 0.2, 0.3], time_limit_s=-1)
    assert e.value.code == "TRACE_TIMEOUT"
    with pytest.raises(T.TraceError) as e:
        T.trace_row(m, [0.1, 0.2, 0.3], max_components=5)
    assert e.value.code == "TRACE_TOO_LARGE"


# ------------------------------------------------------------------ audit of a real published forecast
def _job(client, jid):
    return client.get(f"/api/jobs/{jid}").json()


@pytest.fixture
def demo_run(client, fast_model):
    ss.update("horizons", {"enabled": [7]})
    assert client.post("/api/demo/generate", json={"days": 150, "n_skus": 8, "seed": 3}).status_code == 200
    j = _job(client, client.post("/api/training/run", json={"force": True}).json()["id"])
    assert j["reason_code"] == "FORECAST_PUBLISHED", j
    shap_audit._cache.clear()
    return j["result"]["result"]["forecast_run_id"]


def _rows(run_id, method):
    with db.read() as c:
        return db.rows(c, "SELECT * FROM forecast_values WHERE run_id=? AND method=? ORDER BY sku_id, fc_id, target_date", (run_id, method))


def _p(r):
    return {"sku_id": r["sku_id"], "fc_id": r["fc_id"], "target_date": r["target_date"]}


def test_audit_of_published_rows_is_verified_and_reconciles(client, demo_run):
    rows = _rows(demo_run, "LIGHTGBM")
    for r in (rows[0], rows[6]):                                                       # day 1 and a recursive day
        a = client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params=_p(r)).json()
        assert a["status"] == "VERIFIED", [c for c in a["checks"] if not c["passed"]]
        chk = {c["check"]: c for c in a["checks"]}
        assert chk["SHAP_ENGINE_VS_LIGHTGBM"]["max_abs_diff"] <= 1e-12 and chk["LEAF_PATHS"]["passed"]
        v = a["values"]
        assert v["base_value"] == v["base_value_lightgbm"]
        assert abs(v["raw_score_from_attribution"] - v["raw_score_from_model"]) <= 1e-9
        assert abs(math.exp(v["raw_score_from_model"]) - v["model_output"]) <= 1e-12 * v["model_output"]
        assert round(max(v["model_output"], 0), 4) == pytest.approx(r["yhat"]) and v["published"] == r["yhat"]
        explain = client.get(f"/api/forecast-runs/{demo_run}/explain", params=_p(r)).json()["lightgbm"]
        assert explain["bias"] == pytest.approx(v["base_value"], abs=1e-12)            # same numbers as the explanation panel


def test_feature_view_tree_terms_and_where_contributions_can_come_from(client, demo_run):
    r = _rows(demo_run, "LIGHTGBM")[0]
    a = client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params=_p(r)).json()
    top = max(a["features"], key=lambda f: abs(f["shap_engine"]))
    fv = client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params={**_p(r), "feature": top["name"], "page_size": 500}).json()["feature_view"]
    s = fv["summary"]
    assert sum(t["sum"] for t in fv["trees"]["items"]) == pytest.approx(top["shap_engine"], abs=1e-9)
    assert s["sum_positive_terms"] + s["sum_negative_terms"] == pytest.approx(top["shap_engine"], abs=1e-9)
    # TreeSHAP adds to phi[f] only for path elements created at nodes splitting on f -> only trees that split on f
    with_terms = [t for t in fv["trees"]["items"] if t["n_terms"]]
    assert all(t.get("splits") for t in with_terms) and s["n_trees_with_terms"] <= s["n_trees_splitting_on_feature"]
    t0 = with_terms[0]["tree"]
    tt = client.get(f"/api/forecast-runs/{demo_run}/shap-trace/tree", params={**_p(r), "feature": top["name"], "tree": t0}).json()
    assert tt["sum"] == pytest.approx(with_terms[0]["sum"], abs=1e-12)
    for it in tt["items"]:
        assert it["term"] == pytest.approx(it["weight"] * (it["one_fraction"] - it["zero_fraction"]) * it["leaf_value"], abs=1e-15)
        assert it["one_fraction"] in (0.0, 1.0) and 0 < it["zero_fraction"] <= 1
        assert any(st["is_feature"] for st in it["route"])
    assert sum(1 for it in tt["items"] if it["row_reaches_leaf"]) <= 1


def test_exports_are_independently_checkable(client, demo_run):
    r = _rows(demo_run, "LIGHTGBM")[0]
    a = client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params=_p(r)).json()
    x = client.get(f"/api/forecast-runs/{demo_run}/shap-trace.xlsx", params={**_p(r), "lang": "vi"})
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Thông tin & đối soát", "SHAP 33 feature", "Theo cây", "Thành phần"]
    info = {k: v for k, v in wb.worksheets[0].iter_rows(min_row=2, values_only=True)}
    assert info["model_sha256"] == a["identity"]["model_sha256_file"] and info["feature_row_hash (sha256 of float32 row)"] == a["identity"]["feature_row_hash"]
    terms = list(wb.worksheets[3].iter_rows(min_row=2, values_only=True))
    per = {}
    for t in terms:
        per[t[0]] = per.get(t[0], 0.0) + t[9]
        assert t[9] == pytest.approx(t[7] * t[8] * t[3], abs=1e-15)                    # recomputable from the file alone
    for f in a["features"]:
        assert per.get(f["name"], 0.0) == pytest.approx(f["shap_lightgbm"], abs=1e-9)
    c = client.get(f"/api/forecast-runs/{demo_run}/shap-trace.csv", params={**_p(r), "feature": "lag_1"})
    rows = list(csv.DictReader(io.StringIO(c.content.decode("utf-8-sig"))))
    assert rows and {x["feature"] for x in rows} == {"lag_1"}
    assert sum(float(x["term"]) for x in rows) == pytest.approx(next(f for f in a["features"] if f["name"] == "lag_1")["shap_lightgbm"], abs=1e-9)
    assert rows[0]["model_sha256"] == a["identity"]["model_sha256_file"]


def test_fallback_row_missing_or_changed_artifact_and_names_mismatch(client, demo_run, monkeypatch):
    fb = _rows(demo_run, "FALLBACK_NEW_SERIES")[0]
    assert client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params=_p(fb)).json()["code"] == "NOT_LIGHTGBM"
    r = _rows(demo_run, "LIGHTGBM")[0]
    with db.read() as c:
        mp = paths.models_dir() / c.execute("SELECT t.model_path FROM training_runs t JOIN forecast_runs f ON f.training_run_id=t.id WHERE f.id=?", (demo_run,)).fetchone()[0]
    original = mp.read_bytes()
    mp.write_bytes(original + b"\n")                                                   # artifact changed after publish
    a = client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params=_p(r)).json()
    assert a["status"] == "NOT_REPRODUCIBLE" and a["reason"] == "MODEL_ARTIFACT_CHANGED"
    mp.unlink()                                                                         # artifact gone
    a = client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params=_p(r)).json()
    assert a["status"] == "NOT_REPRODUCIBLE" and a["reason"] == "MODEL_ARTIFACT_MISSING"
    mp.write_bytes(original)
    # feature order of the training run no longer matches the model -> FAILED, never shown as verified
    import json as _json
    orig_loads = _json.loads
    def swapped(s, *a_, **k):
        d = orig_loads(s, *a_, **k)
        if isinstance(d, dict) and isinstance(d.get("features"), list) and "lag_1" in d["features"]:
            d = {**d, "features": list(reversed(d["features"]))}
        return d
    monkeypatch.setattr(shap_audit.json, "loads", swapped)
    shap_audit._cache.clear()
    a = client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params=_p(r)).json()
    assert a["status"] == "FAILED" and any(c["check"] == "FEATURE_NAMES" and not c["passed"] for c in a["checks"])


def test_engine_mismatch_marks_failed_and_cache_keys(client, demo_run, monkeypatch):
    rows = _rows(demo_run, "LIGHTGBM")
    r0, r1 = rows[0], rows[1]
    a0 = shap_audit.audit(demo_run, r0["sku_id"], r0["fc_id"], r0["target_date"])
    assert shap_audit.audit(demo_run, r0["sku_id"], r0["fc_id"], r0["target_date"]) is a0      # cached
    a1 = shap_audit.audit(demo_run, r1["sku_id"], r1["fc_id"], r1["target_date"])
    assert a1 is not a0 and a1["identity"]["feature_row_hash"] != a0["identity"]["feature_row_hash"]
    keys = list(shap_audit._cache)
    assert all(k[0] == a0["identity"]["model_sha256_file"] for k in keys) and len({k[-1] for k in keys}) == 2
    shap_audit._cache.clear()
    real = T.trace_tree
    def broken(t, x, phi, comps, routes):
        leaf = real(t, x, phi, comps, routes)
        if t.index == 5:
            phi[0] += 1e-6                                                             # engine drifts from LightGBM
        return leaf
    monkeypatch.setattr(T, "trace_tree", broken)
    a = shap_audit.audit(demo_run, r0["sku_id"], r0["fc_id"], r0["target_date"])
    assert a["status"] == "FAILED" and not next(c for c in a["checks"] if c["check"] == "SHAP_ENGINE_VS_LIGHTGBM")["passed"]


def test_trace_limits_are_reported_by_the_api(client, demo_run, monkeypatch):
    r = _rows(demo_run, "LIGHTGBM")[0]
    monkeypatch.setattr(shap_audit, "MAX_COMPONENTS", 10)
    resp = client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params=_p(r))
    assert resp.status_code == 422 and resp.json()["code"] == "TRACE_TOO_LARGE"
    monkeypatch.setattr(shap_audit, "MAX_COMPONENTS", 3_000_000)
    monkeypatch.setattr(shap_audit, "TIME_LIMIT_S", -1)
    resp = client.get(f"/api/forecast-runs/{demo_run}/shap-trace", params=_p(r))
    assert resp.status_code == 422 and resp.json()["code"] == "TRACE_TIMEOUT"
