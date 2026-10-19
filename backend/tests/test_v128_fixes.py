"""1.2.8: "how is this calculated" never waits endlessly, Results page accuracy, whole-number forecasts, day-of-week switch.

* fingerprint(): one sequential scan instead of a lookup of every row through the key index - the same hash as before, so
  the fingerprints stored with training runs keep matching (fractional quantities: the original key-order statement).
* explanation: one reproduction at a time - concurrent requests wait for it and reuse it (no parallel panel builds).
* accuracy: the run's own rows are the outer loop of the join (CROSS JOIN) - identical rows.
* forecast quantity shown / exported = round half up of the stored yhat, never below 0 (app/qty.py); totals are sums of
  the rounded daily values; the stored yhat is unchanged and still the input of WAPE / bias.
* feature_weekday: dow + is_weekend of the TARGET day as their own switch; None (every older configuration) = as before.
"""
import concurrent.futures as cf
import hashlib
import io
import json
import sqlite3
import threading
import time
from collections import defaultdict
from datetime import date, datetime, timedelta

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import db, qty, timeutil
from app import settings_service as ss
from app.data.demand import fingerprint
from app.ml import explain as ex_mod
from app.ml import metrics as mt
from app.ml.features import CALENDAR, calendar_names, direct_extra, feature_names, is_direct, weekday_on
from app.timeutil import VN_TZ

OLD_FP = ("SELECT COUNT(*), COALESCE(SUM(qty),0), MIN(order_date), MAX(order_date), MAX(loaded_at) "
          "FROM demand_daily WHERE scope_key=?")
OLD_ACC = ("SELECT f.horizon_day, f.yhat, d.qty FROM forecast_values f JOIN demand_daily d ON d.scope_key=? AND "
           "d.order_date=f.target_date AND d.sku_id=f.sku_id AND d.fc_id=f.fc_id WHERE f.run_id=? AND f.target_date < ?")


def _publish(client):
    assert client.post("/api/demo/generate", json={"days": 150, "n_skus": 8, "seed": 3}).status_code == 200
    j = client.post("/api/training/run", json={}).json()
    assert j["status"] == "SUCCEEDED" and j["reason_code"] == "FORECAST_PUBLISHED", j
    return client.get("/api/system/status").json()["latest_forecast_run"]


def _old_fp(c, sk):
    return hashlib.sha256(repr(tuple(c.execute(OLD_FP, (sk,)).fetchone())).encode()).hexdigest()[:16]


def _traced(monkeypatch) -> list[str]:
    seen: list[str] = []
    real = db.connect

    def connect(*a, **k):
        c = real(*a, **k)
        c.set_trace_callback(seen.append)
        return c
    monkeypatch.setattr(db, "connect", connect)
    return seen


# ------------------------------------------------------------------ fingerprint
def test_fingerprint_is_the_old_hash_from_one_plain_scan(client, fast_model, monkeypatch):
    run = _publish(client)
    sk = run["scope_key"]
    with db.read() as c:
        old = _old_fp(c, sk)
        assert fingerprint(sk, c) == old
        stored = db.one(c, "SELECT data_fingerprint FROM training_runs WHERE id=?", (run["training_run_id"],))
    assert stored["data_fingerprint"] == old                 # the run still counts as "same data" (SAME_LOAD)
    with monkeypatch.context() as m:
        seen = _traced(m)
        with db.read() as c:
            fingerprint(sk, c)
    reads = [s for s in seen if "FROM demand_daily" in s]
    assert len(reads) == 1, reads
    with db.read() as c:
        assert [r[3] for r in c.execute("EXPLAIN QUERY PLAN " + reads[0]).fetchall()] == ["SCAN demand_daily"]


def test_fingerprint_with_fractional_quantities_keeps_the_original_summation_order(client, fast_model):
    run = _publish(client)
    sk = run["scope_key"]
    with db.tx() as c:
        c.execute("UPDATE demand_daily SET qty = qty + 0.1 WHERE rowid IN "
                  "(SELECT rowid FROM demand_daily WHERE scope_key=? ORDER BY rowid DESC LIMIT 9)", (sk,))
    with db.read() as c:
        assert fingerprint(sk, c) == _old_fp(c, sk)


# ------------------------------------------------------------------ explanation: one reproduction at a time
def test_concurrent_explanations_build_the_inputs_and_run_the_model_once(client, fast_model, monkeypatch):
    run = _publish(client)
    rid = run["id"]
    vals = client.get(f"/api/forecast-runs/{rid}/values?horizon={run['max_horizon']}&limit=5000").json()["rows"]
    lgb = [r for r in vals if r["method"] == "LIGHTGBM"]
    assert lgb, "the demo run has LightGBM series"
    sku, fc = lgb[0]["sku_id"], lgb[0]["fc_id"]
    days = sorted({r["target_date"] for r in vals if (r["sku_id"], r["fc_id"]) == (sku, fc)})[:4]
    builds, model_runs = [], []
    real_panel, real_fc = ex_mod.panel_for_run, ex_mod.model.forecast
    gate = threading.Event()

    def slow_panel(tr, *a, **k):                    # a slow build: the other requests arrive while it runs
        builds.append(1)
        gate.wait(10)
        return real_panel(tr, *a, **k)

    with monkeypatch.context() as m:
        m.setattr(ex_mod, "panel_for_run", slow_panel)
        m.setattr(ex_mod.model, "forecast", lambda *a, **k: model_runs.append(1) or real_fc(*a, **k))
        ex_mod._inputs.clear()
        ex_mod._cache.clear()
        with cf.ThreadPoolExecutor(len(days)) as pool:
            futs = [pool.submit(ex_mod.explain, rid, sku, fc, d) for d in days]
            time.sleep(0.5)
            gate.set()
            res = [f.result(timeout=120) for f in futs]
    assert len(builds) == 1 and len(model_runs) == 1
    assert [r["reproduction"]["status"] for r in res] == ["OK"] * len(days)
    assert all(r["lightgbm"]["matches_published"] for r in res)


# ------------------------------------------------------------------ accuracy (Results page)
def test_accuracy_rows_equal_the_old_join_with_the_run_as_outer_loop(client, fast_model, monkeypatch):
    run = _publish(client)
    sk, rid = run["scope_key"], run["id"]
    with db.tx() as c:                               # later actuals for the first 3 forecast days
        rows = c.execute("SELECT sku_id, fc_id, target_date FROM forecast_values WHERE run_id=? AND horizon_day<=3 "
                         "ORDER BY sku_id, fc_id, target_date", (rid,)).fetchall()
        for i, (s, f, d) in enumerate(rows):
            c.execute("INSERT OR IGNORE INTO demand_daily(scope_key, data_source, order_date, sku_id, fc_id, qty, loaded_at) "
                      "VALUES (?, 'DEMO', ?, ?, ?, ?, '2026-01-01T00:00:00+00:00')", (sk, d, s, f, float(i % 5)))
    later = date.fromisoformat(run["forecast_origin"]) + timedelta(days=5)
    timeutil.set_now_override(lambda: datetime(later.year, later.month, later.day, 10, 0, tzinfo=VN_TZ))
    with monkeypatch.context() as m:
        seen = _traced(m)
        acc = client.get(f"/api/forecast-runs/{rid}/accuracy").json()
    with db.read() as c:
        old = sorted(tuple(r) for r in c.execute(OLD_ACC, (sk, rid, str(later))).fetchall())
        stmt = next(s for s in seen if "CROSS JOIN demand_daily" in s)
        plan = [r[3] for r in c.execute("EXPLAIN QUERY PLAN " + stmt).fetchall()]
        assert sorted(tuple(r) for r in c.execute(stmt).fetchall()) == old and len(old) == len(rows)
    assert plan[0].startswith("SEARCH f USING INDEX") and "(run_id=?)" in plan[0]
    q = ss.get("quality")
    df = pd.DataFrame(old, columns=["horizon_day", "yhat", "qty"])
    ref = mt.compute(df["qty"], df["yhat"], q.mape_min_actual)
    assert acc["available"] and acc["overall"]["n_points"] == ref["n_points"] == len(old)
    assert acc["overall"]["wape"] == pytest.approx(ref["wape"])          # WAPE still from the stored (raw) yhat


# ------------------------------------------------------------------ whole-number forecast quantity
@pytest.mark.parametrize("x, want", [(0.0, 0), (0.49, 0), (0.5, 1), (1.4999, 1), (1.5, 2), (2.5, 3), (3.0, 3), (12.345, 12),
                                     (1e6 + 0.5, 1000001), (-0.3, 0), (-2.7, 0)])
def test_quantity_rule_python_and_sql_agree(x, want):
    assert qty.business_qty(x) == want
    c = sqlite3.connect(":memory:")
    assert c.execute(f"SELECT {qty.SQL} FROM (SELECT ? AS yhat)", (x,)).fetchone()[0] == want


def test_quantity_rule_keeps_missing_missing():
    assert qty.business_qty(None) is None and qty.business_qty(float("nan")) is None


def test_values_summary_export_and_explanation_show_the_same_whole_numbers(client, fast_model):
    run = _publish(client)
    rid, h = run["id"], run["max_horizon"]
    with db.read() as c:
        stored = {(r["sku_id"], r["fc_id"], r["target_date"]): r["yhat"] for r in db.rows(
            c, "SELECT sku_id, fc_id, target_date, yhat FROM forecast_values WHERE run_id=?", (rid,))}
    assert any(v != int(v) for v in stored.values())          # the run really has fractional forecasts
    vals = client.get(f"/api/forecast-runs/{rid}/values?horizon={h}&limit=5000").json()
    assert vals["qty_rule"] == qty.RULE and len(vals["rows"]) == len(stored)
    by_day, by_series, by_method = defaultdict(int), defaultdict(int), defaultdict(int)
    for r in vals["rows"]:
        k = (r["sku_id"], r["fc_id"], r["target_date"])
        assert r["yhat"] == stored[k]                                          # stored value unchanged
        assert isinstance(r["forecast_qty"], int) and r["forecast_qty"] == qty.business_qty(stored[k])
        by_day[r["target_date"]] += r["forecast_qty"]
        by_series[(r["sku_id"], r["fc_id"])] += r["forecast_qty"]
        by_method[r["method"]] += r["forecast_qty"]
    s = client.get(f"/api/forecast-runs/{rid}/summary?horizon={h}&limit=5000").json()
    assert {d["target_date"]: d["total"] for d in s["daily"]} == by_day
    assert {(x["sku_id"], x["fc_id"]): x["total"] for x in s["series"]} == by_series
    assert {m["method"]: m["total"] for m in s["methods"]} == by_method
    assert isinstance(s["grand_total"], int) and s["grand_total"] == sum(by_day.values())
    totals = [x["total"] for x in s["series"]]
    assert totals == sorted(totals, reverse=True)
    wb = load_workbook(io.BytesIO(client.get(f"/api/forecast-runs/{rid}/export.xlsx?horizon={h}&lang=en").content))
    daily = list(wb["Daily"].iter_rows(min_row=2, values_only=True))
    assert {(r[0], r[2], r[3].date().isoformat()): r[5] for r in daily} == {k: qty.business_qty(v) for k, v in stored.items()}
    ws = wb[f"Totals {h}d"]
    head = [c.value for c in ws[1]]
    ci = head.index("Period total forecast")
    got = defaultdict(int)
    for r in ws.iter_rows(min_row=2, values_only=True):
        assert isinstance(r[ci], int)
        got[(r[head.index("SKU")], r[head.index("FC/warehouse")])] += r[ci]
    assert dict(got) == dict(by_series)
    k = next(k for k, v in stored.items() if v != int(v))
    ex = client.get(f"/api/forecast-runs/{rid}/explain", params={"sku_id": k[0], "fc_id": k[1], "target_date": k[2]}).json()
    assert ex["result"]["yhat"] == stored[k] and ex["result"]["forecast_qty"] == qty.business_qty(stored[k])
    with db.read() as c:                                                      # nothing stored was rounded
        assert {(r["sku_id"], r["fc_id"], r["target_date"]): r["yhat"] for r in db.rows(
            c, "SELECT sku_id, fc_id, target_date, yhat FROM forecast_values WHERE run_id=?", (rid,))} == stored


# ------------------------------------------------------------------ day-of-week switch
def _legacy_names(cfg) -> list[str]:
    """feature_names of 1.2.7, verbatim: the calendar block is one group."""
    from app.ml.features import LAGS, ROLL_WINDOWS, discount_on
    names = ["sku_code", "fc_code", "series_age"] + ([f"lag_{k}" for k in LAGS] if cfg.feature_lags else [])
    if cfg.feature_rolling:
        names += [f"roll_mean_{w}" for w in ROLL_WINDOWS] + ["roll_std_7", "roll_std_28", "zero_share_28",
                                                             "days_since_last_sale", "roll_max_28"]
    if cfg.feature_calendar:
        names += ["dow", "dom", "month", "weekofyear", "is_weekend", "days_to_month_end"]
    if cfg.feature_price:
        names += ["price_last", "price_rel_28"]
    if discount_on(cfg):
        names += ["disc_last"]
    if cfg.feature_promotion:
        names += ["promo_last", "promo_mean_7"]
    if is_direct(cfg):
        names = [n for n in names if n not in CALENDAR] + (CALENDAR if cfg.feature_calendar else []) + ["horizon"]
    return names


@pytest.mark.parametrize("strategy", ["RECURSIVE", "DIRECT"])
def test_weekday_switch_unset_gives_exactly_the_older_feature_list(strategy):
    cfg = ss.ModelSettings(forecast_strategy=strategy)
    assert cfg.feature_weekday is None and weekday_on(cfg)
    assert feature_names(cfg) == _legacy_names(cfg)
    stored = {k: v for k, v in cfg.model_dump().items() if k != "feature_weekday"}     # saved before 1.2.8
    assert feature_names(ss.ModelSettings.model_validate(stored)) == _legacy_names(cfg)
    no_cal = ss.ModelSettings(forecast_strategy=strategy, feature_calendar=False)
    assert feature_names(no_cal) == _legacy_names(no_cal) and not weekday_on(no_cal)


@pytest.mark.parametrize("strategy", ["RECURSIVE", "DIRECT"])
def test_weekday_switch_off_or_alone(strategy):
    off = ss.ModelSettings(forecast_strategy=strategy, feature_weekday=False)
    names = feature_names(off)
    assert "dow" not in names and "is_weekend" not in names
    assert [n for n in names if n in CALENDAR] == ["dom", "month", "weekofyear", "days_to_month_end"]
    alone = ss.ModelSettings(forecast_strategy=strategy, feature_calendar=False, feature_weekday=True)
    assert [n for n in feature_names(alone) if n in CALENDAR] == ["dow", "is_weekend"]


@pytest.mark.parametrize("day, dow, weekend", [("2026-12-31", 3, 0.0), ("2027-01-01", 4, 0.0), ("2027-01-02", 5, 1.0),
                                               ("2027-01-03", 6, 1.0), ("2026-02-28", 5, 1.0), ("2026-03-01", 6, 1.0),
                                               ("2026-03-02", 0, 0.0)])
def test_weekday_values_are_those_of_the_target_day(day, dow, weekend):
    d = pd.Timestamp(day)
    on = ss.ModelSettings(forecast_strategy="DIRECT")
    extra = direct_extra(on, d, 3)
    vals = dict(zip(calendar_names(on), extra[:-1]))
    assert vals["dow"] == dow and vals["is_weekend"] == weekend and extra[-1] == 3.0
    assert vals["dom"] == d.day and vals["month"] == d.month
    off = ss.ModelSettings(forecast_strategy="DIRECT", feature_weekday=False)
    assert len(direct_extra(off, d, 3)) == len(calendar_names(off)) + 1 == 5


def test_training_and_forecast_with_weekday_off_share_one_feature_list(client, fast_model):
    r = client.put("/api/settings/model", json={**ss.get("model").model_dump(), "feature_weekday": False})
    assert r.status_code == 200 and r.json()["feature_weekday"] is False
    run = _publish(client)
    with db.read() as c:
        fc = json.loads(db.one(c, "SELECT feature_config_json FROM training_runs WHERE id=?",
                               (run["training_run_id"],))["feature_config_json"])
    assert "dow" not in fc["features"] and "is_weekend" not in fc["features"] and "dom" in fc["features"]
    assert fc["config"]["feature_weekday"] is False and fc["settings_switches"]["feature_weekday"] is False
    view = client.get("/api/features/registry").json()
    reg = {e["feature_key"]: e for e in view["entries"]}
    assert reg["dow"]["availability_status"] == "DISABLED" and reg["dow"]["switch"] == "feature_weekday"
    assert reg["is_weekend"]["availability_status"] == "DISABLED" and reg["dom"]["availability_status"] == "ACTIVE"
    assert "dow" not in view["next_training"]["features"] and "dom" in view["next_training"]["features"]
    vals = client.get(f"/api/forecast-runs/{run['id']}/values?limit=5000").json()["rows"]
    lgb = [x for x in vals if x["method"] == "LIGHTGBM"]
    assert lgb, "the demo run has LightGBM series"
    ex = client.get(f"/api/forecast-runs/{run['id']}/explain",
                    params={"sku_id": lgb[0]["sku_id"], "fc_id": lgb[0]["fc_id"], "target_date": lgb[0]["target_date"]}).json()
    assert ex["reproduction"]["status"] == "OK" and ex["lightgbm"]["matches_published"]     # inference = training list
    assert {f["name"] for f in ex["lightgbm"]["features"]} == set(fc["features"])
