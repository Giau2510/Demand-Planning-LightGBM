"""1.2.4: (A) the explanation of a forecast (how it was produced, demand history + Excel) survives a reload of
identical data and refuses a real change; older runs get a data identity only by attestation from a backup that holds
their training data. (B) product names (one per SKU, never duplicating a row or changing a number), keyword search,
Excel. (C) D-1 / D-0 switches: persisted, restored, refused when not ready, no past slot run.

Kind of evidence: integration on an isolated database (conftest: temporary DF_DATA_DIR + process-wide fallback) with a
FAKE BigQuery client. The real production check (17/10/2026) is in the report: the content identity of the copy taken
before the 19:58 D-1 reload equals the live one (c4777152...), 1,349,388 rows."""
import io
import json
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app import backup, db, timeutil
from app import settings_service as ss
from app.bq import product_names
from app.data import identity
from app.jobs import flows
from app.ml import pipeline

from test_v116 import D0_OUT
from test_v118 import _bq_ready_production
from test_v122_bigquery_policy import _bq_history


@pytest.fixture(autouse=True)
def _fresh_caches():
    identity._cache.clear()
    product_names._cache.update(updated_at=None, names={})
    yield
    identity._cache.clear()


def _published():
    fake = _bq_ready_production()
    ss.update("horizons", {"enabled": [7]})
    scope = _bq_history(lam=5.0)          # dense enough for LightGBM series (lam=1 gives only intermittent ones)
    res = pipeline.train_and_publish(force=True)
    assert res["status"] == "PUBLISHED", res
    return fake, scope, res


def _tr(run_id):
    with db.read() as c:
        return db.one(c, "SELECT * FROM training_runs WHERE id=?", (run_id,))


def _lgb_series(client, fr):
    vals = client.get(f"/api/forecast-runs/{fr}/values?limit=5000").json()["rows"]
    row = next(r for r in vals if r["method"] == "LIGHTGBM")
    return row["sku_id"], row["fc_id"], row["target_date"], row["yhat"]


def _explain(client, fr, sku, fc, d):
    r = client.get(f"/api/forecast-runs/{fr}/explain", params={"sku_id": sku, "fc_id": fc, "target_date": d})
    assert r.status_code == 200, r.text
    return r.json()


def _reload_same_values(scope):
    """What a D-1 full reload of unchanged data does: every row rewritten, same values, new loaded_at."""
    with db.tx() as c:
        c.execute("UPDATE demand_daily SET loaded_at=? WHERE scope_key=?", ("2099-01-01T00:00:00+00:00", scope))


# ------------------------------------------------------------------ A. data identity and the two explanation areas
def test_reload_of_identical_data_still_reproduces_exactly_history_and_excel_included(client, fast_model):
    _, scope, res = _published()
    tr = _tr(res["training_run_id"])
    ident = json.loads(tr["feature_config_json"])["data_identity"]
    with db.read() as c:
        n = c.execute("SELECT COUNT(*) FROM demand_daily WHERE scope_key=? AND order_date BETWEEN ? AND ?",
                      (scope, tr["train_start"], tr["train_end"])).fetchone()[0]
    assert ident["algorithm"] == "content-v1" and ident["rows"] == n and [ident["start"], ident["end"]] == [tr["train_start"], tr["train_end"]]
    fr = res["forecast_run_id"]
    sku, fc, d, yhat = _lgb_series(client, fr)
    before = _explain(client, fr, sku, fc, d)
    assert before["reproduction"]["status"] == "OK" and before["reproduction"]["data_check"]["status"] == "SAME_LOAD"
    _reload_same_values(scope)
    with db.read() as c:
        assert identity.data_state(tr, scope, c)["status"] == "SAME_CONTENT"
    ex = _explain(client, fr, sku, fc, d)
    assert ex["reproduction"]["status"] == "OK", ex["reproduction"]
    dc = ex["reproduction"]["data_check"]
    assert dc["status"] == "SAME_CONTENT" and dc["origin"] == "TRAINING" and dc["rows"] == n
    assert ex["history"]["status"] == "OK" and ex["history"]["days"] == before["history"]["days"]
    assert ex["lightgbm"]["status"] == "OK" and ex["lightgbm"]["matches_published"] is True
    assert ex["lightgbm"]["published_yhat"] == pytest.approx(yhat)
    x = client.get(f"/api/forecast-runs/{fr}/explain/history.xlsx", params={"sku_id": sku, "fc_id": fc, "target_date": d})
    assert x.status_code == 200 and load_workbook(io.BytesIO(x.content)).worksheets[1].max_row == len(ex["history"]["days"]) + 1


def test_a_changed_value_inside_the_window_names_the_day_and_nothing_is_reproduced(client, fast_model):
    _, scope, res = _published()
    tr = _tr(res["training_run_id"])
    fr = res["forecast_run_id"]
    sku, fc, d, _ = _lgb_series(client, fr)
    day = str(date.fromisoformat(tr["train_end"]) - timedelta(days=3))
    with db.tx() as c:
        cur = c.execute("UPDATE demand_daily SET qty=qty+1 WHERE rowid=(SELECT rowid FROM demand_daily WHERE scope_key=? "
                        "AND order_date=? LIMIT 1)", (scope, day))
        assert cur.rowcount == 1
    ex = _explain(client, fr, sku, fc, d)
    assert ex["reproduction"]["status"] == "DATA_CHANGED" and ex["history"]["status"] == "DATA_CHANGED"
    assert ex["lightgbm"]["status"] == "DATA_CHANGED"
    dc = ex["reproduction"]["data_check"]
    assert dc["status"] == "CHANGED" and dc["changed_days"] == [day] and dc["n_changed_days"] == 1
    x = client.get(f"/api/forecast-runs/{fr}/explain/history.xlsx", params={"sku_id": sku, "fc_id": fc, "target_date": d})
    assert x.status_code == 409                                              # never an export of other data


def test_a_new_day_after_the_window_keeps_the_run_explainable(client, fast_model):
    _, scope, res = _published()
    tr = _tr(res["training_run_id"])
    nxt = str(date.fromisoformat(tr["train_end"]) + timedelta(days=1))
    with db.tx() as c:          # the next D-1 adds a day after the training window
        c.execute("INSERT INTO demand_daily(scope_key, data_source, order_date, sku_id, fc_id, qty, n_orders, loaded_at) "
                  "VALUES (?, 'BIGQUERY', ?, 'S0', 'F0', 3, 1, ?)", (scope, nxt, timeutil.now_iso()))
    with db.read() as c:
        assert identity.data_state(tr, scope, c)["status"] == "SAME_CONTENT"
    fr = res["forecast_run_id"]
    sku, fc, d, _ = _lgb_series(client, fr)
    assert _explain(client, fr, sku, fc, d)["reproduction"]["status"] == "OK"


def test_an_older_run_is_unverifiable_until_attested_from_a_backup_holding_its_data(client, fast_model, tmp_path):
    _, scope, res = _published()
    rid = res["training_run_id"]
    with db.tx() as c:          # a run trained before 1.2.4: no identity stored with it
        c.execute("UPDATE training_runs SET feature_config_json=json_remove(feature_config_json, '$.data_identity') WHERE id=?", (rid,))
    (tmp_path / "b1").mkdir()
    (tmp_path / "b2").mkdir()
    good = backup.backup_database("manual", dest_dir=tmp_path / "b1")      # taken while the data is unchanged
    _reload_same_values(scope)
    fr = res["forecast_run_id"]
    sku, fc, d, _ = _lgb_series(client, fr)
    ex = _explain(client, fr, sku, fc, d)
    assert ex["reproduction"]["status"] == "DATA_CHANGED" and ex["reproduction"]["data_check"]["status"] == "UNVERIFIABLE"
    late = backup.backup_database("manual", dest_dir=tmp_path / "b2")      # after the reload: not the training data
    from app.messages import AppError
    with pytest.raises(AppError) as e:
        identity.attest_from_backup(rid, late)
    assert "does not hold the training data" in e.value.detail
    out = identity.attest_from_backup(rid, good)
    assert out["current_state"]["status"] == "SAME_CONTENT" and out["identity"]["attestation"]["backup"] == good.name
    ex = _explain(client, fr, sku, fc, d)
    assert ex["reproduction"]["status"] == "OK" and ex["reproduction"]["data_check"]["origin"] == "ATTESTED"
    with pytest.raises(AppError):                                            # stored once, never overwritten
        identity.attest_from_backup(rid, good)
    with db.read() as c:
        assert c.execute("SELECT COUNT(*) FROM audit_log WHERE action='TRAINING_DATA_ATTESTED'").fetchone()[0] == 1


def test_model_page_backtest_flag_follows_the_same_rule(client, fast_model):
    _, scope, res = _published()
    _reload_same_values(scope)
    ov = client.get("/api/model/training-explain", params={"training_run_id": res["training_run_id"]}).json()
    assert ov["data_unchanged_since_training"] is True and ov["data_check"]["status"] == "SAME_CONTENT"


# ------------------------------------------------------------------ B. product names, search, Excel
NAMES = [{"sku": "S0", "name": "Hỗ trợ gan Simarin Plus Stella (H/30V)"}, {"sku": "S1", "name": "Diamicron MR 30mg Servier"},
         {"sku": "S2", "name": "Tên A"}, {"sku": "S2", "name": "Tên B"},              # two names: ambiguous -> none
         {"sku": "S1", "name": "Diamicron MR 30mg Servier"},                            # repeated identical: one name
         {"sku": "ZZZ", "name": "not a local SKU"}]


def _summary(client, fr, **p):
    return client.get(f"/api/forecast-runs/{fr}/summary", params={"horizon": 7, "limit": 5000, **p}).json()


def test_names_are_one_per_sku_never_duplicate_a_row_and_never_change_a_number(client, fast_model):
    fake, _, res = _published()
    fr = res["forecast_run_id"]
    before = _summary(client, fr)
    vals_before = client.get(f"/api/forecast-runs/{fr}/values?limit=5000").json()
    assert all(s["product_name"] is None for s in before["series"])
    fake.responses["product_names"] = NAMES
    st = client.post("/api/bigquery/product-names/refresh").json()
    assert st["skus_local"] == 5 and st["skus_named"] == 2 and st["skus_ambiguous"] == 1 and st["skus_without_name"] == 3
    assert st["source"].endswith(".dim_sku") and "names" not in st
    after = _summary(client, fr)
    vals_after = client.get(f"/api/forecast-runs/{fr}/values?limit=5000").json()
    assert after["n_series"] == before["n_series"] and after["grand_total"] == before["grand_total"]
    assert vals_after["total"] == vals_before["total"] and len(vals_after["rows"]) == len(vals_before["rows"])
    assert [(s["sku_id"], s["fc_id"], s["total"]) for s in after["series"]] == [(s["sku_id"], s["fc_id"], s["total"]) for s in before["series"]]
    names = {s["sku_id"]: s["product_name"] for s in after["series"]}
    assert names["S0"].startswith("Hỗ trợ gan") and names["S1"] == "Diamicron MR 30mg Servier"
    assert names["S2"] is None and names["S3"] is None                     # ambiguous / missing: no name, never a guess
    run = client.get(f"/api/forecast-runs/{fr}").json()
    assert run["product_names"]["skus_named"] == 2
    with db.read() as c:
        assert product_names.for_run({"data_source": "DEMO"}, c) == {}     # never applied to another source


def test_search_partial_case_and_accent_insensitive_combined_with_the_fc_filter(client, fast_model):
    fake, _, res = _published()
    fr = res["forecast_run_id"]
    fake.responses["product_names"] = NAMES
    client.post("/api/bigquery/product-names/refresh")
    skus = lambda r: sorted({s["sku_id"] for s in r["series"]})  # noqa: E731
    assert skus(_summary(client, fr, q="s0")) == ["S0"]                       # part of the SKU, lower case
    assert skus(_summary(client, fr, q="ho tro GAN")) == ["S0"]               # no accents, upper case
    assert skus(_summary(client, fr, q="Hỗ Trợ")) == ["S0"]
    assert skus(_summary(client, fr, q="diamicron 30")) == ["S1"]             # several words: all must match
    r = _summary(client, fr, q="gan diamicron")
    assert r["n_series"] == 0 and r["series"] == [] and r["grand_total"] == 0
    r = _summary(client, fr, q="s", fc_id="F1")
    assert r["n_series"] == 5 and {s["fc_id"] for s in r["series"]} == {"F1"}
    v = client.get(f"/api/forecast-runs/{fr}/values", params={"horizon": 7, "q": "diamicron", "limit": 5000}).json()
    assert v["total"] == 2 * 7 and {x["sku_id"] for x in v["rows"]} == {"S1"}
    assert _summary(client, fr, q="  ")["n_series"] == _summary(client, fr)["n_series"]   # cleared = everything


def test_excel_has_the_name_after_the_sku_same_rows_same_totals(client, fast_model):
    fake, _, res = _published()
    fr = res["forecast_run_id"]
    fake.responses["product_names"] = NAMES
    client.post("/api/bigquery/product-names/refresh")
    wb = load_workbook(io.BytesIO(client.get(f"/api/forecast-runs/{fr}/export.xlsx?horizon=7&lang=vi").content))
    tot, daily = wb.worksheets[1], wb.worksheets[2]
    assert [c.value for c in tot[1]][:3] == ["SKU", "Tên sản phẩm", "FC/kho"]
    rows = list(daily.iter_rows(min_row=2, values_only=True))
    with db.read() as c:
        stored = {(r[0], r[1], r[2]): r[3] for r in c.execute(
            "SELECT sku_id, fc_id, target_date, yhat FROM forecast_values WHERE run_id=?", (fr,))}
    assert len(rows) == len(stored)
    for sku, pname, fc, d, *_ in rows:
        assert pname == {"S0": "Hỗ trợ gan Simarin Plus Stella (H/30V)", "S1": "Diamicron MR 30mg Servier"}.get(sku, "Chưa có tên sản phẩm")
    t = {(r[0], r[2]): r[4] for r in tot.iter_rows(min_row=2, values_only=True)}
    assert sum(t.values()) == pytest.approx(_summary(client, fr)["grand_total"])
    info = {r[0]: r[1] for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True)}
    assert info["Phiên bản định dạng file"] == "5" and info["Nguồn tên sản phẩm"].endswith(".dim_sku")


def test_a_d1_refreshes_the_names_and_a_name_failure_never_fails_it(client):
    from app import sync
    from test_v116 import _daily
    fake = _bq_ready_production()
    y = timeutil.today_vn() - timedelta(days=1)
    fake.responses["daily"] = _daily(y, 4.0)
    fake.errors["product_names"] = RuntimeError("name table unavailable")
    assert sync.run_d1_sync("J1", fallback_mode=True)["status"] == "SUCCEEDED"
    with db.read() as c:
        assert product_names.status(c) is None
    del fake.errors["product_names"]
    fake.responses["product_names"] = [{"sku": "S1", "name": "Tên S1"}]
    assert sync.run_d1_sync("J2", fallback_mode=True)["status"] == "SUCCEEDED"
    with db.read() as c:
        assert product_names.names(c) == {"S1": "Tên S1"}


def test_normalize_removes_vietnamese_accents():
    assert product_names.normalize("  Hỗ trợ GAN   Đà Nẵng ") == "ho tro gan da nang"
    assert product_names.matching_skus(["A1", "B2"], "nang", {"A1": "Đà Nẵng"}) == ["A1"]


# ------------------------------------------------------------------ C. D-1 / D-0 switches
def test_d1_and_d0_switches_persist_are_restored_and_never_run_a_past_slot(client):
    fake = _bq_ready_production()
    fake.dry_schema = D0_OUT
    times_before = ss.get("schedule").model_dump()
    assert client.post("/api/sql-sources/d0/versions/1/validate").json()["validation"]["status"] == "VALID"
    assert flows.readiness("d1")["ready"] and flows.readiness("d0")["ready"]
    for fl in ("d1", "d0"):
        r = client.put(f"/api/schedules/{fl}", json={"enabled": True})
        assert r.status_code == 200 and r.json()["enabled"] is True and r.json()["enabled_at"]
    st = client.get("/api/scheduler/status").json()["flow_states"]
    assert st["d1"]["enabled"] is True and st["d0"]["enabled"] is True
    s = ss.get("schedule")                                                   # what a restart reads back
    assert (s.d1_enabled, s.d0_enabled) == (True, True) and s.d0_times == ["08:00", "10:00", "12:00", "14:00", "16:00", "18:00"]
    assert all(getattr(s, k) == times_before[k] for k in ("d1_signal_poll_start", "d1_fallback_time", "d0_times", "train_time"))
    # slots before the switch-on moment are never run or caught up
    from app.jobs.scheduler import plan
    on_at = timeutil.parse_iso(s.d0_enabled_at)
    with db.read() as c:
        acts = plan(on_at + timedelta(minutes=1), s, "BIGQUERY", None, c)
    assert all(a.get("scheduled_for") is None or timeutil.parse_iso(a["scheduled_for"]) >= on_at for a in acts if a["job_type"] == "D0_SYNC")
    # off stays off (nothing switches it back on), on again works
    assert client.put("/api/schedules/d0", json={"enabled": False}).json()["enabled"] is False
    assert ss.get("schedule").d0_enabled is False and ss.get("schedule").d1_enabled is True
    assert client.put("/api/schedules/d0", json={"enabled": True}).json()["enabled"] is True


def test_a_switch_is_refused_when_the_flow_is_not_ready_and_nothing_is_saved(client):
    _bq_ready_production()
    ss.update("bigquery", {"production_enabled": False}, note="production mode disabled")
    before = ss.get("schedule").model_dump()
    for fl in ("d1", "d0"):
        r = client.put(f"/api/schedules/{fl}", json={"enabled": True})
        assert r.status_code == 409 and r.json()["code"] == "FLOW_NOT_READY"
        assert "SOURCE_NOT_READY" in r.json()["detail"]
    assert ss.get("schedule").model_dump() == before
