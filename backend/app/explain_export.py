"""Excel exports of the explanation: demand history of one forecast row, and the backtest evaluation points.
Dates are real Excel dates shown DD/MM/YYYY, quantities stay numbers, header row frozen with an auto-filter."""
from __future__ import annotations

import io
import json
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .export_i18n import fmt_date, fmt_vn_time, labels
from .ml.methods import label as method_label

WEEKDAYS = {"vi": ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật"],
            "en": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]}
STATUS = {"vi": {"ORDERS": "Có đơn", "ZERO_ORDERS": "Có dữ liệu, 0 đơn", "NO_DATA": "Không có dữ liệu (không rõ)",
                 "BEFORE_START": "Trước khi chuỗi bắt đầu (không rõ)"},
          "en": {"ORDERS": "Orders", "ZERO_ORDERS": "Data present, 0 orders", "NO_DATA": "No data (unknown)",
                 "BEFORE_START": "Before the series started (unknown)"}}
T = {
    "vi": {"hist_sheet": "Lịch sử demand", "info_sheet": "Thông tin", "field": "Trường", "value": "Giá trị",
           "cols": ["Ngày", "Thứ", "Demand (số lượng đặt)", "Trạng thái dữ liệu", "Dùng trong phép tính", "Feature dùng ngày này", "Ghi chú"],
           "yes": "Có", "no": "Không",
           "note_unknown": "Không rõ: không tính là 0, không nằm trong tổng hay mẫu số",
           "note_outside": "Ngoài cửa sổ của phép tính", "note_used": "Được dùng",
           "info": ["Mã lần dự báo (run ID)", "Phiên bản mô hình", "SKU", "FC/kho", "Ngày dự báo", "Phương pháp",
                    "Mã kỹ thuật phương pháp", "Dữ liệu tính đến ngày (cutoff)", "Lịch sử từ", "Lịch sử đến",
                    "Ngày dùng trong phép tính", "Tổng demand các ngày dùng", "Số ngày đã biết được dùng",
                    "Giá trị dự báo đã publish", "Múi giờ nghiệp vụ", "Ghi chú"],
           "tz": "Asia/Ho_Chi_Minh (giờ Việt Nam). Ngày là ngày lịch Việt Nam.",
           "note": "Chỉ có ngày đến hết ngày dữ liệu cuối (không có dữ liệu tương lai). Demand trống = không rõ, không phải 0. Số lượng tổng hợp theo ngày từ dòng đơn (không dùng snapshot trong ngày).",
           "bt_sheet": "Điểm đánh giá", "bt_summary": "Tóm tắt",
           "bt_cols": ["Fold", "Mốc fold (ngày đầu dự báo)", "Ngày", "Thứ", "SKU", "FC/kho", "Ngày thứ trong fold",
                       "Phương pháp", "Mã kỹ thuật phương pháp", "Thực tế", "Dự báo", "Sai số tuyệt đối |A−F|", "Baseline TB 28 ngày"],
           "bt_info": ["Training run", "Backtest", "Số điểm đánh giá", "Σ thực tế (ΣA)", "Σ dự báo (ΣF)", "Σ|A−F|",
                       "WAPE = Σ|A−F| ÷ ΣA", "Bias = (ΣF − ΣA) ÷ ΣA (dương = dự báo dư)", "Khớp chỉ số đã lưu",
                       "Số SKU×FC×ngày khác nhau", "SKU×FC×ngày xuất hiện ở ≥ 2 fold", "Cấu hình (fold × ngày / bước)", "Ghi chú"],
           "bt_note": "Mỗi điểm là một cặp (thực tế, dự báo) của một SKU×FC×ngày trong một fold; fold chồng ngày nên một SKU×FC×ngày có thể xuất hiện ở 2 fold (do 2 mô hình khác nhau dự báo). Chỉ số được tái lập từ đúng cấu hình và dữ liệu của lần huấn luyện và khớp với chỉ số đã lưu."},
    "en": {"hist_sheet": "Demand history", "info_sheet": "Info", "field": "Field", "value": "Value",
           "cols": ["Date", "Weekday", "Demand (ordered qty)", "Data status", "Used in calculation", "Features reading this day", "Note"],
           "yes": "Yes", "no": "No",
           "note_unknown": "Unknown: not counted as 0, not in the sum or the denominator",
           "note_outside": "Outside the calculation window", "note_used": "Used",
           "info": ["Forecast run ID", "Model version", "SKU", "FC/warehouse", "Target date", "Method", "Technical method code",
                    "Data as of (cutoff)", "History from", "History to", "Days used in the calculation",
                    "Total demand of the days used", "Known days used", "Published forecast value", "Business time zone", "Note"],
           "tz": "Asia/Ho_Chi_Minh (Vietnam time). Dates are Vietnam calendar days.",
           "note": "Only days up to the last data day (no future data). Empty demand = unknown, not 0. Daily quantities aggregated from order lines (intraday snapshots are not used).",
           "bt_sheet": "Evaluation points", "bt_summary": "Summary",
           "bt_cols": ["Fold", "Fold origin (first forecast day)", "Date", "Weekday", "SKU", "FC/warehouse", "Day of fold",
                       "Method", "Technical method code", "Actual", "Forecast", "Absolute error |A−F|", "28-day mean baseline"],
           "bt_info": ["Training run", "Backtest", "Evaluation points", "Σ actual (ΣA)", "Σ forecast (ΣF)", "Σ|A−F|",
                       "WAPE = Σ|A−F| ÷ ΣA", "Bias = (ΣF − ΣA) ÷ ΣA (positive = over-forecast)", "Matches stored metrics",
                       "Distinct SKU×FC×dates", "SKU×FC×dates in ≥ 2 folds", "Configuration (folds × days / step)", "Note"],
           "bt_note": "Each point is one (actual, forecast) pair of one SKU×FC×date in one fold; folds overlap, so a SKU×FC×date can appear in 2 folds (forecast by 2 different models). Metrics were recomputed with the training run's exact configuration and data and match the stored metrics."},
}


def _finish(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w


def _info_sheet(wb, title, L, rows):
    ws = wb.active
    ws.title = title
    ws.append([L["field"], L["value"]])
    for k, v in rows:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 90
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in ws.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top")


def history_xlsx(ex: dict, lang: str) -> bytes:
    lang, L = labels(lang)
    X = T[lang]
    h = ex["history"]
    idn = ex["identity"]
    calc = h.get("calc") or {}
    wb = Workbook()
    _info_sheet(wb, X["info_sheet"], L, list(zip(X["info"], [
        idn["run_id"], idn["model_version"], idn["sku_id"], idn["fc_id"], fmt_date(idn["target_date"], L),
        method_label(ex["result"]["method"], lang), ex["result"]["method"], fmt_date(idn["data_as_of"], L),
        fmt_date(h.get("window_start"), L), fmt_date(h.get("window_end"), L),
        f"{fmt_date(calc.get('from'), L)} → {fmt_date(calc.get('to'), L)}" if calc.get("from") else "",
        calc.get("sum"), calc.get("known_days"), ex["result"]["yhat"], X["tz"], X["note"]])))
    ws = wb.create_sheet(X["hist_sheet"])
    ws.append(X["cols"])
    for d in h["days"]:
        known = d["qty"] is not None
        if not known:
            note = X["note_unknown"]
        elif d["in_calc"]:
            note = X["note_used"]
        else:
            note = X["note_outside"] if calc else ""
        ws.append([date.fromisoformat(d["date"]), WEEKDAYS[lang][d["weekday"]], d["qty"], STATUS[lang][d["status"]],
                   X["yes"] if d["in_calc"] else X["no"], ", ".join(d.get("used_by") or []), note])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
    _finish(ws, [13, 12, 20, 32, 18, 60, 60])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def backtest_points_xlsx(res: dict, lang: str) -> bytes:
    lang, L = labels(lang)
    X = T[lang]
    m = res["metrics"]
    sabs = sum(abs(p["actual"] - p["forecast"]) for p in res["points"])
    cfg = res["config"]
    wb = Workbook()
    _info_sheet(wb, X["bt_summary"], L, list(zip(X["bt_info"], [
        res["training_run_id"], res["backtest_id"], m["n_points"], m["sum_actual"], m["sum_forecast"], sabs,
        m["wape"], m["bias"], X["yes"], res["unique_series_days"], res["series_days_in_two_or_more_folds"],
        f"{cfg['folds']} × {cfg['horizon_days']} / {cfg['step_days']}", X["bt_note"]])))
    ws = wb.create_sheet(X["bt_sheet"])
    ws.append(X["bt_cols"])
    for p in res["points"]:
        d = date.fromisoformat(p["date"])
        ws.append([p["fold"], date.fromisoformat(p["origin"]), d, WEEKDAYS[lang][d.weekday()], p["sku_id"], p["fc_id"],
                   p["day_of_fold"], method_label(p["method"], lang), p["method"], p["actual"], p["forecast"],
                   abs(p["actual"] - p["forecast"]), p["baseline"]])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=2).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=3).number_format = "DD/MM/YYYY"
    _finish(ws, [6, 16, 12, 11, 16, 12, 10, 30, 26, 10, 12, 14, 14])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_ = fmt_vn_time   # (kept for symmetry with the other exports)


# ---------------------------------------------------------------- TreeSHAP audit trace
TR = {
    "vi": {"info": "Thông tin & đối soát", "features": "SHAP 33 feature", "trees": "Theo cây", "terms": "Thành phần",
           "info_cols": ["Mục", "Giá trị"],
           "feat_cols": ["Feature", "Giá trị đầu vào", "SHAP (trace engine)", "SHAP (LightGBM pred_contrib)", "Sai lệch tuyệt đối", "Đạt (≤ dung sai)"],
           "tree_cols": ["Feature", "Cây", "Tổng thành phần trong cây", "Phần dương", "Phần âm", "Số thành phần"],
           "term_cols": ["Feature", "Cây", "Lá", "Giá trị lá", "Dòng này rơi vào lá này", "one_fraction", "zero_fraction",
                         "Trọng số Shapley w", "one − zero", "Thành phần = w × (one − zero) × giá trị lá", "Đường đi (* = nhánh dòng này KHÔNG đi)"]},
    "en": {"info": "Info & reconciliation", "features": "SHAP of 33 features", "trees": "By tree", "terms": "Terms",
           "info_cols": ["Item", "Value"],
           "feat_cols": ["Feature", "Input value", "SHAP (trace engine)", "SHAP (LightGBM pred_contrib)", "Absolute difference", "Within tolerance"],
           "tree_cols": ["Feature", "Tree", "Sum of terms in the tree", "Positive part", "Negative part", "Number of terms"],
           "term_cols": ["Feature", "Tree", "Leaf", "Leaf value", "Row reaches this leaf", "one_fraction", "zero_fraction",
                         "Shapley weight w", "one − zero", "Term = w × (one − zero) × leaf value", "Route (* = branch this row does NOT take)"]},
}


def shap_trace_rows(res: dict, feature: str | None):
    from .ml.shap_audit import export_rows
    return export_rows(res, feature)


def shap_trace_xlsx(res: dict, feature: str | None, lang: str) -> bytes:
    from .ml.shap_audit import TOL_SHAP_ABS
    lang = lang if lang in TR else "vi"
    X = TR[lang]
    idn, vals = res["identity"], res["values"]
    wb = Workbook()
    ws = wb.active
    ws.title = X["info"]
    ws.append(X["info_cols"])
    rows = [("status", res["status"]), ("forecast_run_id", idn["run_id"]), ("sku_id", idn["sku_id"]), ("fc_id", idn["fc_id"]),
            ("target_date", idn["target_date"]), ("model_version", idn["model_version"]), ("model_file", idn["model_file"]),
            ("model_sha256", idn["model_sha256_file"]), ("feature_row_hash (sha256 of float32 row)", idn["feature_row_hash"]),
            ("objective / link", f"{idn['objective']} / {idn['link']}"), ("n_trees", idn["n_trees"]), ("feature filter", feature or "ALL"),
            ("engine", "port of LightGBM v4.7.0 Tree::TreeSHAP (src/io/tree.cpp)")]
    rows += [(k, v) for k, v in vals.items()]
    for c in res["checks"]:
        rows.append((f"check {c['check']}", f"{'PASS' if c['passed'] else 'FAIL'} " + json.dumps({k: v for k, v in c.items() if k not in ('check', 'passed')}, default=str)))
    for r in rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 110
    ws2 = wb.create_sheet(X["features"])
    ws2.append(X["feat_cols"])
    for f in res["features"]:
        ws2.append([f["name"], f["value"], f["shap_engine"], f["shap_lightgbm"], f["abs_diff"], f["abs_diff"] <= TOL_SHAP_ABS])
    terms = shap_trace_rows(res, feature)
    agg: dict = {}
    for t in terms:
        a = agg.setdefault((t[0], t[1]), [0.0, 0.0, 0.0, 0])
        a[0] += t[9]
        a[1 if t[9] > 0 else 2] += t[9] if t[9] != 0 else 0.0
        a[3] += 1
    ws3 = wb.create_sheet(X["trees"])
    ws3.append(X["tree_cols"])
    for (f, ti), a in sorted(agg.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        ws3.append([f, ti, a[0], a[1], a[2], a[3]])
    ws4 = wb.create_sheet(X["terms"])
    ws4.append(X["term_cols"])
    for t in terms:
        ws4.append(list(t))
    for w_, widths in ((ws2, [24, 16, 24, 24, 18, 14]), (ws3, [24, 8, 22, 18, 18, 12]), (ws4, [22, 7, 6, 20, 10, 12, 22, 22, 10, 24, 120])):
        _finish(w_, widths)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def shap_trace_csv(res: dict, feature: str | None) -> bytes:
    import csv
    idn = res["identity"]
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["status", "forecast_run_id", "sku_id", "fc_id", "target_date", "model_sha256", "feature_row_hash", "base_value",
                "raw_score_from_model", "feature", "tree", "leaf", "leaf_value", "row_reaches_leaf", "one_fraction", "zero_fraction",
                "weight", "one_minus_zero", "term", "route"])
    for t in shap_trace_rows(res, feature):
        w.writerow([res["status"], idn["run_id"], idn["sku_id"], idn["fc_id"], idn["target_date"], idn["model_sha256_file"],
                    idn["feature_row_hash"], repr(res["values"]["base_value"]), repr(res["values"]["raw_score_from_model"]),
                    t[0], t[1], t[2], repr(t[3]), t[4], repr(t[5]), repr(t[6]), repr(t[7]), repr(t[8]), repr(t[9]), t[10]])
    return out.getvalue().encode("utf-8-sig")
