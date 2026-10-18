"""Forecast & actual history (1.2.5): the D+1 forecast of every forecast session, compared with the actual demand of
that day. Rules confirmed by the owner; full description in docs/LICH-SU-FORECAST.md.

  session    one published forecast run (forecast_runs). Only fully published runs exist: a running or failed training,
             extension or promotion leaves no run, so it can never become history. Session time = start of the job that
             produced the run, else its publish time; run date = the Vietnam date of the session time.
  D+1 only   the values of target_date = run date + 1 are copied to forecast_history_values; the other horizons stay in
             forecast_values only (multi-day forecasts are unchanged).
  official   per scope and run date ONE session: the newest eligible one by (session time, publish time, run rowid).
             A session that started earlier but finished later never replaces a newer one; a session without a value for
             D+1 is NOT_ELIGIBLE and replaces nothing; a newer eligible session of the SAME run date does replace it.
  immutable  sessions and values are never updated or deleted (triggers, migration 0004); only the per-day pointer
             moves, each move is audited (FORECAST_HISTORY_SELECTED).
  actual     read when shown, never stored here and never written back into a forecast (see actual_context):
               D-1 data (demand_daily) when the target day is loaded: a SKU x FC without a row = 0 (the day was read for
               the whole dataset - the same rule as the training panel); a loaded day without any row = unknown;
               else the latest successful D0 snapshot of that day (cumulative so far);
               else none ("no actual data yet") - never 0.
             COMPLETE only when the D-1 load of that day ran on the configured completion signal (bigquery.completion_table);
             1.2.6: a day closed at 23:59 (actual_close.py) without that signal is CLOSED ("closed - needs
             verification"); every other actual is PROVISIONAL. The figure always comes from the source above.
  backup     1.2.6: a run date WITHOUT an official session shows the D+1 of the newest COMPLETE end-of-day snapshot that
             covers it (forecast_backup.py, fallback_days), labelled SNAPSHOT with its snapshot and source run; an
             official session always wins. No snapshot covering it -> no forecast for that day (a gap, never 0).
Recording runs right after a run is published, outside its transaction (a history failure can never undo a forecast),
and is retried idempotently from the immutable runs at the next publish, at app start and by
`python -m app.cli history-reconcile`. Showing the history only reads.
"""
from __future__ import annotations

import io
import json
import logging
from datetime import date, timedelta

from . import db
from .messages import AppError
from .timeutil import VN_TZ, now_iso, parse_iso, today_vn

log = logging.getLogger("df.history")

SELECTED, OLDER, NOT_ELIGIBLE = "SELECTED", "OLDER_THAN_SELECTED", "NOT_ELIGIBLE"
COMPLETE, CLOSED, PROVISIONAL, NO_ACTUAL = "COMPLETE", "CLOSED", "PROVISIONAL", "NO_ACTUAL"
ACTUAL_STATUSES = (COMPLETE, CLOSED, PROVISIONAL, NO_ACTUAL)
PRIMARY, SNAPSHOT = "PRIMARY", "SNAPSHOT"          # 1.2.6 forecast source of a run date: official session / EOD backup
FORECAST_SOURCES = (PRIMARY, SNAPSHOT)
SORTS = {
    "date": "b.run_date DESC, b.yhat DESC, b.sku_id, b.fc_id",
    "abs_error": "ABS(b.yhat - b.actual) DESC NULLS LAST, b.run_date DESC, b.sku_id, b.fc_id",
    "forecast": "b.yhat DESC, b.run_date DESC, b.sku_id, b.fc_id",
    "actual": "b.actual DESC NULLS LAST, b.run_date DESC, b.sku_id, b.fc_id",
    "sku": "b.sku_id, b.fc_id, b.run_date DESC",
}
MAX_EXPORT_ROWS = 1_000_000          # one Excel sheet holds 1,048,576 rows: larger selections are refused, never cut
MAX_RANGE_DAYS = 400


def _vn_date(ts: str) -> date:
    return parse_iso(ts).astimezone(VN_TZ).date()


def _usage(c, training_run_id: str) -> str | None:
    r = db.one(c, "SELECT quality_gate_json FROM training_runs WHERE id=?", (training_run_id,))
    try:
        return ((json.loads((r or {}).get("quality_gate_json") or "{}").get("assessment") or {}).get("usage") or {}).get("status")
    except ValueError:
        return None


def _names(c, data_source: str) -> dict[str, str]:
    """Trusted SKU -> product name mapping (BIGQUERY only; DEMO / Excel SKUs are not the source's SKUs)."""
    if data_source != "BIGQUERY":
        return {}
    from .bq import product_names
    return product_names.names(c)


# ---------------------------------------------------------------- recording
def _record(c, run_id: str, mode: str) -> dict:
    """Inside ONE immediate transaction: record a run (idempotent) and move the day pointer when it is the newest."""
    have = db.one(c, "SELECT * FROM forecast_history_sessions WHERE forecast_run_id=?", (run_id,))
    if have:
        return {**have, "already_recorded": True}
    r = db.one(c, "SELECT r.id, r.rowid AS rid, r.training_run_id, r.model_version, r.scope_key, r.data_source, "
                  "r.data_as_of, r.created_at, COALESCE(j.started_at, r.created_at) AS started "
                  "FROM forecast_runs r LEFT JOIN job_runs j ON j.id = r.job_id WHERE r.id=?", (run_id,))
    if not r:
        raise AppError("BATCH_NOT_FOUND", 404, detail=run_id)
    run_date = _vn_date(r["started"])
    target = str(run_date + timedelta(days=1))
    n, total, hday = c.execute("SELECT COUNT(*), SUM(yhat), MIN(horizon_day) FROM forecast_values WHERE run_id=? "
                               "AND target_date=?", (run_id, target)).fetchone()
    cur = db.one(c, "SELECT d.forecast_run_id, s.session_started_at, s.published_at, fr.rowid AS rid "
                    "FROM forecast_history_days d JOIN forecast_history_sessions s ON s.forecast_run_id = d.forecast_run_id "
                    "JOIN forecast_runs fr ON fr.id = d.forecast_run_id WHERE d.scope_key=? AND d.run_date=?",
                 (r["scope_key"], str(run_date)))
    key = (parse_iso(r["started"]), parse_iso(r["created_at"]), r["rid"])
    if not n:
        decision, reason = NOT_ELIGIBLE, "NO_D1_VALUES"
    elif cur and (parse_iso(cur["session_started_at"]), parse_iso(cur["published_at"]), cur["rid"]) >= key:
        decision, reason = OLDER, f"NEWER_SESSION_SELECTED:{cur['forecast_run_id']}"
    else:
        decision, reason = SELECTED, (f"REPLACES:{cur['forecast_run_id']}" if cur else "FIRST_SESSION_OF_DAY")
    at = now_iso()
    c.execute("INSERT INTO forecast_history_sessions(forecast_run_id, scope_key, data_source, run_date, target_date, "
              "session_started_at, published_at, training_run_id, model_version, data_as_of, horizon_day, n_rows, "
              "total_yhat, usage_status, eligible, decision, decision_reason, capture_mode, captured_at) "
              "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
              (run_id, r["scope_key"], r["data_source"], str(run_date), target, r["started"], r["created_at"],
               r["training_run_id"], r["model_version"], r["data_as_of"], hday, n, total,
               _usage(c, r["training_run_id"]), int(bool(n)), decision, reason, mode, at))
    if decision == SELECTED:
        names = _names(c, r["data_source"])
        vals = [(run_id, s, f, r["scope_key"], str(run_date), target, y, m, names.get(s))
                for s, f, y, m in c.execute("SELECT sku_id, fc_id, yhat, method FROM forecast_values WHERE run_id=? "
                                            "AND target_date=?", (run_id, target))]
        c.executemany("INSERT INTO forecast_history_values(forecast_run_id, sku_id, fc_id, scope_key, run_date, "
                      "target_date, yhat, method, product_name) VALUES (?,?,?,?,?,?,?,?,?)", vals)
        c.execute("INSERT INTO forecast_history_days(scope_key, run_date, data_source, target_date, forecast_run_id, "
                  "selected_at) VALUES (?,?,?,?,?,?) ON CONFLICT(scope_key, run_date) DO UPDATE SET "
                  "forecast_run_id=excluded.forecast_run_id, selected_at=excluded.selected_at",
                  (r["scope_key"], str(run_date), r["data_source"], target, run_id, at))
        db.audit(c, "FORECAST_HISTORY_SELECTED", "forecast_history_day", f"{r['scope_key']}|{run_date}",
                 {"forecast_run_id": run_id, "replaced": (cur or {}).get("forecast_run_id"), "run_date": str(run_date),
                  "target_date": target, "rows": n, "mode": mode})
    return {"forecast_run_id": run_id, "scope_key": r["scope_key"], "run_date": str(run_date), "target_date": target,
            "decision": decision, "decision_reason": reason, "n_rows": n, "capture_mode": mode}


def capture(run_id: str, mode: str = "AT_PUBLISH") -> dict:
    with db.tx() as c:
        return _record(c, run_id, mode)


def reconcile() -> dict:
    """Record every published run that the history has not seen yet (oldest first). Idempotent: running it again
    records nothing. Needs no new forecast: it reads the immutable published runs."""
    with db.read() as c:
        ids = [x[0] for x in c.execute("SELECT r.id FROM forecast_runs r LEFT JOIN forecast_history_sessions s "
                                       "ON s.forecast_run_id = r.id WHERE s.forecast_run_id IS NULL "
                                       "ORDER BY r.created_at, r.rowid")]
    out = {"missing": len(ids), "recorded": [], "errors": []}
    for rid in ids:
        try:
            res = capture(rid, "RECOVERED")
            out["recorded"].append({k: res[k] for k in ("forecast_run_id", "run_date", "decision")})
        except Exception as e:  # noqa: BLE001 - one bad run never stops the others; reported
            log.warning("history: run not recorded", exc_info=True, extra={"ctx": {"forecast_run_id": rid}})
            out["errors"].append({"forecast_run_id": rid, "error": f"{type(e).__name__}: {e}"[:300]})
    return out


def capture_after_publish(run_id: str) -> dict | None:
    """Called right after a forecast run was committed. Never raises: the forecast is published either way. A failure
    is logged and raised as an ERROR alert; reconcile() records the run later from the immutable published run."""
    try:
        res = capture(run_id, "AT_PUBLISH")
    except Exception as e:  # noqa: BLE001
        log.exception("history: recording failed (forecast stays published)", extra={"ctx": {"forecast_run_id": run_id}})
        try:
            with db.tx() as c:
                db.alert(c, "ERROR", "FORECAST_HISTORY_CAPTURE_FAILED",
                         {"forecast_run_id": run_id, "error": f"{type(e).__name__}: {e}"[:500]}, "forecast_run", run_id)
        except Exception:  # noqa: BLE001
            log.exception("history: alert not written")
        return None
    try:
        rec = reconcile()          # earlier runs whose recording had failed
        if rec["missing"]:
            log.info("history: earlier runs recorded", extra={"ctx": {"count": rec["missing"]}})
    except Exception:  # noqa: BLE001
        log.warning("history: reconcile failed", exc_info=True)
    log.info("history: session recorded", extra={"ctx": {k: res.get(k) for k in ("forecast_run_id", "run_date", "decision")}})
    return res


# ---------------------------------------------------------------- reading
def scopes(c) -> list[dict]:
    return db.rows(c, "SELECT scope_key, data_source, COUNT(*) AS n_days, MIN(run_date) AS first_run_date, "
                      "MAX(run_date) AS last_run_date FROM forecast_history_days GROUP BY scope_key, data_source "
                      "ORDER BY MAX(run_date) DESC")


def days(c, scope_key: str) -> list[dict]:
    """Run dates of a scope with their official session (newest first)."""
    return db.rows(c, "SELECT d.run_date, d.target_date, d.forecast_run_id, d.selected_at, s.session_started_at, "
                      "s.published_at, s.training_run_id, s.model_version, s.data_as_of, s.horizon_day, s.n_rows, "
                      "s.total_yhat, s.usage_status, s.capture_mode, s.decision_reason, "
                      "(SELECT COUNT(*) FROM forecast_history_sessions x WHERE x.scope_key = d.scope_key "
                      " AND x.run_date = d.run_date) AS sessions_count "
                      "FROM forecast_history_days d JOIN forecast_history_sessions s ON s.forecast_run_id = d.forecast_run_id "
                      "WHERE d.scope_key=? ORDER BY d.run_date DESC", (scope_key,))


def sessions(c, scope_key: str, run_from: str, run_to: str) -> list[dict]:
    """Every recorded session of the run dates (official or not) with the decision taken: the trace of the choice."""
    return db.rows(c, "SELECT s.forecast_run_id, s.run_date, s.target_date, s.session_started_at, s.published_at, "
                      "s.model_version, s.n_rows, s.eligible, s.decision, s.decision_reason, s.capture_mode, "
                      "s.captured_at, s.usage_status, (d.forecast_run_id = s.forecast_run_id) AS official "
                      "FROM forecast_history_sessions s LEFT JOIN forecast_history_days d ON d.scope_key = s.scope_key "
                      "AND d.run_date = s.run_date WHERE s.scope_key=? AND s.run_date BETWEEN ? AND ? "
                      "ORDER BY s.run_date DESC, s.session_started_at DESC, s.published_at DESC, s.rowid DESC",
                   (scope_key, run_from, run_to))


def fallback_days(c, scope_key: str, today: date | None = None) -> list[dict]:
    """1.2.6: run dates (up to today) WITHOUT an official session whose D+1 is covered by a COMPLETE end-of-day snapshot -
    the newest one, i.e. the largest business date <= the run date that still covers run date + 1 (a snapshot of day S
    covers S+1 .. S+7). Newest first. Reading only; an INCOMPLETE / NO_SOURCE snapshot is never used."""
    today = today or today_vn()
    snaps = db.rows(c, "SELECT id, business_date, last_target_date, created_at, source_run_id, source_run_date, "
                       "source_session_started_at, source_published_at, source_model_version FROM forecast_eod_snapshots "
                       "WHERE scope_key=? AND status='COMPLETE' ORDER BY business_date DESC", (scope_key,))
    if not snaps:
        return []
    primary = {r[0] for r in c.execute("SELECT run_date FROM forecast_history_days WHERE scope_key=?", (scope_key,))}
    out = []
    d = min(today, date.fromisoformat(snaps[0]["last_target_date"]) - timedelta(days=1))
    first = date.fromisoformat(snaps[-1]["business_date"])
    while d >= first:
        run_date, target = str(d), str(d + timedelta(days=1))
        s = next((x for x in snaps if x["business_date"] <= run_date), None)       # newest snapshot taken by that day
        if run_date not in primary and s and s["last_target_date"] >= target:
            out.append({"run_date": run_date, "target_date": target, "snapshot_id": s["id"],
                        "snapshot_business_date": s["business_date"], "snapshot_created_at": s["created_at"],
                        "forecast_run_id": s["source_run_id"], "source_run_date": s["source_run_date"],
                        "session_started_at": s["source_session_started_at"], "published_at": s["source_published_at"],
                        "model_version": s["source_model_version"]})
        d -= timedelta(days=1)
    return out


def _d1_state(c, scope_key: str) -> dict | None:
    r = db.one(c, "SELECT value_json FROM sync_state WHERE name=?", (f"last_d1_sync:{scope_key}",))
    return json.loads(r["value_json"]) if r else None


def signal_days(c, scope_key: str) -> set[str]:
    """Days whose D-1 load started on the source's completion signal (audit of each D-1 sync: mode SIGNAL, watermark =
    the day the signal was checked for). Fallback-time loads (row-count check only) and backfills are not proof."""
    out = set()
    for (dj,) in c.execute("SELECT details_json FROM audit_log WHERE action='D1_SYNC' AND entity_id=?", (scope_key,)):
        try:
            d = json.loads(dj or "{}")
        except ValueError:
            continue
        if d.get("mode") == "SIGNAL" and not d.get("backfill") and d.get("watermark"):
            out.add(str(d["watermark"]))
    return out


def actual_context(c, scope_key: str, data_source: str, dates: list[str], today: date | None = None,
                   closes: bool = True) -> dict[str, dict]:
    """Per target date: where its actual comes from and how final it is (see the module docstring). closes=False: the
    source rule alone (used by the close itself)."""
    today = today or today_vn()
    if not dates:
        return {}
    out = _source_context(c, scope_key, data_source, dates, today)
    if closes:
        _close_overlay(c, scope_key, out)
    return out


def _close_overlay(c, scope_key: str, out: dict[str, dict]) -> None:
    """1.2.6: the latest close revision of each day (actual_closes). A closed day that the completion signal has not
    confirmed is CLOSED instead of PROVISIONAL; COMPLETE and NO_ACTUAL stay as they are. `close.current` is False when
    the figure shown now comes from another source (or another D0 read) than the latest revision - a change that the
    next reconciliation (after each D-1 load and in each close job) records as a new revision."""
    for x in out.values():
        x["close"] = None
    days = list(out)
    for r in db.rows(c, "SELECT a.* FROM actual_closes a WHERE a.scope_key=? AND a.business_date IN "
                        f"({','.join('?' * len(days))}) AND a.revision = (SELECT MAX(b.revision) FROM actual_closes b "
                        "WHERE b.scope_key = a.scope_key AND b.business_date = a.business_date)", (scope_key, *days)):
        x = out[r["business_date"]]
        x["close"] = {"id": r["id"], "revision": r["revision"], "status": r["status"], "closed_at": r["closed_at"],
                      "cutoff_at": r["cutoff_at"], "source": r["source"], "total_qty": r["total_qty"],
                      "current": r["source"] == x["source"] and (x["source"] != "D0" or r["source_ref"] == x["snapshot_id"])}
        if x["status"] == PROVISIONAL:
            x["status"], x["reason"] = CLOSED, f"CLOSED_{x['source']}"


def _source_context(c, scope_key: str, data_source: str, dates: list[str], today: date) -> dict[str, dict]:
    st = _d1_state(c, scope_key) if data_source == "BIGQUERY" else None
    loaded_through = (st or {}).get("watermark") or c.execute(
        "SELECT MAX(order_date) FROM demand_daily WHERE scope_key=?", (scope_key,)).fetchone()[0]
    sig = signal_days(c, scope_key) if data_source == "BIGQUERY" else set()
    out = {}
    for d in sorted(set(dates)):
        x = {"date": d, "source": None, "snapshot_id": None, "as_of": None}
        if d > str(today):
            out[d] = {**x, "status": NO_ACTUAL, "reason": "FUTURE_DATE"}
            continue
        row = db.one(c, "SELECT loaded_at FROM demand_daily WHERE scope_key=? AND order_date=? LIMIT 1", (scope_key, d))
        if row:
            done = d in sig
            out[d] = {**x, "source": "D1", "as_of": row["loaded_at"], "status": COMPLETE if done else PROVISIONAL,
                      "reason": "COMPLETION_SIGNAL" if done else "D1_NO_COMPLETION_SIGNAL"}
            continue
        if loaded_through and d <= loaded_through and d < str(today):
            out[d] = {**x, "status": NO_ACTUAL, "reason": "DAY_WITHOUT_DATA"}
            continue
        snap = db.one(c, "SELECT id, snapshot_at FROM d0_snapshots WHERE scope_key=? AND order_date=? AND "
                         "sync_status='SUCCEEDED' ORDER BY snapshot_at DESC, rowid DESC LIMIT 1", (scope_key, d))
        if snap:
            out[d] = {**x, "source": "D0", "snapshot_id": snap["id"], "as_of": snap["snapshot_at"],
                      "status": PROVISIONAL, "reason": "D0_INTRADAY" if d == str(today) else "D0_LAST_READ"}
        else:
            out[d] = {**x, "status": NO_ACTUAL, "reason": "NO_D0_READ_TODAY" if d == str(today) else "NOT_LOADED_YET"}
    return out


def _range(run_from: str | None, run_to: str | None, dlist: list[dict]) -> tuple[str, str]:
    """Run-date range of a request; default = the last 30 days ending at the newest run date with a stored forecast
    (official session or, 1.2.6, end-of-day backup)."""
    last = dlist[0]["run_date"] if dlist else str(today_vn())
    try:
        end = date.fromisoformat(run_to or last)
        start = date.fromisoformat(run_from) if run_from else end - timedelta(days=29)
    except ValueError:
        raise AppError("VALIDATION_ERROR", 422, detail="run_from / run_to must be dates YYYY-MM-DD") from None
    if start > end:
        raise AppError("VALIDATION_ERROR", 422, detail="run_from must not be after run_to")
    if (end - start).days + 1 > MAX_RANGE_DAYS:
        raise AppError("VALIDATION_ERROR", 422, detail=f"at most {MAX_RANGE_DAYS} run dates per request")
    return str(start), str(end)


class Query:
    """One filtered view of a scope's history: KPI, chart, table and export all use exactly this row set."""

    def __init__(self, c, scope_key: str, run_from: str | None = None, run_to: str | None = None,
                 fc_id: str | None = None, q: str | None = None, actual_status: str | None = None,
                 today: date | None = None, forecast_source: str | None = None):
        self.c = c
        self.today = today or today_vn()
        sc = next((s for s in scopes(c) if s["scope_key"] == scope_key), None)
        self.scope_key = scope_key
        self.data_source = sc["data_source"] if sc else None
        self.days = days(c, scope_key) if sc else []
        self.backup = fallback_days(c, scope_key, self.today) if sc else []
        self.run_from, self.run_to = _range(run_from, run_to, sorted(self.days + self.backup,
                                                                     key=lambda d: d["run_date"], reverse=True))
        statuses = [s for s in (actual_status or "").split(",") if s]
        if any(s not in ACTUAL_STATUSES for s in statuses):
            raise AppError("VALIDATION_ERROR", 422, detail=f"actual_status must be one of {', '.join(ACTUAL_STATUSES)}")
        if forecast_source and forecast_source not in FORECAST_SOURCES:
            raise AppError("VALIDATION_ERROR", 422, detail=f"forecast_source must be one of {', '.join(FORECAST_SOURCES)}")
        self.filters = {"run_from": self.run_from, "run_to": self.run_to, "fc_id": fc_id or None,
                        "q": (q or "").strip() or None, "actual_status": statuses or None,
                        "forecast_source": forecast_source or None}
        # run dates of the range by forecast source (whatever the filters: a filtered-out day is not a missing day)
        self.in_range = [d for d in self.days if self.run_from <= d["run_date"] <= self.run_to]
        self.backup_in_range = [d for d in self.backup if self.run_from <= d["run_date"] <= self.run_to]
        self.use_primary = self.in_range if forecast_source != SNAPSHOT else []
        self.use_backup = self.backup_in_range if forecast_source != PRIMARY else []
        self.names = _names(c, self.data_source) if self.data_source else {}
        targets = [d["target_date"] for d in self.in_range + self.backup_in_range]
        self.ctx = actual_context(c, scope_key, self.data_source, targets, self.today) if targets else {}
        self.skus = self._search(self.filters["q"])

    def name_of(self, sku: str, stored: str | None) -> str | None:
        return self.names.get(sku) or stored

    def _search(self, q: str | None) -> list[str] | None:
        """SKUs whose SKU or product name contains every word of q (partial, case- and accent-insensitive)."""
        if not q:
            return None
        from .bq.product_names import matching_skus
        stored = {s: n for s, n in self.c.execute(
            "SELECT v.sku_id, MAX(v.product_name) FROM forecast_history_days d JOIN forecast_history_values v "
            "ON v.forecast_run_id = d.forecast_run_id WHERE d.scope_key=? AND d.run_date BETWEEN ? AND ? "
            "GROUP BY v.sku_id", (self.scope_key, self.run_from, self.run_to))}
        # a backup day holds the series of its source run, whose D+1 (every series) is in the history values
        for rid in {d["forecast_run_id"] for d in self.backup_in_range}:
            for s, n in self.c.execute("SELECT sku_id, product_name FROM forecast_history_values WHERE forecast_run_id=?",
                                       (rid,)):
                stored[s] = stored.get(s) or n
        mapping = {s: self.name_of(s, n) or "" for s, n in stored.items()}
        return matching_skus(list(stored), q, mapping)

    def _base(self) -> tuple[str, list]:
        """WITH ctx (actual source per target date), bk (backup days: run date -> snapshot), base (one row per SKU x FC
        x run date: the official session's stored D+1, or for a run date without one, its end-of-day backup)."""
        vals = ",".join("(?,?,?,?,?)" for _ in self.ctx)
        params: list = []
        for d, x in self.ctx.items():
            params += [d, x["source"], x["snapshot_id"], x["status"], x["reason"]]
        where, wp = "", []
        if self.filters["fc_id"]:
            where += " AND v.fc_id=?"
            wp.append(self.filters["fc_id"])
        if self.skus is not None:
            where += " AND v.sku_id IN (SELECT value FROM json_each(?))"
            wp.append(json.dumps(self.skus))
        if self.filters["actual_status"]:
            where += f" AND ctx.st IN ({','.join('?' * len(self.filters['actual_status']))})"
            wp += self.filters["actual_status"]
        actual = ("ctx.st AS actual_status, ctx.reason AS actual_reason, ctx.src AS actual_source, CASE ctx.src WHEN 'D1' "
                  "THEN COALESCE(dd.qty, 0) WHEN 'D0' THEN COALESCE(sv.cumulative_qty, 0) END AS actual ")
        joins = ("JOIN ctx ON ctx.target_date = v.target_date "
                 "LEFT JOIN demand_daily dd ON ctx.src = 'D1' AND dd.scope_key = {scope} AND dd.order_date = v.target_date "
                 "AND dd.sku_id = v.sku_id AND dd.fc_id = v.fc_id "
                 "LEFT JOIN d0_snapshot_values sv ON ctx.src = 'D0' AND sv.snapshot_id = ctx.snap AND sv.sku_id = v.sku_id "
                 "AND sv.fc_id = v.fc_id ")
        primary = ("SELECT d.run_date, v.target_date, v.sku_id, v.fc_id, v.yhat, v.method, v.product_name, d.forecast_run_id, "
                   f"s.model_version, s.session_started_at, s.published_at, {actual}, 'PRIMARY' AS forecast_source, "
                   "NULL AS snapshot_id, NULL AS snapshot_business_date, NULL AS snapshot_created_at "
                   "FROM forecast_history_days d JOIN forecast_history_sessions s ON s.forecast_run_id = d.forecast_run_id "
                   "JOIN forecast_history_values v ON v.forecast_run_id = d.forecast_run_id "
                   + joins.format(scope="v.scope_key")
                   + f"WHERE d.scope_key=? AND d.run_date BETWEEN ? AND ?{where}"
                   + ("" if self.use_primary else " AND 0"))
        pp = [self.scope_key, self.run_from, self.run_to] + wp
        sql = f"WITH ctx(target_date, src, snap, st, reason) AS (VALUES {vals})"
        if self.use_backup:
            sql += f", bk(run_date, target_date, snapshot_id) AS (VALUES {','.join('(?,?,?)' for _ in self.use_backup)})"
            for d in self.use_backup:
                params += [d["run_date"], d["target_date"], d["snapshot_id"]]
            primary += (" UNION ALL SELECT bk.run_date, v.target_date, v.sku_id, v.fc_id, v.yhat, v.method, hv.product_name, "
                        "es.source_run_id, es.source_model_version, es.source_session_started_at, es.source_published_at, "
                        f"{actual}, 'SNAPSHOT', es.id, es.business_date, es.created_at "
                        "FROM bk JOIN forecast_eod_snapshots es ON es.id = bk.snapshot_id "
                        "JOIN forecast_eod_snapshot_values v ON v.snapshot_id = bk.snapshot_id AND v.target_date = bk.target_date "
                        "LEFT JOIN forecast_history_values hv ON hv.forecast_run_id = es.source_run_id "
                        "AND hv.sku_id = v.sku_id AND hv.fc_id = v.fc_id "
                        + joins.format(scope="es.scope_key") + f"WHERE es.scope_key=?{where}")
            pp += [self.scope_key] + wp
        return sql + f", base AS ({primary})", params + pp

    def daily(self) -> list[dict]:
        """One point per target date of the range: forecast total, actual total of the SAME rows (None when the day has
        no actual), deviation, status and forecast source. Run dates without any stored forecast (no official session and
        no backup covering them) are returned as gaps (never filled)."""
        by_date = {}
        if self.ctx:
            sql, p = self._base()
            for r in db.rows(self.c, sql + " SELECT run_date, target_date, actual_status, actual_reason, COUNT(*) AS n_rows, "
                                     "SUM(yhat) AS forecast, SUM(actual) AS actual, COUNT(actual) AS n_actual, "
                                     "SUM(CASE WHEN actual IS NOT NULL THEN yhat END) AS forecast_compared, "
                                     "SUM(ABS(yhat - actual)) AS abs_error FROM base GROUP BY run_date, target_date, "
                                     "actual_status, actual_reason", p):
                by_date[r["run_date"]] = r
        out = []
        d, end = date.fromisoformat(self.run_from), date.fromisoformat(self.run_to)
        sel = {x["run_date"]: x for x in self.in_range}
        bk = {x["run_date"]: x for x in self.backup_in_range}
        while d <= end:
            k = str(d)
            r = by_date.get(k)
            day = sel.get(k) or bk.get(k)
            x = self.ctx.get(str(d + timedelta(days=1))) if day else None
            item = {"run_date": k, "target_date": str(d + timedelta(days=1)), "has_history": day is not None,
                    "forecast_source": (PRIMARY if k in sel else SNAPSHOT) if day else None,
                    "forecast_run_id": day["forecast_run_id"] if day else None,
                    "snapshot": {f: bk[k][f] for f in ("snapshot_id", "snapshot_business_date", "snapshot_created_at",
                                                       "source_run_date", "session_started_at", "model_version")}
                    if k in bk else None,
                    "n_rows": (r or {}).get("n_rows", 0), "forecast": (r or {}).get("forecast"),
                    "actual": (r or {}).get("actual") if (r or {}).get("n_actual") else None,
                    "forecast_compared": (r or {}).get("forecast_compared"), "abs_error": (r or {}).get("abs_error"),
                    "actual_status": x["status"] if x else None, "actual_reason": x["reason"] if x else None,
                    "actual_as_of": x["as_of"] if x else None, "actual_source": x["source"] if x else None,
                    "actual_close": x.get("close") if x else None}
            item["deviation"] = (item["forecast_compared"] - item["actual"]) if item["actual"] is not None else None
            item["wape"] = (item["abs_error"] / item["actual"]) if item["actual"] else None
            out.append(item)
            d += timedelta(days=1)
        return out

    def kpi(self, daily: list[dict]) -> dict:
        """Totals of the selected rows. Actual, deviation and WAPE use ONLY the rows that have an actual (never a
        missing actual as 0); WAPE = sum|forecast - actual| / sum(actual) as in ml/metrics.py, undefined when that sum
        is 0. status: COMPLETE when every compared row is complete, PROVISIONAL when any is provisional, otherwise
        CLOSED (closed at 23:59, not confirmed by the source)."""
        f = sum(x["forecast"] or 0 for x in daily)
        comp = [x for x in daily if x["actual"] is not None]
        a = sum(x["actual"] for x in comp)
        fc = sum(x["forecast_compared"] or 0 for x in comp)
        err = sum(x["abs_error"] or 0 for x in comp)
        sts = {x["actual_status"] for x in comp}
        status = NO_ACTUAL if not comp else (COMPLETE if sts == {COMPLETE} else PROVISIONAL if PROVISIONAL in sts
                                             else CLOSED)
        notes = []
        if not comp:
            notes.append("NO_ACTUAL")
        elif a <= 0:
            notes.append("SUM_ACTUAL_ZERO")
        return {"n_rows": sum(x["n_rows"] for x in daily), "n_rows_compared": sum(x["n_rows"] for x in comp),
                "days_with_history": sum(1 for x in daily if x["has_history"]), "days_compared": len(comp),
                "days_primary": sum(1 for x in daily if x["forecast_source"] == PRIMARY),
                "days_backup": sum(1 for x in daily if x["forecast_source"] == SNAPSHOT),
                "forecast_total": f if any(x["has_history"] for x in daily) else None,
                "actual_total": a if comp else None, "forecast_compared": fc if comp else None,
                "deviation": (fc - a) if comp else None, "abs_error": err if comp else None,
                "wape": (err / a) if comp and a > 0 else None, "bias": ((fc - a) / a) if comp and a > 0 else None,
                "status": status, "notes": notes}

    def count(self) -> int:
        if not self.ctx:
            return 0
        sql, p = self._base()
        return self.c.execute(sql + " SELECT COUNT(*) FROM base", p).fetchone()[0]

    def rows(self, sort: str = "date", limit: int | None = 100, offset: int = 0):
        """Table rows (one per official SKU x FC x run date) with deviation and per-row WAPE (|f - a| / a, only when
        the actual is > 0)."""
        if sort not in SORTS:
            raise AppError("VALIDATION_ERROR", 422, detail=f"sort must be one of {', '.join(SORTS)}")
        if not self.ctx:
            return []
        sql, p = self._base()
        sql += (" SELECT b.*, CASE WHEN b.actual IS NOT NULL THEN b.yhat - b.actual END AS deviation, "
                "CASE WHEN b.actual > 0 THEN ABS(b.yhat - b.actual) / b.actual END AS row_wape FROM base b "
                f"ORDER BY {SORTS[sort]}")
        if limit is not None:
            sql += " LIMIT ? OFFSET ?"
            p = p + [int(limit), int(offset)]
        cur = self.c.execute(sql, p)
        cols = [x[0] for x in cur.description]
        for t in cur:
            r = dict(zip(cols, t))
            r["product_name"] = self.name_of(r["sku_id"], r["product_name"])
            r["actual_as_of"] = self.ctx[r["target_date"]]["as_of"]
            yield r

    def fcs(self) -> list[str]:
        return [x[0] for x in self.c.execute(
            "SELECT DISTINCT v.fc_id FROM forecast_history_days d JOIN forecast_history_values v "
            "ON v.forecast_run_id = d.forecast_run_id WHERE d.scope_key=? ORDER BY v.fc_id", (self.scope_key,))]


# ---------------------------------------------------------------- Excel
def export_xlsx(qy: Query, labels: dict, sort: str = "date") -> bytes:
    """All rows of the filter (never only the page on screen), streamed with openpyxl write-only. Refused (never cut)
    when the selection does not fit one sheet."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    n = qy.count()
    if n > MAX_EXPORT_ROWS:
        raise AppError("HISTORY_EXPORT_TOO_LARGE", 422, detail=f"{n} rows > {MAX_EXPORT_ROWS}")
    L = labels
    daily = qy.daily()
    k = qy.kpi(daily)
    wb = Workbook(write_only=True)

    def put(ws, values, date_cols=()):
        row = []
        for i, v in enumerate(values):
            if i in date_cols and v:
                cell = WriteOnlyCell(ws, value=date.fromisoformat(v))
                cell.number_format = L["excel_date_fmt"]
                row.append(cell)
            else:
                row.append(v)
        ws.append(row)

    st_label = lambda s: L["status"].get(s, s) if s else ""               # noqa: E731
    rs_label = lambda s: L["reason"].get(s, s) if s else ""               # noqa: E731
    src_label = lambda s: L["source"].get(s, s) if s else L["source"]["MISSING"]   # noqa: E731
    info = wb.create_sheet(L["sheet_info"])
    info.append([L["field"], L["value"]])
    f = qy.filters
    for key, val in (("export_format_version", "H2"), ("language", L["language_name"]),
                     ("generated_at_vn", L["fmt_time"](now_iso())), ("scope", qy.scope_key),
                     ("data_source", qy.data_source or ""), ("run_from", L["fmt_date"](f["run_from"])),
                     ("run_to", L["fmt_date"](f["run_to"])), ("fc_filter", f["fc_id"] or L["all"]),
                     ("search_filter", f["q"] or L["all"]),
                     ("status_filter", ", ".join(st_label(s) for s in f["actual_status"]) if f["actual_status"] else L["all"]),
                     ("rows", n), ("days_with_history", k["days_with_history"]),
                     ("forecast_total", k["forecast_total"]), ("actual_total", k["actual_total"]),
                     ("forecast_compared", k["forecast_compared"]), ("deviation", k["deviation"]),
                     ("wape", k["wape"]), ("kpi_status", st_label(k["status"])), ("rule_d1", L["rule_d1"]),
                     ("rule_actual", L["rule_actual"]), ("rule_wape", L["rule_wape"]),
                     ("source_filter", src_label(f["forecast_source"]) if f["forecast_source"] else L["all"]),
                     ("days_primary", k["days_primary"]), ("days_backup", k["days_backup"]),
                     ("rule_backup", L["rule_backup"]), ("rule_close", L["rule_close"])):
        info.append([L["info"][key], val])
    ses = wb.create_sheet(L["sheet_sessions"])
    ses.append(list(L["cols_sessions"].values()))
    for s in sessions(qy.c, qy.scope_key, qy.run_from, qy.run_to):
        put(ses, [s["run_date"], s["target_date"], L["fmt_time"](s["session_started_at"]), L["fmt_time"](s["published_at"]),
                  s["forecast_run_id"], s["model_version"], L["yes"] if s["official"] else L["no"], s["decision"],
                  s["decision_reason"], s["capture_mode"], L["fmt_time"](s["captured_at"]), s["n_rows"],
                  s["usage_status"] or ""], date_cols=(0, 1))
    dsh = wb.create_sheet(L["sheet_daily"])
    dsh.append(list(L["cols_daily"].values()))
    for x in daily:
        if not x["has_history"]:
            put(dsh, [x["run_date"], x["target_date"], None, None, None, None, None, L["no_history"], "",
                      src_label(None), None, None], date_cols=(0, 1))
            continue
        snap = x["snapshot"] or {}
        put(dsh, [x["run_date"], x["target_date"], x["forecast"], x["actual"], x["deviation"], x["wape"],
                  x["forecast_run_id"], st_label(x["actual_status"]), rs_label(x["actual_reason"]),
                  src_label(x["forecast_source"]), snap.get("snapshot_id"), snap.get("snapshot_business_date")],
            date_cols=(0, 1, 11))
    det = wb.create_sheet(L["sheet_rows"])
    det.append(list(L["cols_rows"].values()))
    for r in qy.rows(sort=sort, limit=None):
        put(det, [r["run_date"], r["target_date"], r["forecast_run_id"], r["model_version"],
                  L["fmt_time"](r["session_started_at"]), r["sku_id"], r["product_name"] or L["no_product_name"],
                  r["fc_id"], r["yhat"], r["actual"], r["deviation"], r["row_wape"], st_label(r["actual_status"]),
                  rs_label(r["actual_reason"]), L["fmt_time"](r["actual_as_of"]) if r["actual_as_of"] else "",
                  src_label(r["forecast_source"]), r["snapshot_id"], r["snapshot_business_date"],
                  L["fmt_time"](r["snapshot_created_at"]) if r["snapshot_created_at"] else ""],
            date_cols=(0, 1, 17))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
