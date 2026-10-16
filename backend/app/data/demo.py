"""Deterministic synthetic order data for DEMO mode (always labelled DEMO, never shown as production).

Covers: stable/seasonal demand, promotions, intermittent demand, new SKUs, discontinued SKUs, a new FC,
calendar days with no orders at all, late-arriving rows and a few price outliers.
"""
from __future__ import annotations

import io
import json
import logging
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd

from .. import db
from ..timeutil import VN_TZ, now_iso, today_vn
from .demand import aggregate_batch
from .validation import normalize_and_validate

log = logging.getLogger("df.demo")
FCS = ["FC-HCM", "FC-HN", "FC-DN"]


def _series_specs(n_skus: int, rng: np.random.Generator) -> list[dict]:
    specs = []
    for i in range(n_skus):
        sku = f"DEMO-SKU-{i + 1:03d}"
        if i < n_skus * 0.45:
            kind = "stable"
        elif i < n_skus * 0.65:
            kind = "intermittent"
        elif i < n_skus * 0.75:
            kind = "new"
        elif i < n_skus * 0.82:
            kind = "discontinued"
        elif i < n_skus * 0.88:
            kind = "short"
        else:
            kind = "seasonal"
        price = float(rng.choice([15000, 32000, 45000, 78000, 120000, 250000]))
        fcs = FCS if i % 4 else FCS[:2]
        for fc in fcs:
            specs.append({"sku": sku, "fc": fc, "kind": kind, "price": price,
                          "level": float(rng.uniform(4, 40)) * (1.3 if fc == "FC-HCM" else 1.0)})
    # a brand-new FC with two SKUs
    specs.append({"sku": "DEMO-SKU-001", "fc": "FC-CT", "kind": "new", "price": 45000, "level": 6.0})
    specs.append({"sku": "DEMO-SKU-002", "fc": "FC-CT", "kind": "new", "price": 32000, "level": 4.0})
    return specs


def _daily_qty(spec: dict, dates: pd.DatetimeIndex, end: date, rng: np.random.Generator):
    n = len(dates)
    t = np.arange(n)
    dow = dates.dayofweek.values
    weekly = np.where(dow >= 5, 0.7, 1.0) * np.where(dow == 0, 1.15, 1.0)
    yearly = 1 + 0.15 * np.sin(2 * np.pi * dates.dayofyear.values / 365.25)
    trend = 1 + 0.0004 * t
    promo = (rng.random(n) < 0.08).astype(int)
    lam = spec["level"] * weekly * trend * (1 + 0.5 * promo)
    kind = spec["kind"]
    if kind == "seasonal":
        lam = lam * (1 + 0.6 * np.sin(2 * np.pi * dates.dayofyear.values / 365.25))
    else:
        lam = lam * yearly
    q = rng.poisson(np.clip(lam, 0.01, None)).astype(float)
    if kind == "intermittent":
        q = np.where(rng.random(n) < 0.12, rng.integers(1, 6, n), 0).astype(float)
    start_idx = 0
    if kind == "new":
        start_idx = max(0, n - int(rng.integers(8, 21)))
    elif kind == "short":
        start_idx = max(0, n - int(rng.integers(35, 50)))
    q[:start_idx] = 0
    if kind == "discontinued":
        q[n - 90:] = 0
    return q, promo, start_idx


def generate_lines(end: date | None = None, days: int = 730, n_skus: int = 30, seed: int = 42) -> pd.DataFrame:
    """Order lines with raw-ish values (as an Excel/BigQuery source would provide)."""
    rng = np.random.default_rng(seed)
    end = end or (today_vn() - timedelta(days=1))
    dates = pd.date_range(end - timedelta(days=days - 1), end, freq="D")
    holidays = {dates[len(dates) // 3].date(), dates[len(dates) // 3 + 1].date()}  # days with no orders at all
    recs = []
    order_no = 0
    for spec in _series_specs(n_skus, rng):
        q, promo, start_idx = _daily_qty(spec, dates, end, rng)
        if spec["fc"] == "FC-CT":
            q[: len(q) - 10] = 0
        for i, d in enumerate(dates):
            if q[i] <= 0 or d.date() in holidays:
                continue
            n_lines = int(min(q[i], 1 + rng.poisson(1.5)))
            parts = rng.multinomial(int(q[i]), [1 / n_lines] * n_lines)
            for p in parts:
                if p <= 0:
                    continue
                order_no += 1
                minute = int(np.clip(rng.normal(13 * 60, 3 * 60), 7 * 60, 21 * 60 + 59))
                created = datetime(d.year, d.month, d.day, minute // 60, minute % 60, int(rng.integers(0, 60)))
                late = rng.random() < 0.003
                upd = created + (timedelta(days=3) if late else timedelta(minutes=int(rng.integers(1, 30))))
                price = spec["price"] * (12 if rng.random() < 0.0005 else 1)
                disc = round(price * p * 0.1, 0) if promo[i] else 0.0
                recs.append((f"DEMO-SO-{order_no:07d}", created, spec["sku"], spec["fc"], int(p), price, disc,
                             int(promo[i]), f"DEMO-PROMO-{d.month}" if promo[i] else None, 0, 0.0, upd,
                             f"DEMO-B{d:%Y%m%d}"))
    return pd.DataFrame(recs, columns=["order_id", "created_at", "sku_id", "fc_id", "order_qty", "unit_price",
                                       "discount_amount", "promotion_flag", "promotion_id", "voucher_flag",
                                       "voucher_amount", "source_updated_at", "batch_id"])


def create_demo_batch(days: int = 730, n_skus: int = 30, seed: int = 42) -> dict:
    """(Re)generate the DEMO dataset as a new ACTIVE DEMO batch; the previous DEMO batch is retired."""
    from .importer import _insert_lines, _purge_lines
    raw = generate_lines(days=days, n_skus=n_skus, seed=seed)
    raw.insert(0, "row_number", np.arange(1, len(raw) + 1))
    valid, issues, summary = normalize_and_validate(raw, grain_conflict_policy="ERROR")
    batch_id = db.new_id("DEMO")
    # One transaction, parent-first ordering:
    #   1. insert the new batch (not yet ACTIVE) + its lines + daily demand
    #   2. retire the previous ACTIVE DEMO batch, pointing replaced_by_batch_id at the row that now exists
    #   3. activate the new batch and verify invariants (exactly one ACTIVE DEMO, no FK violations)
    # Any failure rolls everything back, so the previous DEMO dataset stays usable.
    with db.tx() as c:
        olds = [r["id"] for r in db.rows(c, "SELECT id FROM import_batches WHERE data_source='DEMO' AND status='ACTIVE'")]
        c.execute(
            """INSERT INTO import_batches(id, data_source, status, file_name, total_rows, valid_rows, error_rows,
                 duplicate_rows, min_order_date, max_order_date, validation_json, created_at, note, replaces_batch_id)
               VALUES (?, 'DEMO', 'PENDING_REVIEW', ?, ?,?,?,?,?,?,?,?,?,?)""",
            (batch_id, f"DEMO synthetic data (seed={seed})", summary["total_rows"], summary["valid_rows"],
             summary["error_rows"], summary["duplicate_rows"], summary["min_order_date"], summary["max_order_date"],
             json.dumps(summary), now_iso(), "DEMO", olds[0] if len(olds) == 1 else None))
        _insert_lines(c, batch_id, valid)
        aggregate_batch(c, batch_id, "DEMO")
        for old_id in olds:
            _purge_lines(c, old_id)
            c.execute("UPDATE import_batches SET status='REPLACED', replaced_by_batch_id=?, closed_at=? WHERE id=?",
                      (batch_id, now_iso(), old_id))
        c.execute("UPDATE import_batches SET status='ACTIVE', confirmed_at=? WHERE id=?", (now_iso(), batch_id))
        n_active = c.execute("SELECT COUNT(*) FROM import_batches WHERE data_source='DEMO' AND status='ACTIVE'").fetchone()[0]
        if n_active != 1:
            raise db.ConsistencyError(f"expected exactly one ACTIVE DEMO batch, found {n_active}")
        bad = db.fk_violations(c, "import_batches")
        if bad:
            raise db.ConsistencyError(f"foreign key violations after DEMO generation: {bad[:5]}")
        db.audit(c, "DEMO_GENERATED", "import_batch", batch_id,
                 {"rows": summary["valid_rows"], "seed": seed, "replaced": olds})
    log.info("demo generated", extra={"ctx": {"batch_id": batch_id, "rows": summary["valid_rows"]}})
    return {"batch_id": batch_id, **summary}


def demo_intraday_lines(day: date, seed: int = 42) -> pd.DataFrame:
    """Synthetic order lines for `day` (used for DEMO D0 snapshots: cumulative up to the snapshot time)."""
    rng = np.random.default_rng(seed + day.toordinal())
    recs = []
    for spec in _series_specs(30, np.random.default_rng(seed)):
        if spec["kind"] in ("discontinued",):
            continue
        dow_f = 0.7 if day.weekday() >= 5 else 1.0
        total = rng.poisson(spec["level"] * dow_f) if spec["kind"] != "intermittent" else (
            int(rng.integers(1, 6)) if rng.random() < 0.12 else 0)
        for _ in range(int(total)):
            minute = int(np.clip(rng.normal(13 * 60, 3 * 60), 7 * 60, 21 * 60 + 59))
            created = datetime(day.year, day.month, day.day, minute // 60, minute % 60, tzinfo=VN_TZ)
            recs.append((spec["sku"], spec["fc"], 1.0, created))
    return pd.DataFrame(recs, columns=["sku_id", "fc_id", "order_qty", "created_at"])


def _xlsx(df: pd.DataFrame, instructions: list[str]) -> bytes:
    from openpyxl import load_workbook
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame({"Instructions": instructions}).to_excel(w, sheet_name="Instructions", index=False)
        df.to_excel(w, sheet_name="Data", index=False)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    ws = wb["Data"]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, datetime):
                c.number_format = "yyyy-mm-dd hh:mm:ss"
    ws.freeze_panes = "A2"
    wb.active = wb.sheetnames.index("Data")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def sample_history_xlsx(days: int = 180) -> bytes:
    """A VALID sample history file (subset of the demo generator) to try the import -> train flow."""
    df = generate_lines(days=days, n_skus=10, seed=7)
    df = df[df["fc_id"].isin(["FC-HCM", "FC-HN"])].reset_index(drop=True)
    df = df.rename(columns={"batch_id": "batch_id"})
    return _xlsx(df, ["SAMPLE DATA (synthetic) — dữ liệu mẫu tổng hợp, KHÔNG phải dữ liệu thật.",
                      "Dùng để thử quy trình import → kiểm tra → huấn luyện → dự báo.",
                      "Use it to try the import → validate → train → forecast flow."])


def sample_errors_xlsx() -> bytes:
    """A sample file with deliberate errors to try validation and the error report."""
    base = datetime(2026, 2, 1, 9, 0)
    good = [(f"SO-{i:05d}", base + timedelta(hours=i), "SKU-A", "FC-HCM", 5, 45000) for i in range(1, 21)]
    bad = [
        ("SO-90001", base, None, "FC-HCM", 3, 45000),                       # missing SKU
        ("SO-90002", base, "SKU-A", None, 3, 45000),                        # missing FC
        ("SO-90003", "31/02/2026 10:00", "SKU-A", "FC-HCM", 3, 45000),      # invalid date
        ("SO-90004", base, "SKU-A", "FC-HCM", -4, 45000),                   # negative qty
        ("SO-90005", base, "SKU-A", "FC-HCM", "abc", 45000),                # non numeric qty
        ("SO-90006", base, "SKU-A", "FC-HCM", None, 45000),                 # missing qty
        ("SO-00001", base + timedelta(hours=1), "SKU-A", "FC-HCM", 5, 45000),  # exact duplicate of SO-00001
        ("SO-90007", base, "SKU-B", "FC-HN", 2, 30000),                     # grain conflict pair
        ("SO-90007", base, "SKU-B", "FC-HN", 9, 30000),
        ("SO-90008", datetime(2099, 1, 1, 9, 0), "SKU-A", "FC-HCM", 1, 45000),  # future date
        ("SO-90009", base, "SKU-A", "FC-HCM", 0, 45000),                    # zero qty (valid)
        ("SO-90010", base, "SKU-A", "FC-HCM", 2, "giá?"),                   # invalid price -> blanked
        ("SO-90011", base, "SKU-A", "FC-HCM", 2, 4500000),                  # price outlier
    ]
    df = pd.DataFrame(good + bad, columns=["order_id", "created_at", "sku_id", "fc_id", "order_qty", "unit_price"])
    return _xlsx(df, ["SAMPLE WITH DELIBERATE ERRORS — file mẫu có lỗi cố ý để thử kiểm định dữ liệu.",
                      "Expected: missing SKU/FC, invalid date, negative/non-numeric/missing qty, duplicate, "
                      "grain conflict, future date, zero qty (valid), invalid price, price outlier."])
