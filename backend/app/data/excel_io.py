"""Safe .xlsx reading (no macros, no formula execution, zip-bomb guard), template and error report."""
from __future__ import annotations

import gc
import io
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..messages import CATALOG, AppError
from .schema import FIELDS

MAX_UNCOMPRESSED = 1024 * 1024 * 1024  # 1 GB


def check_xlsx(path: Path, original_name: str, max_mb: int) -> None:
    if not original_name.lower().endswith(".xlsx"):
        raise AppError("NOT_XLSX", 415, detail=original_name)
    size = path.stat().st_size
    if size == 0:
        raise AppError("EMPTY_FILE")
    if size > max_mb * 1024 * 1024:
        raise AppError("FILE_TOO_LARGE", 413, detail=f"{size / 1024 / 1024:.1f} MB > {max_mb} MB")
    if not zipfile.is_zipfile(path):
        raise AppError("CORRUPT_FILE", detail="not a zip/xlsx container")
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            if "[Content_Types].xml" not in names or not any(n.startswith("xl/workbook") for n in names):
                raise AppError("CORRUPT_FILE", detail="missing workbook parts")
            if any(n.lower().endswith("vbaproject.bin") for n in names):
                raise AppError("MACRO_NOT_ALLOWED")
            ct = z.read("[Content_Types].xml").decode("utf-8", "ignore")
            if "macroEnabled" in ct:
                raise AppError("MACRO_NOT_ALLOWED")
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED:
                raise AppError("FILE_TOO_LARGE", 413, detail="uncompressed content too large")
    except zipfile.BadZipFile as e:
        raise AppError("CORRUPT_FILE", detail=str(e)) from e


def _open(path: Path):
    try:
        return load_workbook(path, read_only=True, data_only=True, keep_vba=False, keep_links=False)
    except AppError:
        raise
    except Exception as e:  # openpyxl raises many types for broken files
        raise AppError("CORRUPT_FILE", detail=type(e).__name__) from e


def inspect_workbook(path: Path, sample_rows: int = 5) -> dict:
    wb = _open(path)
    try:
        sheets = []
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            header = None
            samples = []
            for row in it:
                if header is None:
                    if row and any(v is not None and str(v).strip() for v in row):
                        header = [None if v is None else str(v).strip() for v in row]
                    continue
                if len(samples) >= sample_rows:
                    break
                samples.append([_jsonable(v) for v in row[:len(header)]])
            sheets.append({"name": ws.title, "headers": [h for h in (header or []) if h],
                           "sample_rows": samples, "approx_rows": ws.max_row})
        return {"sheets": sheets, "suggested_sheet": pick_data_sheet([s["name"] for s in sheets])}
    finally:
        wb.close()


def pick_data_sheet(names: list[str]) -> str | None:
    for n in names:
        if n.strip().lower() == "data":
            return n
    non_instr = [n for n in names if n.strip().lower() not in ("instructions", "huong dan", "hướng dẫn", "fields")]
    if len(names) == 1:
        return names[0]
    if len(non_instr) == 1:
        return non_instr[0]
    return None


def _jsonable(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float) and np.isnan(v):
        return None
    return v


def read_sheet(path: Path, sheet: str, mapping: dict[str, str | None], max_rows: int) -> pd.DataFrame:
    """Read the mapped columns of `sheet` into a DataFrame with internal field names + row_number."""
    wb = _open(path)
    try:
        if sheet not in wb.sheetnames:
            raise AppError("SHEET_NOT_FOUND", detail=sheet)
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        header = None
        header_row = 0
        for row in it:
            header_row += 1
            if row and any(v is not None and str(v).strip() for v in row):
                header = [None if v is None else str(v).strip() for v in row]
                break
        if header is None:
            raise AppError("EMPTY_FILE", detail=sheet)
        idx = {}
        for field, colname in mapping.items():
            if not colname:
                continue
            if colname not in header:
                raise AppError("MISSING_REQUIRED_COLUMNS", detail=f"column '{colname}' not in sheet")
            idx[field] = header.index(colname)
        cols = {f: [] for f in idx}
        rownums = []
        rn = header_row
        for row in it:
            rn += 1
            if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                continue
            if len(rownums) >= max_rows:
                raise AppError("TOO_MANY_ROWS", 413, detail=f"> {max_rows}")
            rownums.append(rn)
            for f, i in idx.items():
                cols[f].append(row[i] if i < len(row) else None)
        df = pd.DataFrame({f: pd.Series(v, dtype=object) for f, v in cols.items()})
        df["row_number"] = rownums
        return df
    finally:
        wb.close()


def build_template(lang_note: bool = True, sample_days: int = 7) -> bytes:
    """Excel template: Instructions + Fields + Data (with example rows)."""
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    bold = Font(bold=True)
    lines = [
        ("HƯỚNG DẪN NHẬP DỮ LIỆU LỊCH SỬ ĐẶT HÀNG / HISTORICAL ORDER IMPORT GUIDE", True),
        ("", False),
        ("1. Nhập dữ liệu vào sheet 'Data'. Giữ nguyên dòng tiêu đề (dòng 1). / Enter data in the 'Data' sheet. Keep the header row (row 1).", False),
        ("2. Mỗi dòng = một dòng đơn hàng: đơn × SKU × FC. / Each row = one order line: order × SKU × FC.", False),
        ("3. Bắt buộc: created_at (hoặc order_date), sku_id, fc_id, order_qty. / Required: created_at (or order_date), sku_id, fc_id, order_qty.", False),
        ("4. order_qty = số lượng khách đặt lúc tạo đơn, KHÔNG trừ số lượng hủy. / order_qty = quantity ordered at creation, cancellations NOT subtracted.", False),
        ("5. Ngày giờ: YYYY-MM-DD HH:MM[:SS] hoặc DD/MM/YYYY HH:MM, hoặc ô kiểu Date của Excel. Không ghi múi giờ = giờ Việt Nam (UTC+7). / Date-time: YYYY-MM-DD HH:MM[:SS] or DD/MM/YYYY HH:MM or an Excel date cell. No offset = Vietnam time (UTC+7).", False),
        ("6. order_date có thể để trống: hệ thống lấy ngày của created_at theo giờ Việt Nam. / order_date may be blank: derived from created_at in Vietnam time.", False),
        ("7. sku_id, fc_id là văn bản (cột đã định dạng Text để giữ số 0 ở đầu). / sku_id, fc_id are text (columns formatted as Text to keep leading zeros).", False),
        ("8. Số lượng, giá: chỉ số, dấu chấm thập phân, không ghi đơn vị. Số lượng ≥ 0. / Quantities and prices: numbers only, dot decimal, no units. Quantity ≥ 0.", False),
        ("9. Cờ khuyến mãi/voucher: 1/0, TRUE/FALSE, Y/N, Có/Không. / Promotion/voucher flags: 1/0, TRUE/FALSE, Y/N.", False),
        ("10. Giá/giảm giá/khuyến mãi phải là thông tin đã biết tại thời điểm đặt hàng. / Price/discount/promotion must be known at order time.", False),
        ("11. Dòng trùng hoàn toàn chỉ được tính một lần. Cùng đơn × SKU × FC với giá trị khác nhau sẽ bị báo lỗi (có thể đổi chính sách trong Cấu hình). / Exact duplicate rows count once; same order × SKU × FC with different values is reported as an error (policy configurable).", False),
        ("12. Chỉ chấp nhận .xlsx; không hỗ trợ CSV, .xls, .xlsm (macro). / Only .xlsx; CSV, .xls, .xlsm (macros) are not supported.", False),
        ("13. Có thể đổi tên cột — khi import bạn sẽ ánh xạ cột của mình sang trường nội bộ. / You may rename columns — map them to internal fields during import.", False),
        ("14. Xóa các dòng ví dụ trước khi nhập dữ liệu thật. / Delete the example rows before entering real data.", False),
    ]
    for i, (t, b) in enumerate(lines, start=1):
        c = ins.cell(row=i, column=1, value=t)
        if b:
            c.font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 160

    fs = wb.create_sheet("Fields")
    fs.append(["field", "required / bắt buộc", "type", "mô tả (VI)", "description (EN)", "example"])
    for c in fs[1]:
        c.font = bold
    for f in FIELDS:
        req = "YES (created_at or order_date)" if f["name"] in ("created_at", "order_date") else ("YES" if f["required"] else "no")
        fs.append([f["name"], req, f["kind"], f["vi"], f["en"], f["example"]])
    for col, w in zip("ABCDEF", (20, 30, 10, 80, 80, 22)):
        fs.column_dimensions[col].width = w

    ds = wb.create_sheet("Data")
    headers = [f["name"] for f in FIELDS]
    ds.append(headers)
    fill = PatternFill("solid", fgColor="FFF2CC")
    for i, c in enumerate(ds[1]):
        c.font = bold
        if headers[i] in ("created_at", "sku_id", "fc_id", "order_qty"):
            c.fill = fill
    base = datetime(2026, 1, 1, 8, 30)
    n = 0
    for d in range(sample_days):
        for k, (sku, fc, q, p) in enumerate([("SKU-0001", "FC-HCM", 12, 45000), ("SKU-0002", "FC-HCM", 3, 120000),
                                             ("SKU-0001", "FC-HN", 7, 45000)]):
            n += 1
            ts = base + timedelta(days=d, hours=k * 2)
            ds.append([f"SO-{n:06d}", ts, None, sku, fc, q, p, 0 if d % 3 else 5000, 1 if d % 3 == 0 else 0,
                       "PROMO-T1" if d % 3 == 0 else None, 0, 0, ts + timedelta(minutes=5), f"B{ts:%Y%m%d}"])
    for row in ds.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd hh:mm:ss"
        row[12].number_format = "yyyy-mm-dd hh:mm:ss"
    for i, h in enumerate(headers, start=1):
        ds.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 4)
    for col in ("A", "D", "E", "J", "N"):
        for r in range(1, 2000):
            ds[f"{col}{r}"].number_format = "@"
    ds.freeze_panes = "A2"
    wb.active = wb.index(ds)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_error_report(batch: dict, issues: list[dict], lang: str = "vi") -> bytes:
    """Import error report in the UI language. Dates DD/MM/YYYY; raw values of date columns that are valid ISO
    dates/times are shown as DD/MM/YYYY too, anything unparseable is kept exactly as found in the file."""
    from ..export_i18n import OTHER, fmt_date, fmt_dateish, labels
    lang, L = labels(lang)
    O = OTHER[lang]
    wb = Workbook()
    s = wb.active
    s.title = O["err_sheet_summary"]
    for k in ("id", "file_name", "sheet_name", "status", "total_rows", "valid_rows", "error_rows", "duplicate_rows",
              "min_order_date", "max_order_date", "created_at"):
        v = batch.get(k)
        v = fmt_date(v, L) if k in ("min_order_date", "max_order_date") and v else fmt_dateish(v, L)
        s.append([O["err_summary"][k], v])
    s.column_dimensions["A"].width = 28
    s.column_dimensions["B"].width = 50
    e = wb.create_sheet(O["err_sheet_errors"])
    e.append(O["err_cols"])
    for c in e[1]:
        c.font = Font(bold=True)
    for it in issues:
        vi, en = CATALOG.get(it["code"], (it["code"], it["code"]))
        raw = it["raw_value"]
        if it["field"] in ("created_at", "order_date"):
            raw = fmt_dateish(raw, L)
        e.append([it["row_number"], it["field"], it["code"], it["severity"], raw, vi if lang == "vi" else en])
    for col, w in zip("ABCDEF", (10, 18, 26, 12, 30, 90)):
        e.column_dimensions[col].width = w
    for row in e.iter_rows(min_row=2):
        row[5].alignment = Alignment(wrap_text=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def dataframe_to_xlsx(sheets: dict[str, pd.DataFrame], date_format: str | None = None) -> bytes:
    """date_format (e.g. "DD/MM/YYYY") is the Excel display format of date cells; the values stay real dates.
    1.2.13: the workbook (one object per cell, linked both ways) is cyclic garbage once written; it is collected here,
    not whenever Python's full collection happens to run (forecast export of production: +1.0 GB while building,
    0.67 GB still held after it returned - measured on a copy, 22/10/2026)."""
    content = _dataframe_to_xlsx(sheets, date_format)
    gc.collect()
    return content


def _dataframe_to_xlsx(sheets: dict[str, pd.DataFrame], date_format: str | None) -> bytes:
    import datetime as _dt
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name[:31], index=False)
        if date_format:
            for ws in w.book.worksheets:
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        if isinstance(cell.value, (_dt.date, _dt.datetime)):
                            cell.number_format = date_format
    return buf.getvalue()
